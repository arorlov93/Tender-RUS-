"""Оценка по каждой рассмотренной закупке (20.10.2026): экономика, надёжность расчёта, риски, итоговое решение и приоритет."""
import io, contextlib, json, sys, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
sys.argv = [sys.argv[0], "data/verify_tubes_gorka_2010.json"]
with contextlib.redirect_stdout(io.StringIO()):
    import go_list as g
    import future_1910 as f

T = {x["num"]: x for x in json.load(open("data/triage_1910.json"))}
L = json.load(open("data/go_links_2010.json"))
pct = lambda v: None if v is None else (v / 100 if v > 1 else v)
E = {}


def put(num, **k):
    if num in E: return
    t = T.get(num, {})
    k.setdefault("cust", t.get("cust", "")); k.setdefault("subj", t.get("name", ""))
    k.setdefault("end", (t.get("end") or "")); k.setdefault("law", f'{t.get("law", "")} · {t.get("method", "")}')
    E[num] = dict(num=num, **k)


SKIP = {"32616453117": "Салаватводоканал — решили не идти", "32616452870": "Салаватводоканал — решили не идти"}
for r in g.ROWS:
    net = g.model(r["nmck"] * (1 - r["d"]), r["C"], r.get("days", 60))
    put(r["num"], cust=r["cust"], subj=r["subj"], end=r["end"], law=r["law"], nmck=r["nmck"], C=r["C"], d=r["d"], net=net, conf="высокая (построчно)" if r["num"] in ("32616346943", "B1510261835045") else "средняя–высокая",
        risk=r.get("risk", ""), src="расчёт по позициям")
for r in g.EST:
    put(r["num"], cust=r["cust"], subj=r["subj"], end=r["end"], law=r["law"], nmck=r["nmck"], C=r["C"], d=r["d"], net=g.model(r["nmck"] * (1 - r["d"]), r["C"]), conf="низкая (оценка)", risk="нужна спецификация", src="оценка по категории")
for fn in ("data/scan/scan_out_1.json", "data/scan/scan_out_2.json", "data/scan/scan_out_3.json", "data/big/big_other.json", "data/big/big_rosseti.json"):
    for x in json.load(open(fn)):
        if not x.get("nmck"): continue
        C = (x.get("cost_total") or 0) + (x.get("logistics") or 0)
        put(x["num"], nmck=x["nmck"], C=C or None, d=x.get("exp_disc") or 0.1, net=x.get("net_expected"), conf={"low": "низкая (оценка)", "med": "средняя", "high": "высокая"}.get(x.get("confidence"), "—"),
            risk=str(x.get("risks") or "")[:200], src="расчёт агента", dec0=x.get("decision", ""))
for r in f.out:
    if r["net"] is not None and r["nmck"]:
        C = r["C"]
        if r["cat"].startswith("Медрасход") and r["num"] in g.U:  # пробирки — по проверенным ценам МиниМед
            C = sum(g.tube_cost(p["name"]) * (p.get("qty") or 0) for p in g.U[r["num"]])
        hard = "Мособлэнерго" in r["cust"] or r["num"] == "32616425363"
        put(r["num"], nmck=r["nmck"], C=C, d=r["d"], net=g.model(r["nmck"] * (1 - r["d"]), C), conf="средняя", src="расчёт по позициям",
            risk="жёсткие условия: рамочный договор, логотип, требования к капиталу и гарантиям" if hard else "", dec0="Осторожно" if hard else "")
now = dt.datetime(2026, 10, 20, 12, 0)
rows = []
for e in E.values():
    m = 1 - e["C"] / e["nmck"] if e.get("C") and e.get("nmck") else None
    try: end = dt.datetime.strptime(e["end"][:16], "%d.%m.%Y %H:%M")
    except Exception:
        try: end = dt.datetime.strptime(e["end"][:16], "%Y-%m-%d %H:%M")
        except Exception: end = None
    dec0 = str(e.get("dec0", ""))
    if e["num"] in SKIP: dec, pr = SKIP[e["num"]], "—"
    elif end and end < now: dec, pr = "Срок подачи прошёл", "—"
    elif dec0.startswith("Нет") or "аукцион" in dec0.lower(): dec, pr = "Не идём: " + dec0[:60], "—"
    elif m is None or e.get("net") is None: dec, pr = "Нет расчёта", "C"
    elif e["net"] <= 0 or m < 0.10: dec, pr = "Не идём — нет маржи", "—"
    elif e["conf"].startswith("низкая"): dec, pr = "Запросить спецификацию / КП, потом решать", "B" if m >= 0.25 else "C"
    elif m >= 0.30 and e["net"] >= 100000: dec, pr = "ИДЁМ", "A"
    elif m >= 0.20: dec, pr = "Идём, если подтвердятся цены", "B"
    else: dec, pr = "Осторожно — малая маржа", "C"
    if dec0 == "Осторожно" and pr == "A": dec, pr = "Осторожно — жёсткие условия (рамочный, логотип, капитал)", "C"
    if e["num"] == "B1510261835045": dec, pr = "Идём только при подтверждённом наличии (10 дней)", "B"
    if e["num"] == "0373100001026000845": dec, pr = "Осторожно: дёшево только б/у Cisco; с новым маржа ~35%", "B"
    rows.append(dict(e, m=m, end_dt=end, dec=dec, pr=pr))
ORD = {"A": 0, "B": 1, "C": 2, "—": 3}
rows.sort(key=lambda r: (ORD[r["pr"]], -(r.get("net") or -1e12)))

thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; P = "0%"
FILL = {"A": "C6EFCE", "B": "FFEB9C", "C": "F2F2F2", "—": "FFC7CE"}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Оценка"
ws.append([f"Оценка по каждой рассмотренной закупке — 20.10.2026. Всего {len(rows)}. A — идём, B — идём после подтверждения цен/КП, C — под вопросом, «—» — не идём"]); ws["A1"].font = Font(bold=True, size=12)
ws.append(["Приоритет", "Решение", "Подача до", "Номер", "Заказчик", "Предмет", "Способ", "НМЦК, ₽", "Себестоимость, ₽", "Маржа при НМЦК", "Ожид. снижение", "Прибыль после УСН, ₽",
           "Надёжность расчёта", "Риски / комментарий", "Ссылка ЕИС", "Площадка"])
for c in ws[2]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
sh = lambda t, n: (t or "") if len(t or "") <= n else t[:n - 1] + "…"
for r in rows:
    lk = L.get(r["num"], {})
    ws.append([r["pr"], r["dec"], r["end_dt"].strftime("%d.%m.%Y %H:%M") if r["end_dt"] else r["end"], r["num"], sh(r["cust"], 50), sh(r["subj"], 80), sh(r["law"], 36),
               r.get("nmck"), r.get("C"), r["m"], r.get("d"), r.get("net"), r["conf"], sh(r.get("risk"), 170), lk.get("href") if "zakupki" in str(lk.get("href")) else "", lk.get("etp_url") or ""])
    X = ws.max_row; ws.cell(X, 4).number_format = "@"
    for j in (8, 9, 12): ws.cell(X, j).number_format = M
    for j in (10, 11): ws.cell(X, j).number_format = P
    ws.cell(X, 1).fill = PatternFill("solid", fgColor=FILL[r["pr"]]); ws.cell(X, 2).fill = PatternFill("solid", fgColor=FILL[r["pr"]])
    for j in (15, 16):
        v = ws.cell(X, j).value
        if v: ws.cell(X, j).hyperlink = v if str(v).startswith("http") else "https://" + str(v)
    for c in ws[X]: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([9, 30, 15, 20, 28, 36, 22, 14, 14, 9, 9, 14, 16, 48, 34, 34], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "E3"; ws.auto_filter.ref = f"A2:P{ws.max_row}"
tidy(wb, row_h=46)
wb.save("../EVAL_ALL_2010.xlsx")
for r in rows:
    if r["pr"] in ("A", "B"): print(r["pr"], r["num"], r["end"][:16], round(r.get("nmck") or 0), round((r["m"] or 0) * 100), round(r.get("net") or 0), r["dec"][:40], r["conf"][:10])
import collections; print(collections.Counter(r["pr"] for r in rows))
