"""АСТ ГОЗ 29.10.2026: оценка присланных закупок (закрытые аукционы 44-ФЗ и малые закупки 223-ФЗ) — условия, построчная себестоимость, маржа. → ../GOZ_2910.xlsx"""
import json, os, glob, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
pct = lambda v: None if v is None else (v / 100 if abs(v) > 1.5 else v)
D = []
for f in sorted(glob.glob("data/goz2910/*.json")):
    for x in json.load(open(f)): D.append(x)
D.sort(key=lambda x: -(x.get("net_10") or -1e9))
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); M = '#,##0" ₽"'
head = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row] if c.value is not None]
txt = lambda v: "; ".join(f"{k}: {w}" for k, w in v.items()) if isinstance(v, dict) else ("; ".join(map(str, v)) if isinstance(v, list) else str(v or ""))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Сводка"
ws.append(["АСТ ГОЗ — закупки от 29.10.2026: условия и потенциальная маржа (УСН 15%, закупка с НДС)"]); ws["A1"].font = Font(bold=True, size=12)
H = ["№", "Номер", "Способ", "Заказчик", "Предмет", "НМЦК, ₽", "Подача до", "Себестоимость + логистика, ₽", "Маржа при НМЦК", "Прибыль при −10%, ₽", "Прибыль при −20%, ₽", "Безубыточное снижение",
     "Обеспечение заявки / контракта", "Поставка", "Оплата", "Нацрежим", "Требования", "Решение", "Риски"]
ws.append(H); head(ws)
for i, x in enumerate(D, 1):
    ws.append([i, x["num"], x.get("method", "закрытый аукцион 44-ФЗ"), str(x.get("customer"))[:80], str(x.get("object"))[:90], x.get("nmck"), x.get("end"), x.get("cost_total"), pct(x.get("margin")),
               x.get("net_10"), x.get("net_20"), pct(x.get("breakeven")), f'{x.get("security_app")} / {x.get("security_contract")}', txt(x.get("delivery"))[:250], txt(x.get("payment"))[:150],
               txt(x.get("nat_regime"))[:200], txt(x.get("requirements"))[:300], txt(x.get("decision"))[:120], txt(x.get("risks"))[:400]])
    X = ws.max_row; ws.cell(X, 2).number_format = "@"
    for j in (6, 8, 10, 11): ws.cell(X, j).number_format = M
    for j in (9, 12): ws.cell(X, j).number_format = "0%"
    d = str(x.get("decision") or "")
    ws.cell(X, 18).fill = PatternFill("solid", fgColor="C6EFCE" if d.startswith(("Идти", "Участвовать", "Да")) else "F8CBAD" if d.startswith(("Не", "Нет")) else "FFEB9C")
    for c in ws[X]: c.border = BOX
for i, w in enumerate([4, 22, 18, 30, 34, 13, 15, 14, 9, 13, 13, 10, 18, 36, 26, 30, 36, 24, 50], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "F3"
ln = wb.create_sheet("Построчно")
ln.append(["Номер", "Позиция", "Кол-во", "Ед.", "НМЦ за ед., ₽", "Себестоимость за ед., ₽", "Маржа", "Источник цены"]); head(ln)
for x in D:
    for l in x.get("lines") or []:
        nu, cu = l.get("nmck_unit"), l.get("cost_unit")
        ln.append([x["num"], str(l.get("item"))[:120], l.get("qty"), l.get("unit"), nu, cu, (1 - cu / nu) if nu and cu else None, str(l.get("src"))[:250]])
        X = ln.max_row; ln.cell(X, 1).number_format = "@"; ln.cell(X, 5).number_format = M; ln.cell(X, 6).number_format = M; ln.cell(X, 7).number_format = "0%"
for i, w in enumerate([22, 60, 9, 8, 13, 13, 8, 70], 1): ln.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ln.freeze_panes = "B2"; ln.auto_filter.ref = f"A1:H{ln.max_row}"
tidy(wb, row_h=60); wb.save("../GOZ_2910.xlsx")
for x in D: print(x["num"], x.get("nmck"), pct(x.get("margin")), x.get("net_10"), str(x.get("decision"))[:40])
