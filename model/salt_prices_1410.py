"""Цены на техническую соль (открытые источники, 14.10.2026) и маржа лотов по сценариям себестоимости."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
def P(nm,c,d):
    p=nm*(1-d); pre=p-c-p*0.008-c*0.2*60/365; return pre-max(pre,0)*0.15
def be(nm,c):
    lo,hi=-0.5,1.0
    for _ in range(60):
        m=(lo+hi)/2
        if P(nm,c,m)>0: lo=m
        else: hi=m
    return lo
REF=[
("esolk.ru (ЕСК)","Павлодарсоль / Аралтуз (Казахстан), Уралкалий, Белкалий, Тыреть, АСК","Галит тип B/C, помол 2–3, навалом",1990,"без НДС, без доставки, от вагона (~70 т), станция отгрузки","https://www.esolk.ru/produktsiya/kontsentrat-mineralnyy-galit/mkr-1-tonna/"),
("nsmgr.ru (Люберцы)","без указания","Тип D, от 40 т, россыпь/биг-бэг",4350,"самовывоз, НДС не указан","https://nsmgr.ru/reagenty/tekhnicheskaya-sol/big-begi"),
("nsmgr.ru","без указания","Тип D, до 40 т, россыпь/биг-бэг",4550,"самовывоз","https://nsmgr.ru/reagenty/tekhnicheskaya-sol/big-begi"),
("nsmgr.ru","Илецкая (Руссоль)","Тип C, россыпь",5750,"самовывоз","https://nsmgr.ru/reagenty/tekhnicheskaya-sol/big-begi"),
("nsmgr.ru","без указания","Тип C, россыпь/биг-бэг",6300,"самовывоз","https://nsmgr.ru/reagenty/tekhnicheskaya-sol/big-begi"),
("nsmgr.ru","Уралкалий / Белкалий","Тип A с антислеживателем, биг-бэг",6500,"самовывоз","https://nsmgr.ru/reagenty/tekhnicheskaya-sol/tip-a/uralkaliy-big-beg-s-antislezhivatelem"),
("sknerud.ru (Москва)","Астрахань","МКР 1 т",6500,"«от», доставка отдельно","https://sknerud.ru/reagent/sol-texnicheskaya/"),
("sknerud.ru","Аралтуз / Индертуз / Алтайсоль (КЗ)","МКР 1 т",7000,"«от», доставка отдельно","https://sknerud.ru/reagent/sol-texnicheskaya/"),
("sknerud.ru","Беларуськалий / Баскунчак","МКР 1 т",7500,"«от», доставка отдельно","https://sknerud.ru/reagent/sol-texnicheskaya/"),
("sknerud.ru","Уралкалий","МКР 1 т",8500,"«от», доставка отдельно","https://sknerud.ru/reagent/sol-texnicheskaya/"),
("sknerud.ru","Тыреть / Илецксоль","МКР 1 т",10000,"«от», доставка отдельно","https://sknerud.ru/reagent/sol-texnicheskaya/"),
("perevalka-tsk.ru (Москва)","Аралтуз (КЗ)","навалом",7300,"самовывоз; доставка 8–11 тыс. за машину 20–25 м³ до 60 км","https://perevalka-tsk.ru/product/pg-materialy/tekhnicheskaya-sol/tekhnicheskaya-sol-araltuz/"),
("perevalka-tsk.ru","Баскунчак","навалом / биг-бэг",8700,"самовывоз","https://perevalka-tsk.ru/product/pg-materialy/tekhnicheskaya-sol/tekhnicheskaya-sol-araltuz/"),
("vektor-e.ru (Москва)","Аралтуз (КЗ)","МКР / россыпь, от 2 т",7800,"с НДС, без доставки","https://vektor-e.ru/sol/"),
("vektor-e.ru","Баскунчак","МКР / россыпь",8000,"с НДС, без доставки","https://vektor-e.ru/sol/"),
("vektor-e.ru","Уралкалий","МКР / россыпь",8500,"с НДС, без доставки","https://vektor-e.ru/sol/"),
("vektor-e.ru","Илецкая","МКР / россыпь",10000,"с НДС, без доставки","https://vektor-e.ru/sol/"),
("perevalka24.ru","Уралкалий «Галит» марка А","биг-бэг 1000 кг",7000,"«от»","https://perevalka24.ru/product/peskosol/kontsentrat-mineralnyy-galit-marki-a-big-beg-1000kg-1-tonna/"),
("opt6.ru","—","МКР, помол 2 тип Б",9990,"«от»","https://opt6.ru/products/tekhnicheskaya_sol_mkr/"),
("Руссоль-центр","Руссоль","мешки/МКР, от 10 т",11300,"опт, по данным выдачи поиска (сайт 403)","https://www.xn----jtbqkblfcapd5b2e.xn--p1ai/catalog/k-6070503-sol_tekhnicheskaya"),
("тендер Алтайский край (saby.ru)","—","НМЦК 2 430 т = 13,37 млн",5500,"начальная цена заказчика, ориентир","https://saby.ru/tenders/card/245190387"),
("Ваша цена","Казахстан","МКР, НДС вкл., с доставкой в РФ",2500,"данные клиента — проверить, куда именно доставка","—"),
]
# сценарии себестоимости, ₽/т с доставкой до заказчика
SC=[("A: ваша КЗ 2 500",2500),("B: опт тип D + доставка ≈5 000",5000),("C: 4 500 (не подтверждено)",4500),("D: дилер МКР + доставка ≈8 000",8000)]
LOTS=[("ДГХ, Московская обл.","32616451106","16.10",6724000,400,"223-ФЗ"),
("Новосибирск","0851300006226000106","16.10",4400000,400,"44-ФЗ"),
("Волоколамск, МО","32616447519","16.10",1593800,130,"223-ФЗ"),
("Абинск, Краснодарский край","0318300065126000256","19.10",1368000,150,"44-ФЗ"),
("Воронеж","0131300000626001316","16.10",3296150,461,"44-ФЗ"),
("Струнино, Владимирская обл.","0128600006826000407","19.10",1342500,150,"44-ФЗ"),
("Школа № 525, СПб (мешки 20–25 кг)","0372200193426000027","16.10",106470,6.5,"44-ФЗ"),
("Самара (нужен опыт 24,2 млн)","0142200001326018965","19.10",121000000,11000,"44-ФЗ"),]
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Маржа по сценариям"
ws.append(["Соль: прибыль лотов при разной себестоимости (снижение на торгах 15%, чистая после УСН 15%)"]); ws["A1"].font=Font(bold=True,size=13)
h=["Лот","№","Подача","НМЦК, ₽","Тонн","Цена в НМЦК, ₽/т"]
for n,_ in SC: h+= [f"{n}: маржа",f"{n}: прибыль, ₽",f"{n}: безубыток"]
ws.append(h)
for c in ws[2]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True,vertical="center")
tot=[0]*len(SC)
for name,num,end,nm,t,law in LOTS:
    row=[name,num,end,nm,t,nm/t]
    for i,(_,cpt) in enumerate(SC):
        C=t*cpt; pe=P(nm,C,0.15); row+=[(nm-C)/nm,pe,be(nm,C)]
        if "Самара" not in name: tot[i]+=pe
    ws.append(row); R=ws.max_row
    ws.cell(R,2).number_format="@"; ws.cell(R,4).number_format=MONEY; ws.cell(R,6).number_format=MONEY
    for i in range(len(SC)):
        b=7+3*i; ws.cell(R,b).number_format="0%"; ws.cell(R,b+1).number_format=MONEY; ws.cell(R,b+2).number_format="0%"
        v=ws.cell(R,b+1).value
        ws.cell(R,b+1).fill=PatternFill("solid",fgColor="C6EFCE" if v>300000 else ("FFEB9C" if v>0 else "FFC7CE"))
ws.append(["ИТОГО без Самары"]+[""]*5+sum([["",t_,""] for t_ in tot],[]))
R=ws.max_row
for i in range(len(SC)): ws.cell(R,8+3*i).number_format=MONEY
for c in ws[R]: c.font=Font(bold=True)
for i,w in enumerate([30,21,8,15,8,12]+[10,15,10]*len(SC),1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="C3"
for row in ws.iter_rows(min_row=3):
    for c in row: c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
w2=wb.create_sheet("Цены на соль (источники)")
w2.append(["Источник","Производитель","Продукт / фасовка","Цена от, ₽/т","Условия","Ссылка"])
for c in w2[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
for r in sorted(REF,key=lambda x:x[3]):
    w2.append(list(r)); R=w2.max_row; w2.cell(R,4).number_format=MONEY
    if r[5].startswith("http"): w2.cell(R,6).hyperlink=r[5]; w2.cell(R,6).font=Font(color="0563C1",underline="single")
for i,w in enumerate([26,40,34,13,52,60],1): w2.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
for row in w2.iter_rows():
    for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
w2.auto_filter.ref=f"A1:F{w2.max_row}"
from tidy import tidy; tidy(wb)
wb.save("../SALT_PRICES_1410.xlsx")
print([round(x) for x in tot])
for name,num,end,nm,t,law in LOTS: print(name, round(nm/t), [round(P(nm,t*c,0.15)) for _,c in SC])
