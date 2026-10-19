"""Крупные лоты 19.10: Камчатскэнерго, Россети МР (кабель), прочие ≥3 млн — маржа, прибыль, решение, позиции."""
import json, re, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

L = []
for f in ("kamch", "rosseti", "other"):
    L += json.load(open(f"data/big/big_{f}.json"))
pct = lambda v: None if v is None else (v / 100 if v > 1 else v)
ORD = {"Участвовать": 0, "Осторожно": 1, "Запросить": 2, "Нет": 3}
L.sort(key=lambda x: (next(v for k, v in ORD.items() if x["decision"].startswith(k)), -(x["nmck"] or 0)))
FILL = {"Участвовать": "C6EFCE", "Осторожно": "FFEB9C", "Запросить": "DDEBF7", "Нет": "FFC7CE"}
CONF = {"high": "высокая", "med": "средняя", "low": "низкая"}
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; P = "0.0%"
short = lambda t, n: (t or "") if len(t or "") <= n else t[:n - 1] + "…"


def head(ws):
    for c in ws[ws.max_row]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")


wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Крупные лоты"
ws.append(["Крупные лоты (подача с 20.10.2026, без аукционов) — маржа и решение"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Прибыль — после УСН 15% при ожидаемом снижении, с логистикой и стоимостью денег. Надёжность «низкая» = спецификация недоступна (ЕИС/площадка закрыты), оценка по категории/истории."])
ws.append(["№", "Номер", "Заказчик", "Предмет", "Способ", "Подача", "НМЦК, ₽", "Себестоимость, ₽", "Логистика, ₽", "Маржа при НМЦК", "Ожид. снижение",
           "Прибыль после УСН, ₽", "Безубыточное снижение", "Решение", "Надёжность", "Риски / условия"]); head(ws)
for i, x in enumerate(L, 1):
    ws.append([i, x["num"], short(x.get("cust"), 50), short(x.get("name"), 80), short(x.get("method"), 30), x.get("end"), x["nmck"], x.get("cost_total"), x.get("logistics"),
               pct(x.get("margin_pct")), x.get("exp_disc"), x.get("net_expected"), x.get("breakeven_disc"), x["decision"], CONF.get(x.get("confidence"), x.get("confidence")),
               short(x.get("risks"), 150)])
    R = ws.max_row; ws.cell(R, 2).number_format = "@"
    for j in (7, 8, 9, 12): ws.cell(R, j).number_format = M
    for j in (10, 11, 13): ws.cell(R, j).number_format = P
    ws.cell(R, 14).fill = PatternFill("solid", fgColor=next(v for k, v in FILL.items() if x["decision"].startswith(k)))
    for c in ws[R]: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([4, 17, 26, 40, 18, 15, 15, 15, 12, 9, 9, 15, 11, 18, 10, 70], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "D4"

ps = wb.create_sheet("Позиции")
ps.append(["Лот", "Позиция", "Кол-во", "Ед.", "НМЦК за ед., ₽", "Себес за ед., ₽", "Источник цены"]); head(ps)
for x in L:
    for p in (x.get("positions") or []):
        if not (p.get("item") and (p.get("qty") or p.get("cost_unit"))): continue
        ps.append([x["num"], short(p.get("item"), 90), p.get("qty"), p.get("unit"), p.get("nmck_unit"), p.get("cost_unit"), short(p.get("cost_src"), 140)])
        ps.cell(ps.max_row, 1).number_format = "@"
        for j in (5, 6): ps.cell(ps.max_row, j).number_format = M
for row in ps.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([17, 50, 10, 8, 14, 14, 70], 1): ps.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

R0 = next(x for x in L if x["num"] == "32616480459")
rs = wb.create_sheet("Россети — чувствительность")
rs.append(["Закупка у завода (K = доля от цены дилера)", "Себестоимость; маржа; прибыль при НМЦК; прибыль при −10%; безубыточное снижение"]); head(rs)
for k, v in (R0.get("sensitivity_cost_factor") or {}).items():
    rs.append([k, v])
rs.append([]); rs.append(["Потребность в капитале", json.dumps(R0.get("capital_need"), ensure_ascii=False)[:600]])
rs.append(["Металлы", json.dumps(R0.get("metal_prices"), ensure_ascii=False)[:400]])
for i, w in enumerate([34, 40, 60], 1): rs.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ks = wb.create_sheet("Камчатка — история")
K = next(x for x in L if x["num"] == "32616346943")
ks.append(["Прошлые лоты Камчатскэнерго (арматура)"]); head(ks)
for h in K.get("history") or []:
    ks.append([h if isinstance(h, str) else "; ".join(f"{k}: {v}" for k, v in h.items())])
ks.append([]); ks.append(["Барьеры: " + str(K.get("barriers") or "")])
ks.column_dimensions["A"].width = 160

src = wb.create_sheet("Источники")
src.append(["Лот", "Ссылка"]); head(src)
for x in L:
    for u in (x.get("sources") or [])[:12]:
        src.append([x["num"], u]); src.cell(src.max_row, 1).number_format = "@"
src.column_dimensions["A"].width = 17; src.column_dimensions["B"].width = 120
tidy(wb, row_h=60)
wb.save("../BIG_LOTS_1910.xlsx")
print(len(L), "лотов")
