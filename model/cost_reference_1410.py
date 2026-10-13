"""Справочник: средняя себестоимость с доставкой — одежда (сайты РФ) и соль (по регионам)."""
import json, statistics as st, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
TA="https://files.technoavia.ru/files/price/11/price_Technoavia_moscow_2025_06_27_08_46_34.pdf"
IVU="https://www.ivuniforma.ru/catalog/rabochaya-odezhda/"; SPT="https://www.spets.ru/products/spetsodezhda/"; TR="https://www.trakt.ru/"
# товар: вес кг, [(сайт, производитель/город, цена от, цена до, тип: опт|прайс|розница, url)]
G={
"Костюм летний рабочий":(1.2,[("ivuniforma.ru","Иваново, фабрика",938,1715,"опт",IVU+"letnyaya-spetsodezhda/kostumy/"),("ivsoteks.ru","Иваново",1600,2490,"опт","https://ivsoteks.ru/specodejda"),("technoavia.ru","Техноавиа",2321,3065,"прайс",TA),("trakt.ru","Тракт",3086,3746,"прайс",TR+"product/11954/"),("spets.ru","Спецобъединение",3367,4984,"прайс",SPT+"letnyaya/")]),
"Костюм летний спасателя/МЧС":(1.3,[("productcenter.ru","«Аврора», Барнаул",2490,2490,"опт","https://productcenter.ru/products/45461/kostiumy-dlia-sotrudnikov-mchs"),("voen36.ru","Воен36, Воронеж (собств. пр-во)",3100,5995,"розница","https://voen36.ru/catalog/forma_mchs/"),("barrakuda.biz","Барракуда, Москва",4900,6675,"розница","https://barrakuda.biz/formennaya-odezhda/kostyumy-mchs")]),
"Костюм зимний рабочий (I–III пояс)":(2.6,[("fabrikant-textile.ru","Фабрикант, Иваново",2040,2460,"опт","https://fabrikant-textile.ru/catalog/spetsodezhda/"),("specteks37.ru","СпецТекс 37, Иваново",2650,4600,"опт","https://www.specteks37.ru/katalog/zimnyaya-spetsodezhda/"),("spec-standart.ru","Стандарт Спецодежда, Иваново",2570,4968,"опт","https://spec-standart.ru/zimnyaya-specodezhda"),("ivsoteks.ru","Ивсотекс, Иваново",2900,5830,"опт","https://ivsoteks.ru/specodejda"),("ivuniforma.ru","Иваново, фабрика",3549,5810,"опт",IVU+"zimnjaja/kostyumy-1/"),("trakt.ru","Тракт",3588,7793,"прайс",TR+"product/12066/"),("spets.ru","Спецобъединение",6084,9394,"прайс",SPT+"odezhda_zimnyaya/"),("technoavia.ru","Техноавиа",8262,9158,"прайс",TA)]),
"Костюм зимний северный (IV/особый пояс)":(3.2,[("specteks37.ru","СпецТекс 37 №265 (IV пояс)",4600,4600,"опт","https://www.specteks37.ru/katalog/zimnyaya-spetsodezhda/"),("spec-standart.ru","«Север» (тройка), Иваново",8543,8543,"опт","https://spec-standart.ru/zimnyaya-specodezhda"),("trakt.ru","Тракт «Беринг» / «Северный шельф»",9659,23998,"прайс",TR+"product/13086/"),("spets.ru","Спецобъединение «Надым» / «Ураган»",9760,15250,"прайс",SPT+"odezhda_zimnyaya/")]),
"Комбинезон летний":(1.1,[("ivuniforma.ru","Иваново, фабрика",3766,4102,"опт",IVU+"letnyaya-spetsodezhda/kombinezony/")]),
"Комбинезон зимний":(2.4,[("spets.ru","Спецобъединение «Пилот»",10199,10577,"прайс",SPT+"odezhda_zimnyaya/")]),
"Куртка утепленная":(1.6,[("fabrikant-textile.ru","Фабрикант, Иваново",1125,1695,"опт","https://fabrikant-textile.ru/catalog/spetsodezhda/"),("specteks37.ru","СпецТекс 37",1200,2620,"опт","https://www.specteks37.ru/katalog/zimnyaya-spetsodezhda/"),("ivuniforma.ru","Иваново, фабрика",1764,4480,"опт",IVU+"zimnjaja/"),("technoavia.ru","Техноавиа",2501,4669,"прайс",TA),("trakt.ru","Тракт (медиана)",4995,4995,"прайс",TR)]),
"Полукомбинезон / брюки зимние":(1.0,[("fabrikant-textile.ru","Фабрикант",720,765,"опт","https://fabrikant-textile.ru/catalog/spetsodezhda/"),("spec-standart.ru","Стандарт Спецодежда",958,958,"опт","https://spec-standart.ru/zimnyaya-specodezhda"),("trakt.ru","Тракт (медиана)",2799,2799,"прайс",TR),("technoavia.ru","Техноавиа",3190,3750,"прайс",TA)]),
"Футболка":(0.2,[("productcenter.ru","фабрики Иваново",211,270,"опт","https://productcenter.ru/products/r-ivanovskaia-obl-168/c-ivanovo-2550/catalog-muzhskiie-futbolki-1062"),("technoavia.ru","Техноавиа",478,478,"прайс",TA),("trakt.ru","Тракт",484,484,"прайс",TR+"product/12012/"),("spets.ru","Спецобъединение",488,488,"прайс",SPT+"trikotazh/")]),
"Халат рабочий":(0.5,[("fabrikant-textile.ru","Фабрикант",400,400,"опт","https://fabrikant-textile.ru/catalog/spetsodezhda/"),("ivuniforma.ru","Иваново, фабрика",609,1064,"опт",IVU+"letnyaya-spetsodezhda/khalati/"),("technoavia.ru","Техноавиа",1051,1466,"прайс",TA)]),
"Жилет утепленный":(0.7,[("ivuniforma.ru","Иваново, фабрика",1190,1372,"опт",IVU+"zimnjaja/"),("trakt.ru","Тракт",1866,1866,"прайс",TR+"product/12021/"),("spets.ru","Спецобъединение «Сатурн»",2788,2788,"прайс",SPT+"odezhda_dlya_okhrannykh_struktur/")]),
"Жилет сигнальный":(0.15,[("fabrikant-textile.ru","Фабрикант",260,320,"опт","https://fabrikant-textile.ru/catalog/spetsodezhda/"),("trakt.ru","Тракт",316,316,"прайс",TR+"product/138206/"),("spets.ru","Спецобъединение",390,677,"прайс",SPT+"odezhda_spetsialnaya/"),("technoavia.ru","Техноавиа",390,753,"прайс",TA)]),
"Ботинки кожаные с МП (летние)":(1.6,[("technoavia.ru","Техноавиа «Билд»",3018,3351,"прайс",TA),("trakt.ru","Тракт (медиана 42 моделей)",3416,3416,"прайс",TR)]),
"Ботинки утепленные":(1.9,[("trakt.ru","Тракт «Металлан»…«Тягач»",2174,4579,"прайс",TR+"product/12102/"),("technoavia.ru","Техноавиа «Билд» утепл.",3775,3775,"прайс",TA),("trakt.ru","Тракт «Алтай» натур. мех",5511,5511,"прайс",TR+"product/12598/")]),
"Сапоги ПВХ":(1.4,[("technoavia.ru","Техноавиа",620,1184,"прайс",TA),("spets.ru","Спецобъединение «Ямал» (утепл.)",1122,1122,"прайс",SPT)]),
"Костюм сварщика":(3.0,[("technoavia.ru","Техноавиа КС01/КС21",3801,7020,"прайс",TA),("trakt.ru","Тракт (медиана)",6531,6531,"прайс",TR),("ivuniforma.ru","Иваново, утепл. со спилком",6993,6993,"опт",IVU+"zimnjaja/kostyumy-1/")]),
"Костюм от электродуги":(2.0,[("trakt.ru","Тракт летний (медиана) / зимний",17008,33155,"прайс",TR+"product/61043/"),("technoavia.ru","Техноавиа «Вольт»",24195,24195,"прайс",TA)]),
"Костюм от нефтепродуктов":(2.0,[("trakt.ru","Тракт летний / «Пластойл» зимний",6731,11253,"прайс",TR+"product/12720/"),("spets.ru","Спецобъединение «Мегастатойл»",13219,24199,"прайс",SPT+"odezhda_spetsialnaya/")]),
"Термобельё (комплект)":(0.4,[("trakt.ru","Тракт",1428,2892,"прайс",TR+"product/12922/")]),
"Перчатки х/б с ПВХ (пара)":(0.08,[("technoavia.ru","Техноавиа «Гриф»…«Техно»",34,47,"прайс",TA)]),
"Рукавицы (пара)":(0.15,[("technoavia.ru","Техноавиа",88,149,"прайс",TA),("trakt.ru","Тракт (медиана)",140,140,"прайс",TR)]),
"Каска":(0.4,[("technoavia.ru","Техноавиа СОМЗ-55",303,303,"прайс",TA),("trakt.ru","Тракт (медиана)",488,488,"прайс",TR)]),
}
CLIENT={"Костюм летний рабочий":3500,"Костюм летний спасателя/МЧС":3500,"Костюм зимний рабочий (I–III пояс)":5000}
bench={b["item"]:b for b in json.load(open("data/deep/bench_clothes.json"))}
TEN={"Костюм летний рабочий":"костюм летний","Костюм летний спасателя/МЧС":"костюм летний","Костюм зимний рабочий (I–III пояс)":"костюм зимний","Костюм зимний северный (IV/особый пояс)":"костюм зимний",
     "Куртка утепленная":"куртка утепленная","Полукомбинезон / брюки зимние":"полукомбинезон","Футболка":"футболка","Халат рабочий":"халат","Жилет сигнальный":"жилет сигнальный",
     "Ботинки кожаные с МП (летние)":"ботинки кожаные с МП","Ботинки утепленные":"сапоги утепленные","Сапоги ПВХ":"сапоги ПВХ","Костюм сварщика":"костюм сварщика","Костюм от электродуги":"костюм от ЭДП",
     "Термобельё (комплект)":"термобелье","Перчатки х/б с ПВХ (пара)":"перчатки х/б с ПВХ","Рукавицы (пара)":"рукавицы","Каска":"каска"}
DLV={"Центр (Москва/МО, ЦФО)":15,"Регионы (Урал, Юг, Сибирь)":35,"Крайний Север / Якутия / ДВ":120}  # ₽/кг ТК [оценка]
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Одежда — средняя себест."
ws.append(["Одежда: средняя себестоимость с доставкой по открытым ценам российских фабрик и оптовиков (14.10.2026)"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["Средняя = среднее по серединам диапазонов всех сайтов. Опт = среднее только по фабричным/оптовым ценам (если есть). Доставка ТК: центр 15, регионы 35, Север 120 ₽/кг × вес [оценка]."])
H=["Товар","Сайтов","Мин., ₽","Макс., ₽","Средняя по рынку, ₽","Средняя опт/фабрика, ₽","Себестоимость для расчёта, ₽","Вес, кг"]+[f"С доставкой: {k}, ₽" for k in DLV]+["Ваша себестоимость, ₽","Медиана НМЦК 44-ФЗ, ₽","Наценка к медиане НМЦК (центр)"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True,vertical="center")
for g,(w,rows) in G.items():
    mids=[(a+b)/2 for _,_,a,b,_,_ in rows]; opt=[(a+b)/2 for _,_,a,b,t,_ in rows if t=="опт"]
    avg=st.mean(mids); aopt=st.mean(opt) if opt else None
    base=CLIENT.get(g) or (avg*0.9 if g.startswith('Полукомбинезон') else (min(aopt,avg*0.9) if aopt else avg*0.9))
    dl=[base+w*r for r in DLV.values()]
    tm=bench.get(TEN.get(g,""),{}).get("tender_median")
    ws.append([g,len({r[0] for r in rows}),min(r[2] for r in rows),max(r[3] for r in rows),avg,aopt,base,w]+dl+[CLIENT.get(g),tm,(tm/dl[0]-1) if tm else None])
    R=ws.max_row
    for k in list(range(3,8))+[9,10,11,12,13]: ws.cell(R,k).number_format=MONEY
    ws.cell(R,14).number_format="0%"
    if tm: ws.cell(R,14).fill=PatternFill("solid",fgColor="C6EFCE" if tm/dl[0]-1>0.3 else ("FFEB9C" if tm/dl[0]-1>0 else "FFC7CE"))
for i,w in enumerate([34,7,10,10,13,13,14,7,14,14,14,13,13,12],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
for row in ws.iter_rows(min_row=4):
    for c in row: c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
ws.freeze_panes="B4"; ws.auto_filter.ref=f"A3:N{ws.max_row}"
w2=wb.create_sheet("Одежда — цены по сайтам")
w2.append(["Товар","Сайт","Производитель / город","Цена от, ₽","Цена до, ₽","Тип цены","Ссылка"])
for c in w2[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
for g,(w,rows) in G.items():
    for s,m,a,b,t,u in sorted(rows,key=lambda r:r[2]):
        w2.append([g,s,m,a,b,t,u]); R=w2.max_row; w2.cell(R,4).number_format=MONEY; w2.cell(R,5).number_format=MONEY
        w2.cell(R,7).hyperlink=u; w2.cell(R,7).font=Font(color="0563C1",underline="single")
for i,w in enumerate([34,20,36,11,11,10,70],1): w2.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
w2.auto_filter.ref=f"A1:G{w2.max_row}"; w2.freeze_panes="A2"
# соль по регионам
import importlib.util, sys, io
spec=importlib.util.spec_from_file_location("ss","salt_strategy_1410.py"); ss=importlib.util.module_from_spec(spec); _o=sys.stdout; sys.stdout=io.StringIO(); spec.loader.exec_module(ss); sys.stdout=_o
w3=wb.create_sheet("Соль — средняя себест.")
w3.append(["Соль техническая: себестоимость с доставкой до заказчика, ₽/т"]); w3["A1"].font=Font(bold=True,size=13)
w3.append(["Средняя реальная = среднее вариантов 2 (завод вагоном) и 3 (склад авто). Ваша КЗ-цена 2 500 — требует подтверждения точки доставки."])
w3.append(["Лот / регион","Тонн","Цена в НМЦК, ₽/т","1. Казахстан (ваша)","2. Ближайший завод вагоном","3. Ближайший склад авто","4. Дилер с доставкой","Средняя реальная (2–3), ₽/т","Источник (ближайший)"])
for c in w3[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True)
for L in ss.LOTS:
    v=ss.variants(L); real=[c for _,c,_ in v[1:3] if c]
    w3.append([L[0],L[4],L[3]/L[4]]+[c for _,c,_ in v]+[st.mean(real),L[7][0]]); R=w3.max_row
    for k in range(3,9): w3.cell(R,k).number_format=MONEY
for i,w in enumerate([30,8,13,13,15,15,13,16,46],1): w3.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
w4=wb.create_sheet("Соль — цены по сайтам")
w4.append(["Источник","Производитель","Продукт / фасовка","Цена от, ₽/т","Условия","Ссылка"])
for c in w4[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
for r in sorted(ss.REF,key=lambda x:x[3]):
    w4.append(list(r)); R=w4.max_row; w4.cell(R,4).number_format=MONEY
    if r[5].startswith("http"): w4.cell(R,6).hyperlink=r[5]; w4.cell(R,6).font=Font(color="0563C1",underline="single")
for i,w in enumerate([26,40,34,13,52,60],1): w4.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
for wsx in (w2,w3,w4):
    for row in wsx.iter_rows():
        for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
from tidy import tidy; tidy(wb)
wb.save("../COST_REFERENCE_1410.xlsx")
for row in ws.iter_rows(min_row=4,values_only=True): print(row[0],'|',round(row[4]),'|',row[5] and round(row[5]),'|',round(row[6]),'|',[round(x) for x in row[8:11]],'|',row[12],'|',row[13] and f"{row[13]:.0%}")
print(len({s for _,(w,r) in G.items() for s,*_ in r}),"сайтов по одежде")
