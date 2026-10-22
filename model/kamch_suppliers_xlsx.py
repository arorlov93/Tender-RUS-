"""Камчатскэнерго 32616346943 — отдельная таблица: заводы для запроса КП по группам + позиции по группам."""
import json, re, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from tidy import tidy
from kamch_groups import group

Q = json.load(open("data/kamch/kam_qty.json"))
SUP = {x["group"]: x["suppliers"] for x in json.load(open("data/kamch/suppliers.json"))}
P = {}
for g in "ABC":
    for x in json.load(open(f"data/kamch/price_{g}.json")):
        P[(re.search(r"1\.\d", x["spec"]).group(0), x["n"])] = x
GV = {}
for q in Q:
    g = group(q["name"]); p = P.get((q["table"], q["n"]), {})
    d = GV.setdefault(g, dict(pos=[], val=0.0)); d["pos"].append((q["table"], q["n"], q["name"], q["qty"], p.get("price")))
    d["val"] += (p.get("price") or 0) * (q["qty"] or 0)
ORDER = sorted(GV, key=lambda g: -GV[g]["val"])
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'
def head(ws):
    for c in ws[ws.max_row]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
def box(ws, r0):
    for row in ws.iter_rows(min_row=r0):
        for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)

wb = openpyxl.Workbook()
a = wb.active; a.title = "Обзвон — топ-2"
a.append(["Камчатскэнерго 32616346943 — кому звонить первым (по группам, от большей суммы). Подача 29.10.2026 10:00 МСК"]); a["A1"].font = Font(bold=True, size=12)
a.append(["Что просить: КП на позиции группы (лист «Позиции по группам»), номер записи в реестре Минпромторга (ПП 719) по каждой позиции, сертификаты ТР ТС 032/2013 и 010/2011, срок поставки до Петропавловска-Камчатского"])
a.append(["№", "Группа", "Позиций", "Сумма по рынку, ₽", "Поставщик", "Телефон", "E-mail", "Сайт", "Звонок: результат", "Цена КП, ₽", "Срок", "Реестр МПТ (№)"]); head(a)
for i, g in enumerate(ORDER, 1):
    for j, s in enumerate(SUP.get(g, [])[:2]):
        a.append([i if j == 0 else "", g if j == 0 else "", len(GV[g]["pos"]) if j == 0 else "", GV[g]["val"] if j == 0 else "", s.get("company"), (s.get("phone") + (f"\n[{s['checked']}]" if s.get("checked") else "")), s.get("email"), s.get("site"), "", "", "", ""])
        if j == 0: a.cell(a.max_row, 4).number_format = M
box(a, 4)
for i, w in enumerate([4, 32, 8, 14, 32, 22, 26, 20, 22, 14, 12, 16], 1): a.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
a.freeze_panes = "C4"

ks = wb.create_sheet("Все поставщики")
ks.append(["Группа", "Позиций", "Сумма по рынку, ₽", "Поставщик", "Город", "Телефон", "E-mail", "Сайт", "Что закрывает", "Реестр МПТ", "Срок / склад", "Примечание"]); head(ks)
for g in ORDER:
    for i, sp in enumerate(SUP.get(g, [{}])):
        ks.append([g if i == 0 else "", len(GV[g]["pos"]) if i == 0 else "", GV[g]["val"] if i == 0 else "", sp.get("company"), sp.get("city"), (sp.get("phone") + (f"\n[{sp['checked']}]" if sp.get("checked") else "")), sp.get("email"), sp.get("site"),
                   str(sp.get("covers") or "")[:110], str(sp.get("registry") or "").split(";")[0].split("(")[0][:45], f'{sp.get("lead_time") or ""}; {sp.get("stock") or ""}'[:90], str(sp.get("note") or "")[:90]])
        if i == 0: ks.cell(ks.max_row, 3).number_format = M
box(ks, 2)
for i, w in enumerate([28, 7, 13, 28, 13, 20, 24, 18, 46, 22, 32, 32], 1): ks.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ks.freeze_panes = "D2"

kp = wb.create_sheet("Позиции по группам")
kp.append(["Группа", "Таблица ТТ", "№", "Позиция", "Кол-во", "Наша оценка цены за ед., ₽"]); head(kp)
for g in ORDER:
    for t, n, nm, q, pr in GV[g]["pos"]:
        kp.append([g, t, n, nm, q, pr]); kp.cell(kp.max_row, 6).number_format = M
box(kp, 2)
for i, w in enumerate([34, 9, 5, 70, 8, 16], 1): kp.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
kp.auto_filter.ref = f"A1:F{kp.max_row}"; kp.freeze_panes = "D2"
tidy(wb, row_h=34)
for r in range(2, ks.max_row + 1): ks.row_dimensions[r].height = 66
wb.save("../KAMCHATKA_SUPPLIERS_KP.xlsx")
print(len(ORDER), "групп")
