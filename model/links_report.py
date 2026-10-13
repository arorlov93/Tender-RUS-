"""Отчёт по проверке ссылок-источников цен (14.10.2026)."""
import json, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
d=json.load(open("data/link_check.json"))
OV={ # ручная проверка (WebFetch) — перекрывает автомат
"esolk.ru":("работает","есть (от 1 990 ₽/т)","+7 (343) 288-79-88, 8-800-707-08-35, info@esolk.ru","НЕ подтверждена: даты цены нет, сайт © 2013–2019","Не опираться без КП. На этой цене построен вариант «завод вагоном» по соли — запросить КП"),
"fabrikant-textile.ru":("работает","есть","8-800-200-41-30, sales@fabrikant-textile.ru","ПОДТВЕРЖДЕНА: прайс-лист от 07.10.2026","Можно использовать"),
"vektor-e.ru":("работает","есть (7 800–10 000 ₽/т)","+7 (499) 403-19-06, info@vektor-e.ru, Москва, Вагоноремонтная 10А","не указана, сайт © 2005–2023","Сомнительно — уточнить по телефону"),
"spec-standart.ru":("работает","есть","+7 (4932) 26-33-55, info@spec-standart.ru (Иваново)","не указана, сайт © 2025; есть прайс /price.xlsx","Уточнить актуальность прайса"),
"specteks37.ru":("работает","есть (цены «от»)","+7 (493) 293-97-90 (Иваново); e-mail не указан","не указана, сайт © 2019–2026","Живой каталог 2026 — использовать с проверкой"),
"opt6.ru":("открывается, но цена не отображается","нет (загружается скриптом)","не найдены","—","Исключить из источников"),
"perevalka24.ru":("НЕ работает (ошибка 502)","—","—","—","Исключить из источников"),
"files.technoavia.ru":("НЕ открывается (таймаут)","—","—","УСТАРЕЛ: прайс от 27.06.2025","Исключить; запросить прайс 2026 у Техноавиа"),
"xn----jtbqkblfcapd5b2e.xn--p1ai":("НЕ открывается (403)","—","—","—","Исключить (Руссоль-центр)"),
"saby.ru":("закрыт для автоматической проверки","—","—","тендер 2026","Не прайс — ориентир НМЦК заказчика"),
"bicotender.ru":("работает","—","—","тендер 2026","Не прайс — карточка тендера Мосгортура"),
}
def dom(u): return u.split("//")[1].split("/")[0].replace("www.","")
rows=[]
for x in d:
    D=dom(x["url"])
    if D in OV:
        st,pr,ct,act,ver=OV[D]
    else:
        ok=x["code"]=="200"
        st="работает" if ok else f"НЕ работает ({x['code']})"
        pr="есть" if x["prices"]>0 else "не найдены автоматически"
        ct=", ".join(v for v in (x["phone"],x["email"]) if v) or "не найдены"
        act=("сайт обновлён в 2026 (© 2026), дата цены не указана" if "2026" in x["years"] else f"год на сайте: {x['years'] or '—'}")
        ver="Можно использовать" if ok and x["prices"]>0 and "2026" in x["years"] and (x["phone"] or x["email"]) else ("Проверить: нет контактов на странице" if ok else "Проверить")
    rows.append((D,x["url"],x["used"],st,pr,ct,act,ver))
rows.sort(key=lambda r:(0 if r[7].startswith("Исключ") or r[7].startswith("Не опир") else (1 if r[7].startswith(("Сомн","Уточн","Провер","Живой")) else 2),r[0]))
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Проверка ссылок"
ws.append(["Проверка источников цен (14.10.2026): доступность, цены, контакты, актуальность"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["Автопроверка всех страниц + ручная проверка сомнительных. «© 2026» = сайт живой в 2026, но явной даты цены нет (кроме Фабриканта)."])
H=["Сайт","Ссылка","Где использована цена","Ссылка работает","Цены на странице","Контакты","Актуальность цен","Вывод"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
thin=Side(style="thin",color="BFBFBF")
for r in rows:
    ws.append(list(r)); R=ws.max_row
    ws.cell(R,2).hyperlink=r[1]; ws.cell(R,2).font=Font(color="0563C1",underline="single")
    col="FFC7CE" if r[7].startswith(("Исключ","Не опир")) else ("FFEB9C" if r[7].startswith(("Сомн","Уточн","Провер","Живой","Не прайс")) else "C6EFCE")
    ws.cell(R,8).fill=PatternFill("solid",fgColor=col)
    for cc in ws[R]: cc.border=Border(top=thin,bottom=thin,left=thin,right=thin)
for i,w in enumerate([20,50,36,22,22,40,38,40],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="B4"; ws.auto_filter.ref=f"A3:H{ws.max_row}"
from tidy import tidy; tidy(wb,row_h=30)
wb.save("../LINKS_CHECK_1410.xlsx")
import collections; print(collections.Counter(r[7].split(";")[0][:22] for r in rows)); print(len(rows))
