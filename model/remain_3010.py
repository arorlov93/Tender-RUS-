"""Финал на 30.10.2026: что осталось закрыть по позициям — Камчатка (32616346943, подача 02.11) и Сахалин лот 2 (11202, подача 05.11). → ../OSTALOS_ZAKRYT_3010.xlsx"""
import json, re, runpy, io, contextlib, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
with contextlib.redirect_stdout(io.StringIO()):
    rows = runpy.run_path("kamch_status.py")["rows"]; FIN = runpy.run_path("kamch_final.py")["SEL"]
REG = json.load(open("data/kamch/registry.json"))
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); M = '#,##0" ₽"'
head = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row] if c.value is not None]
COL = {"Нет КП": "DDEBF7", "КП дорого / не по ТЗ": "F8CBAD", "Торг": "FFEB9C", "Подтвердить условие": "E2EFDA", "Нет № реестра": "FCE4D6", "Не закрыто — разъяснения": "C00000"}
wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Сводка"
sv.append(["Камчатскэнерго 32616346943 — что осталось закрыть по позициям (на 30.10.2026, подача 02.11.2026 10:00 МСК)" if "kamch" in __import__("sys").argv else "Что осталось закрыть по позициям — на 30.10.2026"]); sv["A1"].font = Font(bold=True, size=13)
sv.append(["Закупка", "Что", "Позиций", "Сумма (наша цена / оценка), ₽", "Что сделать", "Срок"]); head(sv)
K = []
for d in sorted(rows, key=lambda d: (d["tt"].split("/")[0], int(d["tt"].split("/")[1]))):
    f = FIN.get(d["tt"], {}); num = REG.get(d["tt"], {}).get("num")
    if f.get("src") == "НЕ ЗАКРЫТО": st = "Не закрыто — разъяснения"
    elif d["act"] == "Не отработано": st = "Нет КП"
    elif d["act"] == "Заказать у других": st = "КП дорого / не по ТЗ"
    elif d["act"] == "Торг: наш бюджет": st = "Торг"
    elif d["act"] == "Брать при условии": st = "Подтвердить условие"
    elif not num: st = "Нет № реестра"
    else: continue
    K.append(dict(tt=d["tt"], name=d["name"], q=d["q"], st=st, our=d["our"], inv=d["inv"], sup=d["sup"], who=f.get("who") or "", cont=f.get("contact") or "",
                  ask=(f.get("ask") or d["why"] or "")[:250], num=num or ""))
for st, w in [("Нет КП", "ответы заводов (KAMCHATKA_FINAL_ZAVODY)"), ("КП дорого / не по ТЗ", "заводы с реестром / торг"), ("Торг", "письма KAMCHATKA_TORG_2910"),
              ("Подтвердить условие", "письменные подтверждения поставщиков"), ("Нет № реестра", "выписки ГИСП (ДН.РУ — изготовитель)"), ("Не закрыто — разъяснения", "запрос разъяснений заказчику / эквивалент")]:
    g = [x for x in K if x["st"] == st]
    sv.append(["Камчатка", st, len(g), sum(x["our"] * x["q"] for x in g), w, "до 01.11 (подача 02.11 10:00)"]); sv.cell(sv.max_row, 4).number_format = M; sv.cell(sv.max_row, 2).fill = PatternFill("solid", fgColor=COL[st])
ws = wb.create_sheet("Камчатка")
ws.append(["Табл./№", "Позиция ТЗ", "Кол-во", "Что осталось", "Наша цена, ₽/шт", "Цена КП, ₽/шт", "Сумма по нашей цене, ₽", "Текущий поставщик", "Кому писать / завод", "Контакт", "№ реестра", "Что именно закрыть"]); head(ws)
for x in K:
    ws.append([x["tt"], x["name"], x["q"], x["st"], x["our"], x["inv"], x["our"] * x["q"], x["sup"], x["who"], x["cont"], x["num"], x["ask"]]); X = ws.max_row
    for j in (5, 6, 7): ws.cell(X, j).number_format = M
    ws.cell(X, 4).fill = PatternFill("solid", fgColor=COL[x["st"]])
    if x["st"].startswith("Не закрыто"): ws.cell(X, 4).font = Font(color="FFFFFF", bold=True)
    for c in ws[X]: c.border = BOX
for i, w in enumerate([8, 40, 7, 22, 13, 13, 15, 16, 34, 30, 11, 60], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:L{ws.max_row}"
import sys
ONLY_K = "kamch" in sys.argv
# Сахалин лот 2
if not ONLY_K:
    P = json.load(open("data/sakh/nmc_positions.json"))["11202"]; PR = {(str(y["part"]), y["n"]): y for y in json.load(open("data/sakh/sakh_prices.json"))["11202"]}
    AS = {(x["part"], x["n"]): x for x in json.load(open("data/sakh/assign.json"))["11202"]}
    FAC = {x["company"]: x for x in json.load(open("data/sakh/factories.json"))}
    ph = lambda s: str(s or "").split(";")[0].strip()
    S = []
    for p in P:
        y = PR[(p["part"][-1], p["n"])]; a = AS.get((p["part"], p["n"]), {})
        if y.get("kp"):
            if 1 - y["cost_unit"] / p["nmc_unit"] < 0.05: st = "Торг"; why = f'КП есть, маржа {1 - y["cost_unit"] / p["nmc_unit"]:.0%}: {y.get("kp_note") or ""}'
            else: continue
        elif not a.get("main"): st = "Не закрыто — разъяснения"; why = a.get("why", "")
        else: st = "Нет КП"; why = "цена — " + ("оценка" if "ОЦЕНКА" in str(y.get("src", "")).upper() else "публичный прайс") + "; запрос: SAKHALIN_ZAPROS"
        f = FAC.get(a.get("main") or "", {})
        S.append(dict(n=f'{p["part"][-1]}/{p["n"]}', name=p["name"], q=p["qty"], st=st, nmc=p["nmc_unit"], cost=y["cost_unit"], who=a.get("main") or "", cont=f'{ph(f.get("phone"))} {ph(f.get("email"))}'.strip(), why=why[:250]))
    for st in ("Нет КП", "Торг", "Не закрыто — разъяснения"):
        g = [x for x in S if x["st"] == st]
        sv.append(["Сахалин лот 2", st, len(g), sum(x["cost"] * x["q"] for x in g), {"Нет КП": "ответы заводов (SAKHALIN_ZAPROS)", "Торг": "АДЛ Гранлок DN100/150 и др.", "Не закрыто — разъяснения": "Flamco ×2, Valtec — запрос разъяснений"}[st], "до 03.11 (подача 05.11 10:00)"])
        sv.cell(sv.max_row, 4).number_format = M; sv.cell(sv.max_row, 2).fill = PatternFill("solid", fgColor=COL[st])
for row in sv.iter_rows(min_row=2):
    for c in row:
        if c.value is not None: c.border = BOX
for i, w in enumerate([16, 28, 9, 22, 46, 28], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
if not ONLY_K:
    s2 = wb.create_sheet("Сахалин лот 2")
    s2.append(["Партия/№", "Позиция ТТ", "Кол-во", "Что осталось", "НМЦ за ед., ₽", "Себестоимость сейчас, ₽/шт", "Сумма НМЦ, ₽", "Кому писать / завод", "Контакт", "Комментарий"]); head(s2)
    for x in sorted(S, key=lambda x: -x["nmc"] * x["q"]):
        s2.append([x["n"], x["name"], x["q"], x["st"], x["nmc"], x["cost"], x["nmc"] * x["q"], x["who"], x["cont"], x["why"]]); X = s2.max_row
        for j in (5, 6, 7): s2.cell(X, j).number_format = M
        s2.cell(X, 4).fill = PatternFill("solid", fgColor=COL[x["st"]])
        if x["st"].startswith("Не закрыто"): s2.cell(X, 4).font = Font(color="FFFFFF", bold=True)
        for c in s2[X]: c.border = BOX
    for i, w in enumerate([8, 50, 7, 22, 13, 13, 14, 34, 30, 60], 1): s2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    s2.freeze_panes = "C2"; s2.auto_filter.ref = f"A1:J{s2.max_row}"
tidy(wb, row_h=34); wb.save("../KAMCHATKA_OSTALOS_3010.xlsx" if ONLY_K else "../OSTALOS_ZAKRYT_3010.xlsx")
import collections
print(collections.Counter(x["st"] for x in K))
