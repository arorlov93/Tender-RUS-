"""Камчатка 32616346943: «прожать» поставщиков — целевые цены по каждой позиции их счетов и письма на торг.
Цель: группа 2 (дороже > +10% к нашей цене) — наш бюджет (наша цена −10%); группа 3 (0…+10%) — наша цена −5%; группа 1 (уже ниже нашей цены) — скидка за объём −5%.
Не по ТЗ — замена на товар по ТЗ. → ../KAMCHATKA_TORG_2910.xlsx"""
import runpy, io, contextlib, re, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

with contextlib.redirect_stdout(io.StringIO()):
    m = runpy.run_path("kamch_scheta.py")
q, grp = m["q"], m["grp"]
def target(d):
    g = grp(d)
    if any(f.startswith("BAD") for f in d["flags"]): return None, "замена на товар строго по ТЗ: " + "; ".join(f[4:] for f in d["flags"] if f.startswith("BAD"))
    if g == 2: return round(d["our"] * 0.9, -1), "цена выше рынка: прошу до бюджета"
    if g == 3: return round(min(d["inv"], d["our"] * 0.95), -1), "небольшое превышение: прошу снизить"
    return round(d["inv"] * 0.95, -1), "скидка за объём заказа"
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); M = '#,##0" ₽"'
head = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row] if c.value is not None]
COL = {1: "C6EFCE", 2: "F8CBAD", 3: "FFEB9C"}
wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Сводка торга"
sv.append(["Камчатскэнерго 32616346943 — торг с поставщиками по счетам (29.10.2026)"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["Поставщик", "Позиций в счёте", "Сумма счёта, ₽", "Целевая сумма, ₽", "Экономия, ₽", "Экономия, %", "Позиций — замена (не по ТЗ)", "Главное в торге"]); head(sv)
LET = {}
for s in sorted({d["sup"] for d in q}):
    G = sorted([d for d in q if d["sup"] == s], key=lambda d: (-grp(d) if grp(d) != 1 else 0, -(d["inv"] * d["q"])))
    a = sum(d["inv"] * d["q"] for d in G); t = 0; bad = 0
    ws = wb.create_sheet(re.sub(r"[^\wА-Яа-я ]", "", s)[:28])
    ws.append([f"{s}: позиции для торга — наша целевая цена"]); ws["A1"].font = Font(bold=True, size=12)
    ws.append(["№", "Табл./№ ТЗ", "Позиция ТЗ", "Кол-во", "Что в счёте", "Цена в счёте, ₽/шт", "Целевая цена, ₽/шт", "Снижение", "Сумма в счёте, ₽", "Целевая сумма, ₽", "Обоснование", "Группа"]); head(ws)
    lines = []
    for i, d in enumerate(G, 1):
        tg, why = target(d)
        if tg is None: bad += 1; tsum = 0
        else: tsum = tg * d["q"]; t += tsum
        ws.append([i, d["tt"], d["name"], d["q"], d["offer"][:120], d["inv"], tg, (tg / d["inv"] - 1) if tg else None, d["inv"] * d["q"], tsum or None, why, {1: "1 точно есть", 2: "2 выше бюджета", 3: "3 продать 7–10%"}[grp(d)]])
        X = ws.max_row
        for j in (6, 7, 9, 10): ws.cell(X, j).number_format = M
        ws.cell(X, 8).number_format = "+0%;-0%"; ws.cell(X, 12).fill = PatternFill("solid", fgColor=COL[grp(d)])
        for c in ws[X]: c.border = BOX
        lines.append(f'{d["tt"]} {d["name"]} — {d["q"]:g} шт: ' + (f'{tg:,.0f} ₽/шт (в счёте {d["inv"]:,.0f})'.replace(",", " ") if tg else "нужна замена по ТЗ"))
    ok = [d for d in G if target(d)[0] is not None]; a_ok = sum(d["inv"] * d["q"] for d in ok)
    ws.append(["", "", "ИТОГО", "", "", None, None, None, a, t]); [setattr(c, "font", Font(bold=True)) for c in ws[ws.max_row]]
    for j in (9, 10): ws.cell(ws.max_row, j).number_format = M
    for i, w in enumerate([4, 8, 34, 7, 44, 13, 13, 9, 14, 14, 40, 14], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "D3"
    main = {"ДН.РУ": "30с64нж и 15с52нж11 в 2–6 раз дороже рынка; заказ 14+ млн — скидка за объём; № реестра ПП 719 по всем позициям",
            "Тулаэлектропривод": "Н-В-02 на 57% дороже — цель 117 тыс.", "ТД АДЛ (Гранвэл)": "скидка за объём; Ду65/80 до нашей цены",
            "Лидер-Энерго": "уже дёшево — скидка 5% за объём", "КИПИКА (прайс 01.10)": "цена от 51 шт уже минимальная — −5% за объём"}.get(s, "")
    sv.append([s, len(G), a, t, a_ok - t, (a_ok - t) / a_ok if a_ok else None, bad, main]); X = sv.max_row
    for j in (3, 4, 5): sv.cell(X, j).number_format = M
    sv.cell(X, 6).number_format = "0%"
    LET[s] = (f"Добрый день!\n\nООО «ГИДРОПАС» готово разместить заказ по вашему счёту (закупка ПАО «Камчатскэнерго», {len(G)} позиций, сумма более {a / 1e6:.1f} млн ₽). "
              f"Сравнили с другими предложениями заводов и просим пересмотреть цены до уровня ниже — при согласии оплатим быстро одной партией:\n\n" + "\n".join(lines[:60]) +
              ("\n…полный список — во вложении (лист с вашим названием)." if len(lines) > 60 else "") +
              "\n\nТакже по каждой позиции просим указать завод-изготовитель и действующий номер реестровой записи ПП 719 (с выпиской ГИСП) — без этого товар не проходит нацрежим.\n\nС уважением,\nООО «ГИДРОПАС»")
tot_a = sum(r[2].value or 0 for r in sv.iter_rows(min_row=3)); tot_t = sum(r[3].value or 0 for r in sv.iter_rows(min_row=3)); tot_e = sum(r[4].value or 0 for r in sv.iter_rows(min_row=3))
sv.append(["ИТОГО", None, tot_a, tot_t, tot_e]); [setattr(c, "font", Font(bold=True)) for c in sv[sv.max_row]]
for j in (3, 4, 5): sv.cell(sv.max_row, j).number_format = M
for row in sv.iter_rows(min_row=2):
    for c in row:
        if c.value is not None: c.border = BOX
for i, w in enumerate([26, 10, 16, 16, 14, 10, 12, 70], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
lt = wb.create_sheet("Письма", 1); lt.append(["Поставщик", "Текст письма на торг"]); head(lt)
for s, t in LET.items(): lt.append([s, t])
lt.column_dimensions["A"].width = 24; lt.column_dimensions["B"].width = 140
tidy(wb, row_h=34)
for r in range(2, lt.max_row + 1): lt.row_dimensions[r].height = 409
wb.save("../KAMCHATKA_TORG_2910.xlsx")
for r in sv.iter_rows(min_row=3, values_only=True): print([x if not isinstance(x, float) else round(x, 2) for x in r[:7]])
