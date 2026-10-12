"""Итоговая таблица по выгрузкам клиента 12.10.2026 -> TENDERS_1210.xlsx"""
import json, os, glob, sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from subgroups import calc, be
SUB="/tmp/claude-0/-home-user-Tender-RUS-/e981d7a6-b89f-5df6-824c-c1907918e938/scratchpad/subg2"
tri=json.load(open("data/triage_1210.json"))
UN=json.load(open("data/units_1210.json")) if os.path.exists("data/units_1210.json") else {}
def fmt(x): return f"{x:,.2f}".replace(",", " ").replace(".00","")
def unitinfo(num,nmck,positions=None):
    ps=positions if positions is not None else UN.get(num)
    if not ps: return "нет данных (коммерческая/малая закупка или позиции не опубликованы)",None
    parts=[]; tq=0; units=set(); allp=True
    for p in ps:
        q=p.get("qty"); pr=p.get("price") if "price" in p else p.get("unit_nmck")
        if pr is None: allp=False
        if q: tq+=q; units.add(p.get("unit",""))
        parts.append(f"{(p.get('name') or '')[:45]}: {fmt(q) if q else '?'} {p.get('unit','')} × {fmt(pr) if pr else '?'}")
    avg=(nmck/tq) if tq and len(units)<=1 else None
    txt="; ".join(parts[:8])+(f"; … ещё {len(parts)-8}" if len(parts)>8 else "")
    if not allp and avg: txt+=" (цена за ед. = НМЦК / кол-во)"
    if avg is None and allp and ps and len(ps)>1: avg="разные позиции — см. колонку слева"
    elif avg is None: avg="разные ед. изм." if tq else None
    return txt,avg
rows=[]
def P(nm,c,d):
    p=nm*(1-d); pre=p-c-p*0.008-c*0.2*60/365; return pre-max(pre,0)*0.15
# соль
for n,cust,t,nm,end,note in [("32616451106","ООО «Департамент городского хозяйства», МО",400,6724000,"16.10.2026 10:00","223-ФЗ; обеспечение исполнения 10%"),
    ("32616436306","ООО «ИКС Орехово-Зуево», МО",948,13869240,"13.10.2026 10:00","223-ФЗ; «Концентрат минеральный Галит» — сверить ТЗ на эквивалент"),
    ("32616447519","АО «Волоколамское ПТП РЖКХ», МО",130,1593800,"16.10.2026 10:00","223-ФЗ; «Концентрат минеральный Галит» — сверить ТЗ")]:
    c=t*2500; rows.append(["Соль",n,cust,f"Соль техническая {t} т","223-ФЗ",end,nm,f"{t} т × {fmt(nm/t)} ₽/т",nm/t,c,(nm-c)/nm,0.15,P(nm,c,0),P(nm,c,.15),"Участвовать",note+"; снижение 15% — консервативно (медиана по соли 3%)"])
# одежда по ценам клиента
xs=[x for x in tri if x['cat']=='Одежда' and x['ok'] and 100000<=(x['nmck'] or 0)<=20e6]; xs.sort(key=lambda x:-x['nmck'])
for x in xs:
    nm=x['nmck']; c=0.75*nm
    ui,avg=unitinfo(x["num"],nm)
    rows.append(["Одежда",x['num'],str(x['cust']),x['name'],x['law'],x['end'],nm,ui,avg,c,0.25,0.05,P(nm,c,.05),P(nm,c,.05),"Участвовать, если ≤2 заявок","себестоимость = 75% НМЦК (данные клиента), снижение ≤5%"])
# агентские категории
names={"med":"Медрасходники","cable":"Кабель (ЭОМ)","ss":"СС","gear":"Амуниция","plumb":"Сантехника и ЖКХ","med2a":"Медрасходники","med2b":"Медрасходники","other2":None}
REL={x["num"]:x for x in json.load(open("data/relevant_1210.json"))}
for k,cat in names.items():
    p=f"{SUB}/{k}.json"
    if not os.path.exists(p): continue
    for t in json.load(open(p,encoding="utf-8")):
        if not t.get("id") or not t.get("nmck"): continue
        cat=names[k] or REL.get(t["id"],{}).get("cat","Прочее")
        if t.get("cost")==[] and "не наш товар" in (t.get("notes") or ""):
            continue
        ui,avg=unitinfo(t["id"],t["nmck"],t.get("positions"))
        if not t.get("cost") or not t.get("covered_nmck") or t.get("spec_known") is False:
            rows.append([cat,t["id"],t.get("customer",""),t.get("subgroup",""),t.get("law",""),t.get("deadline",""),t["nmck"],ui,avg,None,None,None,None,None,"Нет цен — запросить КП",(t.get("notes") or t.get("risks") or "")[:300]]); continue
        tt=dict(t)
        if t.get("covered_nmck"): tt["nmck"]=t["covered_nmck"]
        disc=t.get("disc") or [0,.15,.3]
        P0,C,gm,n0=calc(tt,0); Pb,_,_,nb=calc(tt,disc[1]); b=be(tt)
        dec="Нет" if nb<=0 else ("Осторожно" if b<0.3 else "Участвовать")
        hyp="[ГИПОТЕЗА по себестоимости] " if any(c.get("hyp") and not c.get("alt") for c in t["cost"]) else ""
        rows.append([cat,t["id"],t.get("customer",""),t.get("subgroup",""),t.get("law",""),t.get("deadline",""),t["nmck"],ui,avg,C,gm,disc[1],n0,nb,dec,hyp+(t.get("risks") or "")[:300]])
TRI={x["num"]:x for x in tri}
def sec(v,nm):
    if v in (None,""): return "не установлено / не указано"
    v=str(v)
    if v.endswith("%"): 
        try: return f"{v} (≈{float(v[:-1].strip())/100*nm:,.0f} ₽)".replace(","," ")
        except Exception: return v
    try: return f"{float(v):,.0f} ₽ ({float(v)/nm*100:.1f}%)".replace(","," ") if nm else v
    except Exception: return v
def adv(v):
    if v in (None,"","RUB"): return "нет / не указан"
    return str(v).replace("Есть. См. документацию","есть, размер в документации")
rows2=[]
for r in rows:
    t=TRI.get(str(r[1]),{}); nm=r[6] or 0
    rows2.append(r[:15]+[adv(t.get("advance")),sec(t.get("sec_app"),nm),sec(t.get("sec_contr"),nm),t.get("smp") or "",t.get("etp") or ""]+[r[15]])
rows=rows2
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Тендеры 12.10"
H=["Категория","№ закупки","Заказчик","Предмет","Закон","Подача","НМЦК, ₽","Позиции: кол-во × цена за ед. в НМЦК, ₽","Средняя цена за ед. в НМЦК, ₽","Себестоимость, ₽","Валовая маржа при НМЦК","Ожид. снижение","Прибыль при НМЦК, ₽","Прибыль при ожид. снижении, ₽","Решение","Аванс","Обеспечение заявки","Обеспечение исполнения контракта","Только СМП","Площадка (ЭТП)","Комментарий"]
ws.append(H)
for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True,vertical="center")
OVR={"0373200024426001489":("Нет","ТЗ под поршневую систему S-Monovette (Sarstedt), импорт пробирок запрещён — заявку с российским аналогом, скорее всего, отклонят. ")}
for r in rows:
    if str(r[1]) in OVR: r[14]=OVR[str(r[1])][0]; r[15]=OVR[str(r[1])][1]+str(r[15] or "")
NORM={"Амуниция":"Амуниция и снаряжение","ЭОМ (кабель)":"Кабель (ЭОМ)","Соль и сыпучие":"Соль"}
for r in rows: r[0]=NORM.get(r[0],r[0])
_d={}
for r in rows: _d[str(r[1])]=r
rows=list(_d.values())
order={"Участвовать":0,"Участвовать, если ≤2 заявок":1,"Осторожно":2,"Нет цен — запросить КП":3,"Нет":4}
rows.sort(key=lambda r:(order.get(r[14],9),-(r[13] or 0)))
for r in rows: ws.append(r)
for row in ws.iter_rows(min_row=2):
    for i in (6,8,9,12,13): row[i].number_format='# ##0'
    for i in (10,11): row[i].number_format='0%'
    row[7].alignment=Alignment(wrap_text=True,vertical="top")
    col={"Участвовать":"C6EFCE","Участвовать, если ≤2 заявок":"FFEB9C","Осторожно":"FFEB9C","Нет цен — запросить КП":"D9D9D9","Нет":"FFC7CE"}.get(row[14].value)
    if col: row[14].fill=PatternFill("solid",fgColor=col)
for i,w in enumerate([16,22,40,45,10,18,14,70,14,14,10,10,14,16,20,22,22,26,14,20,60],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="C2"; ws.auto_filter.ref=ws.dimensions
# итоги
s=wb.create_sheet("Итог")
s.append(["Категория","Лотов","«Участвовать»","Прибыль по «Участвовать», ₽"])
cats={}
for r in rows:
    c=cats.setdefault(r[0],[0,0,0]); c[0]+=1
    if r[14].startswith("Участвовать"): c[1]+=1; c[2]+=r[13] or 0
for k,v in cats.items(): s.append([k]+v)
for c in s[1]: c.font=Font(bold=True)
s.column_dimensions["A"].width=24; s.column_dimensions["D"].width=26
for row in s.iter_rows(min_row=2): row[3].number_format='# ##0'
c2=wb.create_sheet("Коммерческие без НМЦК")
c2.append(["Категория","№","Предмет","Заказчик","Регион","Подача","ЭТП","Что делать"])
for c in c2[1]: c.font=Font(bold=True)
for x in sorted([x for x in REL.values() if x["kind"]!="с НМЦК"],key=lambda x:(x["cat"],x["end"])):
    c2.append([x["cat"],x["num"],x["name"],str(x["cust"]),x["region"],x["end"],x["etp"],"Цена не опубликована: запросить ТЗ на площадке и подать КП"])
for i,w in enumerate([20,22,70,40,22,18,22,40],1): c2.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
c2.freeze_panes="C2"; c2.auto_filter.ref=c2.dimensions
top=[r for r in rows if r[14] and r[14].startswith(("Участвовать","Осторожно")) and r[13]]
top.sort(key=lambda r:-r[13]); top=top[:30]
t=wb.create_sheet("Топ-30",0)
cols=[0,1,2,3,5,6,8,13,14,15,16,17,18,20]
t.append(["#"]+[H[i] for i in cols])
for c in t[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True)
for i,r in enumerate(top,1): t.append([i]+[r[j] for j in cols])
for row in t.iter_rows(min_row=2):
    for k in (6,7,8):
        if isinstance(row[k].value,(int,float)): row[k].number_format='# ##0'
for i,w in enumerate([4,16,22,40,40,18,14,14,16,22,22,22,26,14,60],1): t.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
t.freeze_panes="D2"
wb.save("../TENDERS_1210.xlsx"); print(len(rows), {k:v for k,v in cats.items()})
