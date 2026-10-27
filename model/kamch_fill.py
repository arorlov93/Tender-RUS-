"""Камчатка 32616346943: каждой позиции (315) — предполагаемый поставщик; по каждому поставщику отдельный xlsx для заполнения
(марка, производитель, страна, № реестра ПП 719, цена, срок). Без наших цен. → ../KAMCHATKA_ZAPOLNENIE/*.xlsx + сводка"""
import json, os, re, runpy, io, contextlib, shutil, openpyxl
from collections import defaultdict
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

with contextlib.redirect_stdout(io.StringIO()):
    rows = runpy.run_path("kamch_status.py")["rows"]; C = runpy.run_path("kamch_open.py")["C"]
    KR = {o["tt"]: o for o in runpy.run_path("kamch_registry.py")["out"]}
REGF = json.load(open("data/kamch/registry_found.json")) if os.path.exists("data/kamch/registry_found.json") else []
TT = {f'{t["table"]}/{t["n_table"]}': t for t in json.load(open("data/kamch/tz_verbatim.json"))}
SUPJ = {s["company"]: s for g in json.load(open("data/kamch/suppliers.json")) for s in g["suppliers"]}
ph = lambda s: str(s or "").split(";")[0].split("\n")[0].strip()
# контакты текущих поставщиков (из их КП/писем)
CUR = {"ДН.РУ": ("ООО «ДН.РУ» (ИНН 7724452117)", "менеджер по счетам 819892/819896", ""), "ТД АДЛ (Гранвэл)": ("ООО «Торговый Дом АДЛ»", "+7 (495) 937-89-68", "info@adl.ru"),
       "КИПИКА (прайс 01.10)": ("ООО «Пензапромарматура» (КИПИКА), Хрипунова Светлана", "+7 (8412) 909-200", "b18bk@11b18bk.ru"),
       "Лидер-Энерго": ("ООО «Лидер-Энерго» / «Транс-Энерго»", "+7 (8482) 21-66-32", "trans-energo@bk.ru"), "Тулаэлектропривод": ("АО «Тулаэлектропривод»", "по КП 10/160", "")}
NAMES = {"СИБЗТА": ("ООО «Сибирский завод трубопроводной арматуры» (СИБЗТА)", "8-800-350-91-13", "sales@sibzta.ru")}


def supj(name):
    s = next((v for k, v in SUPJ.items() if k.startswith(name[:12])), None)
    return (s["company"], ph(s.get("phone")), ph(s.get("email"))) if s else (name, "", "")


A = defaultdict(list); ROLE = {}
# 1) котированные и годные — у текущего поставщика
for d in rows:
    if d["sup"] and d["act"] != "Заказать у других":
        A[CUR[d["sup"]]].append((d, "КП получено — подтвердить и дополнить реестр"))
# 2) «Заказать у других» — первый завод из колонки «Кому ещё»
for d in rows:
    if d["act"] == "Заказать у других":
        if d["sup"] == "Тулаэлектропривод": A[CUR["Тулаэлектропривод"]].append((d, "торг: наш бюджет")); continue
        first = d["alt"].split("; ")[0].split(" +")[0].split(" 8-")[0].split(" 8 ")[0].strip()
        key = NAMES.get(first) or supj(first)
        A[key].append((d, f"у ДН.РУ: {d['why']}"))
# 3) не закрытые — жадное покрытие проверенными заводами (меньше поставщиков)
left = {d["tt"]: d for d in rows if d["act"] == "Не отработано" and not d["sup"]}
pend = [d for d in rows if d["act"] == "Не отработано" and d["sup"]]
for d in pend: A[CUR[d["sup"]]].append((d, d["why"]))
cs = sorted(C, key=lambda c: (c["type"] != "завод", -len(c["cov"])))
while left:
    best = max(cs, key=lambda c: (bool([t for t in c["cov"] if t in left]), c["type"] == "завод", len([t for t in c["cov"] if t in left])))
    got = [t for t in best["cov"] if t in left]
    if not got: break
    for t in got: A[(best["company"], best["phone"], best["email"])].append((left.pop(t), "новый запрос КП"))
for t, d in left.items(): A[supj("ООО «Ридан»")].append((d, "проверенного поставщика нет — Ридан (оригинал NVD)"))

F = ["Наименование", "ГОСТ/ТУ", "Ду, мм", "Проход", "Давление, МПа", "Температура, °С", "Среда", "Материал корпуса", "Климатическое исполнение", "Присоединение",
     "Приводное устройство: наименование", "Приводное устройство: технические характеристики", "Установочное положение", "Материал штока, шара, уплотнений", "Габариты", "Примечание"]
FILL = ["Предлагаемая марка / обозначение", "Производитель (завод)", "Страна происхождения", "№ реестровой записи (ПП 719) или евразийского реестра", "Срок действия записи",
        "Соответствие ТЗ: «да» или отличия", "Цена за ед. с НДС, ₽", "Сумма с НДС, ₽", "Срок отгрузки", "Вес ед., кг / габариты", "Сертификат / декларация ТР ТС №"]
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); Y = PatternFill("solid", fgColor="FFF2CC")
OUT = "../KAMCHATKA_ZAPOLNENIE"; shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT)
M = {}
for key in list(A):  # склеить одного поставщика под разными названиями
    base = re.sub(r"\s*\(.*", "", key[0]).replace("Торговый Дом", "ТД").strip()
    if base in M and M[base] != key: A[M[base]] += A.pop(key)
    else: M[base] = key
summ = []
for k, (key, L) in enumerate(sorted(A.items(), key=lambda x: -len(x[1])), 1):
    comp, phone, email = key
    L.sort(key=lambda x: (x[0]["tt"].split("/")[0], int(x[0]["tt"].split("/")[1])))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Заполнить"
    ws.append([f"Запрос КП — ООО «ГИДРОПАС» (ИНН 9729231965) для ПАО «Камчатскэнерго», закупка 32616346943. Получатель: {comp}"]); ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Просим заполнить жёлтые графы по каждой позиции. Характеристики — дословно из технических требований заказчика; при отличии — указать в графе «Соответствие ТЗ»."])
    ws.append(["Графы марка / производитель / страна / № записи частично ПРЕДЗАПОЛНЕНЫ нами по вашему КП и открытым данным — проверьте и исправьте, если не так."])
    ws.append(["Нацрежим ПП 1875 (ограничение): обязательны производитель, страна и ДЕЙСТВУЮЩИЙ номер реестровой записи ПП 719 (для ЕАЭС — евразийский реестр). Приложите выписку из ГИСП и сертификаты. Доставка — до Петропавловска-Камчатского (укажите вес)."])
    ws.append(["№", "Табл./№ ТЗ", *F, "Кол-во", "Ед.", *FILL])
    for c in ws[5]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for j in range(len(F) + 5, len(F) + 5 + len(FILL)): ws.cell(5, j).fill = PatternFill("solid", fgColor="BF8F00")
    comp_is_plant = any(c["company"] == comp and c["type"] == "завод" for c in C)
    for i, (d, _) in enumerate(L, 1):
        t = TT[d["tt"]]
        pre = [None] * len(FILL); kr = KR.get(d["tt"], {})
        if d["sup"] and d["act"] != "Заказать у других" and kr.get("kind") in ("ok", "eaes", "done"):  # предзаполнение из КП: марка, завод, страна, номер (если найден в открытых источниках)
            pre[0] = d["offer"][:120]; pre[1] = kr["man"]; pre[2] = kr["country"]
        elif not d["sup"] and comp_is_plant: pre[1] = comp; pre[2] = "Россия"
        m = pre[1] or comp
        hit = next((r for r in REGF if r.get("reg_number") and r.get("confidence", "").startswith("номер") and str(r.get("manufacturer", ""))[:12].lower() in (m or "").lower()
                    and any(a.lower() in (d["name"] + " " + d["offer"]).lower() for a in re.findall(r"[\wА-Яа-я-]{4,}", str(r.get("applies_to", "")))[:6])), None)
        if hit: pre[3] = hit["reg_number"]; pre[4] = hit.get("valid_to")
        ws.append([i, d["tt"], *[ILLEGAL_CHARACTERS_RE.sub("", str(t.get(f) or "")) for f in F], t["qty"], t.get("unit", "шт")] + pre)
        X = ws.max_row; ws.cell(X, len(F) + 12).value = f"={openpyxl.utils.get_column_letter(len(F) + 11)}{X}*{openpyxl.utils.get_column_letter(len(F) + 3)}{X}"
        for j in range(len(F) + 5, len(F) + 5 + len(FILL)): ws.cell(X, j).fill = Y
        for j in (len(F) + 11, len(F) + 12): ws.cell(X, j).number_format = '#,##0.00" ₽"'
        for c in ws[X]: c.border = BOX
    X = ws.max_row; L2 = openpyxl.utils.get_column_letter(len(F) + 12)
    ws.append([""] * (len(F) + 11) + [f"=SUM({L2}6:{L2}{X})"]); ws.cell(ws.max_row, len(F) + 12).number_format = '#,##0.00" ₽"'; ws.cell(ws.max_row, len(F) + 11).value = "ИТОГО"
    for i, w in enumerate([4, 9, 34, 14, 7, 10, 9, 12, 14, 14, 10, 14, 14, 16, 10, 22, 12, 14, 7, 5] + [22, 24, 14, 26, 13, 24, 13, 14, 12, 14, 20], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "D6"
    tidy(wb, row_h=60)
    fn = f"{k:02d}_" + re.sub(r"[^\wА-Яа-яЁё-]+", "_", re.sub(r"[«»\"()]", "", comp.split("(")[0] if not comp.startswith("(") else comp))[:40].strip("_") + ".xlsx"
    wb.save(f"{OUT}/{fn}")
    why = defaultdict(int)
    for _, w in L: why[w.split(":")[0][:40]] += 1
    summ.append([k, comp, phone, email, len(L), sum(TT[d["tt"]]["qty"] for d, _ in L), ", ".join(d["tt"] for d, _ in L)[:400], "; ".join(f"{a} — {b}" for a, b in why.items()), fn])

wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Кому что отправить"
sv.append(["Камчатскэнерго 32616346943 — предполагаемые поставщики: каждой из 315 позиций назначен один поставщик; файл для заполнения — в папке KAMCHATKA_ZAPOLNENIE"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["№", "Поставщик", "Телефон", "E-mail", "Позиций", "Шт", "Позиции (табл./№ ТЗ)", "Основание", "Файл"])
for c in sv[2]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for r in summ:
    sv.append(r)
    for c in sv[sv.max_row]: c.border = BOX
sv.append(["", "ИТОГО", "", "", sum(r[4] for r in summ), sum(r[5] for r in summ)]); [setattr(c, "font", Font(bold=True)) for c in sv[sv.max_row]]
for i, w in enumerate([4, 40, 22, 28, 9, 8, 60, 50, 40], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=48)
wb.save(f"{OUT}/00_СВОДКА_кому_что_отправить.xlsx")
for r in summ: print(r[0], r[1][:45], r[2], r[3], r[4])
print("позиций", sum(r[4] for r in summ))
