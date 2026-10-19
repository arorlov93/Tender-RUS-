"""Камчатскэнерго 32616346943 — арматура: позиции, количества, себестоимость (если есть price_*.json), маржа к НМЦК, условия, история."""
import json, os, re, sys, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from tidy import tidy

PR = sys.argv[1] if len(sys.argv) > 1 else ""
NMCK = 101102908
Q = {(x["table"], x["n"]): x for x in json.load(open("data/kamch/kam_qty.json"))}
S = json.load(open("data/kamch/kam_spec.json"))["sheet2"]
P = {}
for g in "ABC":
    f = f"{PR}/price_{g}.json"
    if PR and os.path.exists(f):
        for x in json.load(open(f)):
            m = re.search(r"1\.\d", str(x.get("spec", "")))
            if m: P[(m.group(0), int(x["n"]))] = x
K = json.load(open("data/big/big_kamch.json"))[0]
TABLES = {"1.1": "Спец. №1 · Возобновляемая энергетика · ремонт", "1.2": "Спец. №2 · КТЭЦ · ремонт арендованного", "1.3": "Спец. №2 · КТЭЦ · ремонт",
          "1.4": "Спец. №2 · КТЭЦ · эксплуатация", "1.5": "Спец. №2 · КТЭЦ · ТО", "1.6": "Спец. №3 · Коммунальная энергетика · ремонт",
          "1.7": "Спец. №3 · Коммунальная энергетика · ТО", "1.8": "Спец. №4 · ЦЭС · ремонт"}
FREIGHT = 55  # ₽/кг до Петропавловска-Камчатского (оценка: 45–60 ₽/кг, контейнер морем через Владивосток)
TORG = 0.10   # торг к ценам сайтов/прайсов
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; PCT = "0.0%"


def head(ws, r):
    for c in ws[r]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")


wb = openpyxl.Workbook()
sv = wb.active; sv.title = "Итог"
ws = wb.create_sheet("Позиции")
ws.append(["Камчатскэнерго 32616346943 — позиции, количества, себестоимость (заполненные цены — рынок 2026 с НДС; пустые — ещё ищем, можно вписать вручную)"]); ws["A1"].font = Font(bold=True, size=12)
H = ["Таблица", "№", "Позиция", "Кол-во", "Ед.", "DN", "PN, МПа", "Материал корпуса", "Привод", "Цена за ед. (рынок), ₽", "Цена после торга −10%, ₽", "Вес ед., кг",
     "Сумма закупки, ₽", "Доставка, ₽", "Себестоимость позиции, ₽", "Поставщик", "Тип цены / подтверждение 2026", "Аналог / реестр МПТ", "Ссылка"]
ws.append(H); head(ws, 2)
first = 3
for x in S:
    m = re.search(r"Таблица\s*(1\.\d)", x["spec"]); t = m.group(1); n = x["n"]
    q = Q[(t, n)]; p = P.get((t, n), {})
    r = ws.max_row + 1
    ws.append([t, n, q["name"], q["qty"], q["unit"], x.get("dn"), x.get("pn"), str(x.get("mat") or "")[:40], str(x.get("drive") or "")[:30],
               p.get("price"), f"=IF(J{r}=\"\",\"\",J{r}*(1-{TORG}))", p.get("weight_kg"),
               f"=IF(J{r}=\"\",\"\",K{r}*D{r})", f"=IF(L{r}=\"\",\"\",L{r}*D{r}*{FREIGHT})", f"=IF(M{r}=\"\",\"\",M{r}+N(N{r}))",
               p.get("supplier"), " · ".join(str(v) for v in (p.get("price_type"), p.get("date_evidence")) if v)[:80],
               " · ".join(str(v) for v in (p.get("analog"), p.get("registry")) if v)[:80], p.get("url")])
    if p.get("url"): ws.cell(r, 19).hyperlink = p["url"]
    for j in (10, 11, 13, 14, 15): ws.cell(r, j).number_format = M
    if not p.get("price"): ws.cell(r, 10).fill = PatternFill("solid", fgColor="FFF2CC")
last = ws.max_row
ws.append(["", "", "ИТОГО", f"=SUM(D{first}:D{last})", "", "", "", "", "", "", "", "", f"=SUM(M{first}:M{last})", f"=SUM(N{first}:N{last})", f"=SUM(O{first}:O{last})"])
for j in (13, 14, 15): ws.cell(ws.max_row, j).number_format = M
for c in ws[ws.max_row]: c.font = Font(bold=True)
for row in ws.iter_rows(min_row=3):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([7, 5, 46, 7, 5, 6, 7, 18, 16, 13, 13, 8, 14, 12, 14, 22, 26, 26, 30], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "D3"; ws.auto_filter.ref = f"A2:S{last}"
TOT = f"Позиции!O{last + 1}"; CNT = f"COUNT(Позиции!J{first}:J{last})"

# Итог
sv.append(["Камчатскэнерго 32616346943 — арматура среднего и низкого давления (Лот 0200-РЕМ ПРОД-2027-КамчЭн)"]); sv["A1"].font = Font(bold=True, size=13)
sv.append(["Конкурс в эл. форме, 223-ФЗ, Росэлторг (секция РусГидро). Подача до 29.10.2026. НМЦК с НДС. Наша цена = цена договора (УСН)."])
sv.append(["Показатель", "Значение", "Комментарий"]); head(sv, 3)
rows = [
    ("НМЦК, ₽", NMCK, "весь лот, 315 позиций, 3 383 ед., 4 грузополучателя"),
    ("Позиций с ценой", f"={CNT}", "из 315; пустые — ищем, можно вписать вручную на листе «Позиции»"),
    ("Себестоимость по найденным позициям (после торга −10%, с доставкой), ₽", f"={TOT}", "сумма столбца O листа «Позиции»"),
    ("Оценка себестоимости всего лота, ₽", f"=IF({CNT}=0,{K['cost_total'] + K['logistics']},B6/MAX(B5,1)*315)",
     "пока цен нет — оценка по истории лотов (60% НМЦК + 5 млн доставка); по мере заполнения — экстраполяция по найденным"),
    ("Маржа при НМЦК", "=1-B7/B4", ""),
    ("Ожидаемое снижение на конкурсе", 0.15, "история: −7% … −23% (2025–2026), большинство заявок отклоняют"),
    ("Цена заявки, ₽", "=B4*(1-B9)", ""),
    ("Прибыль до налога, ₽", "=B10-B7-0.008*B10-B7*0.2*120/365", "0,8% ЭП/банк, деньги 20% годовых на 120 дней"),
    ("Прибыль после УСН 15%, ₽", "=B11-0.15*MAX(B11,0)", ""),
    ("Маржа к цене заявки", "=B12/B10", ""),
]
for r in rows: sv.append(list(r))
for rr in (4, 6, 7, 10, 11, 12): sv.cell(rr, 2).number_format = M
for rr in (8, 9, 13): sv.cell(rr, 2).number_format = PCT
sv.append([])
sv.append(["Сценарии снижения", "Прибыль после УСН, ₽", ""]); head(sv, sv.max_row)
for d in (0, 0.07, 0.15, 0.23, 0.30):
    r = sv.max_row + 1
    sv.append([f"−{int(d * 100)}%", f"=(B4*(1-{d})-B7-0.008*B4*(1-{d})-B7*0.2*120/365)*IF(B4*(1-{d})-B7-0.008*B4*(1-{d})-B7*0.2*120/365>0,0.85,1)", ""])
    sv.cell(r, 2).number_format = M
sv.append([])
sv.append(["Условия", "", ""]); head(sv, sv.max_row)
for a, b in [("Поставка", "в течение 120 календарных дней с подписания договора; Петропавловск-Камчатский (ул. Пограничная 75А, ул. Степная 50, пр. Победы 24) и Елизовский р-н, пос. Светлый"),
             ("Обеспечение заявки", "1 005 499,42 ₽ (БГ или деньги)"), ("Аванс", "предусмотрен; обеспечение возврата аванса — 100% суммы аванса (БГ); размер — в проекте договора"),
             ("Гарантия", "по заводу, не менее 12 мес.; изготовление — год поставки или предыдущий"),
             ("Нацрежим", "ПП 1875 — ограничение + преимущество: по каждой позиции номер записи в реестре Минпромторга/ЕАЭС"),
             ("Документы", "сертификаты/декларации ТР ТС 032/2013 и 010/2011 со схемами 1с/5Д и протоколами испытаний; техпредложение по Прил. 1"),
             ("Аккредитация", "аккредитация в РусГидро (Прил. 10)"), ("Подача на часть лота", "нельзя — лот неделимый"),
             ("Эквиваленты", "допускаются, марка носит описательный характер; нужно обоснование эквивалентности")]:
    sv.append([a, b, ""])
sv.append([])
sv.append(["История лотов Камчатскэнерго", "", ""]); head(sv, sv.max_row)
for h in K.get("history") or []:
    if isinstance(h, dict):
        sv.append([h.get("lot") or h.get("num"), f'НМЦК {h.get("nmck"):,.0f} ₽ → победитель {h.get("winner_price") or "—"}; снижение {h.get("disc")}'.replace(",", " "),
                   str(h.get("rejected") or h.get("contract") or "")[:150]])
    else: sv.append([str(h)[:60], "", ""])
for row in sv.iter_rows(min_row=3):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
sv.column_dimensions["A"].width = 52; sv.column_dimensions["B"].width = 70; sv.column_dimensions["C"].width = 70

# По таблицам
st = wb.create_sheet("По таблицам")
st.append(["Таблица", "Грузополучатель / статья", "Позиций", "Единиц", "Цен найдено", "Себестоимость найденных, ₽"]); head(st, 1)
for t, nm in TABLES.items():
    r = st.max_row + 1
    st.append([t, nm, f'=COUNTIF(Позиции!A{first}:A{last},"{t}")', f'=SUMIF(Позиции!A{first}:A{last},"{t}",Позиции!D{first}:D{last})',
               f'=COUNTIFS(Позиции!A{first}:A{last},"{t}",Позиции!J{first}:J{last},">0")', f'=SUMIF(Позиции!A{first}:A{last},"{t}",Позиции!O{first}:O{last})'])
    st.cell(r, 6).number_format = M
for row in st.iter_rows():
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([9, 48, 10, 10, 12, 22], 1): st.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

tidy(wb, row_h=34)
for r in range(4, 14): sv.row_dimensions[r].height = 30
wb.save("../KAMCHATKA_32616346943.xlsx")
print("позиций", len(S), "с ценой", len(P))
