"""Маржа: МУП «Уфимские инженерные сети» 32616450774 — отводы в ППУ-ПЭ (ЗК 223-ФЗ, подача 20.10.2026 12:00 Уфа)."""
import json, statistics, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from tidy import tidy

P = json.load(open("data/ppu_prices_1910.json"))
MAIN = ("Сайверхот", "ПрофТехноПласт", "ТЕПЛО", "АЛЬФАВОДОСТРОЙ")
ITEMS = [  # диаметр, наименование, кол-во, цена НМЦК за шт (из ТЗ), вес шт, кг (оценка; 530 — 333 кг по awstroy)
    ("108", "Отвод 108х4,5-90-1-ППУ-ПЭ", 3, 10370.00, 30),
    ("377", "Отвод 377х8-90-1-ППУ-ПЭ", 2, 111267.66, 190),
    ("426", "Отвод 426х8-90-1-ППУ-ПЭ", 3, 80063.72, 230),
    ("530", "Отвод 530х9-90-1-ППУ-ПЭ", 1, 146012.04, 333),
]
NMCK = 639848.52
TORG = 0.10       # цены с сайтов = прайс/розница, закладываем торг 10%
FREIGHT = 50000   # доставка ~1,5 т до Уфы с разгрузкой (из Екатеринбурга ~35–45 тыс., из Ярославля/Ставрополя 60–90 тыс.)


def src(it):
    return [x for x in P if x["item"] == it and any(m in x["company"] for m in MAIN)]


def model(P_, C, L):
    # аванс 50% → финансируем половину закупки; остаток у банкрота может задержаться — 60 дней
    pre = P_ - C - L - 0.008 * P_ - (C + L) * 0.5 * 0.20 * 60 / 365
    return pre - 0.15 * max(pre, 0)


rows = []
for it, name, q, pn, w in ITEMS:
    s = src(it); pr = sorted(x["price"] for x in s)
    rows.append(dict(it=it, name=name, q=q, pn=pn, w=w, mn=pr[0], med=statistics.median(pr), mx=pr[-1],
                     best=min(s, key=lambda x: x["price"]), near=next(x for x in s if "Сайверхот" in x["company"]), s=s))
C_med = sum(r["med"] * r["q"] for r in rows); C_min = sum(r["mn"] * r["q"] for r in rows)
C_sav = sum(r["near"]["price"] * r["q"] for r in rows)
W = sum(r["w"] * r["q"] for r in rows)

thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; PCT = "0.0%"
H = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row]]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Позиции и себес"
ws.append(["32616450774 — МУП «Уфимские инженерные сети»: отводы в ППУ-ПЭ. Запрос котировок 223-ФЗ (не аукцион), подача до 20.10.2026 12:00 по Уфе (10:00 МСК)"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["НМЦК 639 848,52 ₽ с НДС, включая доставку и разгрузку в Уфе. Цены поставщиков — с сайтов 2026 (прайс/розница, с НДС), торг 10% заложен отдельно."])
ws.append(["№", "Позиция", "Кол-во, шт", "НМЦК за шт", "НМЦК сумма", "Себес за шт — мин", "Себес за шт — медиана", "Себес за шт — макс", "Себес за шт после торга −10%", "Сумма себеса (медиана)", "Маржа по позиции (медиана)", "Самый дешёвый", "Ближайший к Уфе (Екатеринбург)"]); H(ws)
for i, r in enumerate(rows, 1):
    ws.append([i, r["name"], r["q"], r["pn"], r["pn"] * r["q"], r["mn"], r["med"], r["mx"], r["med"] * (1 - TORG), r["med"] * r["q"],
               1 - r["med"] / r["pn"], f'{r["best"]["company"]} — {r["best"]["price"]:,.0f} ₽'.replace(",", " "), f'Сайверхот — {r["near"]["price"]:,.0f} ₽'.replace(",", " ")])
ws.append(["", "ИТОГО", sum(r["q"] for r in rows), None, NMCK, C_min, C_med, sum(r["mx"] * r["q"] for r in rows), C_med * (1 - TORG), C_med, 1 - C_med / NMCK, "", ""])
for c in ws[ws.max_row]: c.font = Font(bold=True)
for row in ws.iter_rows(min_row=4):
    for j, c in enumerate(row, 1):
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        if j in (4, 5, 6, 7, 8, 9, 10): c.number_format = M
        if j == 11: c.number_format = PCT
ws.cell(ws.max_row, 6).value = C_min; ws.cell(ws.max_row, 8).value = sum(r["mx"] * r["q"] for r in rows)
ws.append([]); ws.append(["Ⓘ НМЦК по 377 (111 268 ₽) дороже, чем 426 (80 064 ₽) — завышена у заказчика, рынок 53–63 тыс. На этой позиции основная маржа."])
ws.append([f"Ⓘ Вес партии ≈ {W/1000:.1f} т (оценка), доставка до Уфы с разгрузкой ≈ {FREIGHT:,.0f} ₽ — отдельная машина.".replace(",", " ")])
NOTE_ROWS = (ws.max_row - 1, ws.max_row)
for i, w in enumerate([4, 30, 9, 13, 14, 13, 14, 13, 15, 15, 12, 34, 26], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

s2 = wb.create_sheet("Маржа по сценариям")
s2.append(["Прибыль после УСН 15% при разной цене заявки (НМЦК 639 848,52 ₽)"]); s2["A1"].font = Font(bold=True, size=13)
s2.append(["Модель: P − себес − доставка − 0,8% (ЭП/банк) − стоимость денег 20% годовых на 50% закупки × 60 дн (аванс 50%), налог 15% с прибыли."])
s2.append(["Снижение от НМЦК", "Цена заявки", "Себес: медиана сайтов", "Прибыль, медиана", "Маржа, медиана", "Себес: после торга −10%", "Прибыль, после торга", "Маржа, после торга", "Себес: Сайверхот (Екб)", "Прибыль, Сайверхот"]); H(s2)
for d in (0, 0.05, 0.10, 0.15, 0.20, 0.25, 0.30):
    Pp = NMCK * (1 - d); a = model(Pp, C_med, FREIGHT); b = model(Pp, C_med * (1 - TORG), FREIGHT); c = model(Pp, C_sav, FREIGHT * 0.8)
    s2.append([d, Pp, C_med, a, a / Pp, C_med * (1 - TORG), b, b / Pp, C_sav, c])


def be(C, L):
    lo, hi = 0, 0.9
    for _ in range(60):
        m = (lo + hi) / 2
        lo, hi = (m, hi) if model(NMCK * (1 - m), C, L) > 0 else (lo, m)
    return lo
s2.append(["Безубыточное снижение", None, None, None, be(C_med, FREIGHT), None, None, be(C_med * (1 - TORG), FREIGHT), None, be(C_sav, FREIGHT * 0.8)])
for c in s2[s2.max_row]: c.font = Font(bold=True)
for row in s2.iter_rows(min_row=4):
    for j, c in enumerate(row, 1):
        c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
        c.number_format = PCT if j in (1, 5, 8) or (row[0].row == s2.max_row and c.value is not None and j > 1) else M
for i, w in enumerate([20, 15, 16, 16, 12, 18, 16, 12, 18, 16], 1): s2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

s3 = wb.create_sheet("Условия и риски")
s3.append(["Условие", "Что в документации", "Что значит для нас"]); H(s3)
for r in [
    ["Способ", "Запрос котировок в эл. форме, 223-ФЗ, ЭТП «РЭСТ» (etp.r-est.ru); не только для МСП", "Без торгов: одна цена, побеждает минимальная. Переторжки нет."],
    ["Подача", "до 12:00 по местному времени Уфы 20.10.2026 (10:00 МСК)", "Подать до утра; форма заявки — Прил. 2, цена × кол-во без округлений"],
    ["Обеспечение заявки / договора", "не требуется", "Деньги не замораживаем"],
    ["Аванс", "ПРЕДОПЛАТА 50%", "Половину закупки финансирует заказчик"],
    ["Оплата остатка", "7 раб. дней после подписания документа о приёмке; приёмка — до 20 раб. дней", "Фактически до ~1 мес. после поставки"],
    ["⚠ Заказчик — банкрот", "МУП УИС признан банкротом 02.09.2022 (А07-13106/2018), конкурсное производство. Наши требования — текущие платежи, гасятся в календарной очерёдности (ст. 134 127-ФЗ)", "Риск задержки второй половины оплаты. Аванс 50% сильно снижает риск: отгружать только после аванса."],
    ["Срок поставки", "в течение 7 дней с момента заявки заказчика, партиями", "Отводы делаются под заказ 5–15 дн. — договориться с заводом о наличии/резерве до подачи"],
    ["Доставка", "Объекты МУП УИС в Уфе, разгрузка силами поставщика, будни 08:00–16:45", "Доставка и разгрузка — в нашей цене (≈50 тыс.)"],
    ["Требования к товару", "ГОСТ 30732-2020, отводы ГОСТ 17375/17380 (не спиральношовные), Ст20/09Г2С/17Г1С; 3–4 медных проводника ОДК; концы 150/210 мм; изготовление не ранее 2025 г.; сертификаты ТР ТС 032/2013 и др.", "Запросить у завода паспорт и комплект сертификатов"],
    ["Нестыковка по 108", "В ТЗ 108х4,5 «тип 1», но оболочка 200 = тип 2 по ГОСТ 30732-2020 (тип 1 = 180)", "Предлагать как в ТЗ (тип 1); цену брали по 200 — запас"],
    ["Гарантия", "не менее 12 мес. по договору; срок службы не менее 10 лет", "Стандарт заводов"],
    ["Нацрежим", "ПП 1875 — национальный режим", "Только российский товар (заводы РФ — ок)"],
    ["Контакты", "Данилова М. А., zakupki@mup-uis.ru, +7 917 760-77-50; снабжение +7 917 760-77-90, -85", ""],
]: s3.append(r)
for row in s3.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([24, 70, 60], 1): s3.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

s4 = wb.create_sheet("Поставщики и ссылки")
s4.append(["Диаметр", "Компания", "Цена за шт, ₽", "НДС", "Подтверждение 2026", "Телефон", "E-mail", "Примечание", "Ссылка"]); H(s4)
for r in rows:
    for x in sorted(r["s"], key=lambda x: x["price"]):
        s4.append([r["it"], x["company"], x["price"], x.get("vat"), x.get("date_evidence"), x.get("phone"), x.get("email"), x.get("note"), x["url"]])
        s4.cell(s4.max_row, 9).hyperlink = x["url"]; s4.cell(s4.max_row, 3).number_format = M
for row in s4.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([8, 28, 12, 10, 34, 18, 22, 40, 50], 1): s4.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=40)
for r in NOTE_ROWS:
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
    ws.cell(r, 1).alignment = Alignment(wrap_text=False, vertical="center"); ws.row_dimensions[r].height = 20
wb.save("../PPU_UFA_32616450774.xlsx")
print("себес мин/мед/Сайверхот", round(C_min), round(C_med), round(C_sav), "вес", W)
for d in (0, 0.10, 0.20):
    Pp = NMCK * (1 - d); print(d, round(Pp), round(model(Pp, C_med, FREIGHT)), round(model(Pp, C_med * .9, FREIGHT)))
print("безубыток", round(be(C_med, FREIGHT), 3), round(be(C_med * .9, FREIGHT), 3))
