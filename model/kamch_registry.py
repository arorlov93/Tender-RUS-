"""Камчатка 32616346943: таблица по всему лоту (315 позиций) для формы КП / структуры НМЦ — производитель, страна происхождения,
№ реестровой записи (ПП 719 / евразийский реестр). ТТ п. 1.5: режим ОГРАНИЧЕНИЯ по ПП 1875, номер записи обязателен и должен быть действующим.
Номера реестра заполняются по ответам поставщиков (data/kamch/registry.json: {"1.3/89": {"num": "...", "valid_to": "...", "src": "..."}}). → KAMCHATKA_REESTR_MPT.xlsx"""
import json, os, re, runpy, io, contextlib, openpyxl
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

with contextlib.redirect_stdout(io.StringIO()):
    rows = runpy.run_path("kamch_status.py")["rows"]; C = runpy.run_path("kamch_open.py")["C"]
REG = json.load(open("data/kamch/registry.json")) if os.path.exists("data/kamch/registry.json") else {}
# бренд в предложении → (производитель, страна, вывод)
BR = [(r"СибЗТА", "ООО «Сибирский завод трубопроводной арматуры» (Новосибирск)", "Россия", "ok"),
      (r"Авангард", "АО «Авангард» / САЗ (Старый Оскол)", "Россия", "ok"), (r"Ридан", "ООО «Ридан»", "Россия", "ok"),
      (r"ЛМЗ", "ОАО «Литейно-механический завод»", "Россия", "ok"), (r"ЧАЗ", "Чебоксарский арматурный завод", "Россия", "ok"),
      (r"Гранвэл|ГРАНВЭЛ|Гранлок", "ООО «ТД АДЛ» (Гранвэл/Гранлок, производство Лыткарино)", "Россия", "ok"),
      (r"Тулаэлектропривод|Н-В", "АО «Тулаэлектропривод»", "Россия", "ok"), (r"11б18бк|11б27пм", "ООО «Пензапромарматура» (КИПИКА)", "Россия", "ok"),
      (r"Цветлит", "УП «Цветлит»", "Беларусь (ЕАЭС)", "eaes"),
      (r"Китай|BasicStandard|Benarmo|Мегастрой", "—", "КНР", "bad"),
      (r"Dendor|SantechSystems|Луидор|ТехГидроМаш|DN\.ru", None, "уточнить", "ask")]
ASKMSG = {"ДН.РУ": "ДН.РУ (менеджеру, выставившему счета 819892/819896) — запросить завод-изготовитель, страну и № записи по каждой строке счёта"}
out = []
for d in sorted(rows, key=lambda d: (d["tt"].split("/")[0], int(d["tt"].split("/")[1]))):
    o = dict(tt=d["tt"], name=d["name"], q=d["q"], sup=d["sup"], offer=d["offer"], act=d["act"])
    b = next(((m, c, k) for rx, m, c, k in BR if re.search(rx, d["offer"] + " " + d["sup"])), None)
    if d["sup"] == "Лидер-Энерго": b = ("уточнить у Лидер-Энерго (ЗДП по ТУ 28.14.13-012-09338969-2019)", "Россия (уточнить)", "ask")
    if not d["sup"]:
        cand = [c for c in C if d["tt"] in c["cov"]][:3]
        o.update(man="не выбран", country="", kind="open", who="; ".join(f'{c["company"].split("(")[0].strip()} {c["phone"]} {c["email"]}' for c in cand) or "поставщик не найден")
    else:
        m, c, k = b or (None, "уточнить", "ask")
        brand = re.search(r"DN\.ru|Dendor|SantechSystems|Луидор|ТехГидроМаш", d["offer"])
        o.update(man=m or (f"{brand.group(0)} — завод-изготовитель уточнить" if brand else "уточнить у поставщика"), country=c, kind=k, who=ASKMSG.get(d["sup"], d["sup"]))
    r = REG.get(d["tt"], {}); o.update(num=r.get("num", ""), valid=r.get("valid_to", ""), src=r.get("src", ""))
    if o["num"]: o.update(kind="done", man=r["man"], country="Россия", who=f'{r.get("note", "")}; источник {r["src"]} — сверить с ГИСП')
    elif r.get("note"): o["who"] = (o["who"] + " | " + r["note"]).strip(" |")
    out.append(o)
ST = {"done": ("Номер есть", "C6EFCE"), "ok": ("Российский производитель — запросить номер", "E2EFDA"), "eaes": ("ЕАЭС — нужен № евразийского реестра", "FFF2CC"),
      "ask": ("Страна/завод неизвестны — уточнить", "FFEB9C"), "bad": ("Иностранный — заменить на российский", "F8CBAD"), "open": ("Поставщик не выбран", "DDEBF7")}
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
head = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row] if c.value is not None]
wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Итог"
sv.append(["Камчатскэнерго 32616346943 — номера реестра Минпромторга по всему лоту (315 позиций)"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["ТТ п. 1.5: режим ОГРАНИЧЕНИЯ (ПП 1875). В форме КП и структуре НМЦ по каждой позиции — производитель, страна, № реестровой записи (ПП 719; для ЕАЭС — евразийский реестр), запись действующая."])
sv.append(["Без номера товар считается иностранным: при наличии заявки с российским товаром — отклонение. Номера даёт завод (выписка из ГИСП) — сверить на gisp.gov.ru перед подачей."])
sv.append(["Статус", "Позиций", "Что делать"]); head(sv)
DO = {"done": "сверить срок действия в ГИСП", "ok": "запросить у поставщика/завода № записи и выписку ГИСП", "eaes": "запросить № евразийского реестра",
      "ask": "запросить у поставщика завод и страну; если КНР — заменить", "bad": "заменить на российский аналог (позиции из счёта ДН.РУ)", "open": "выбрать завод (KAMCHATKA_OPEN_ZAPROS_2710) — номер запросить вместе с КП"}
for k, (t, col) in ST.items():
    sv.append([t, sum(o["kind"] == k for o in out), DO[k]]); sv.cell(sv.max_row, 1).fill = PatternFill("solid", fgColor=col)
sv.append(["ИТОГО", len(out), ""]); sv.cell(sv.max_row, 1).font = Font(bold=True)
for row in sv.iter_rows(min_row=4):
    for c in row:
        if c.value is not None: c.border = BOX
for i, w in enumerate([46, 10, 80], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws = wb.create_sheet("Все позиции (315)")
ws.append(["№", "Табл./№ ТЗ", "Наименование по ТЗ", "Кол-во", "Предлагаемый товар (марка)", "Производитель", "Страна происхождения", "№ реестровой записи (ПП 719 / ЕАЭС)",
           "Действует до", "Статус", "Где взять номер / кому писать", "Поставщик"]); head(ws)
for i, o in enumerate(out, 1):
    ws.append([i, o["tt"], o["name"], o["q"], o["offer"][:150], o["man"], o["country"], o["num"], o["valid"], ST[o["kind"]][0], o["who"][:300], o["sup"]])
    X = ws.max_row; ws.cell(X, 10).fill = PatternFill("solid", fgColor=ST[o["kind"]][1]); ws.cell(X, 8).fill = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[X]: c.border = BOX
for i, w in enumerate([4, 9, 40, 7, 44, 32, 14, 22, 11, 26, 44, 16], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:L{ws.max_row}"

rq = wb.create_sheet("Запрос номеров")
rq.append(["Кому", "Позиций", "Какие позиции (табл./№ ТЗ)", "Текст запроса"]); head(rq)
G = defaultdict(list)
for o in out:
    if o["kind"] in ("ok", "eaes", "ask", "bad"): G[o["sup"]].append(o)
for s, L in sorted(G.items(), key=lambda x: -len(x[1])):
    txt = ("Добрый день! Для заявки в ПАО «Камчатскэнерго» (32616346943, нацрежим ПП 1875 — ограничение) просим по каждой позиции вашего счёта/КП указать: "
           "завод-изготовитель, страну происхождения, номер реестровой записи из реестра российской промышленной продукции (ПП 719) или евразийского реестра, "
           "срок действия записи, и приложить выписку из ГИСП. Позиции иностранного производства без записи в реестре — прошу заменить на российские аналоги.")
    rq.append([s, len(L), ", ".join(o["tt"] for o in L), txt])
    for c in rq[rq.max_row]: c.border = BOX
for i, w in enumerate([30, 9, 60, 90], 1): rq.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=44)
wb.save("../KAMCHATKA_REESTR_MPT.xlsx")
for k in ST: print(ST[k][0], sum(o["kind"] == k for o in out))
print({s: len(L) for s, L in G.items()})
