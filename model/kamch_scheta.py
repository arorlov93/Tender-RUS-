"""Камчатка 32616346943: разбор всех полученных счетов/КП попозиционно (ДН.РУ 819892/819896, АДЛ, Лидер-Энерго, КИПИКА, Тулаэлектропривод; ЭПАС — альтернатива).
Группы: 1) точно есть — цена КП не выше нашей проверенной цены и товар по ТЗ; 2) выше нашего бюджета (> +10%) или не по ТЗ; 3) дороже нашей цены на 0–10% — можно продать с меньшей наценкой (7–10%).
Ориентир цены заказчика за позицию — НМЦК лота, распределённая пропорционально нашей цене (оценка). → ../KAMCHATKA_SCHETA_2910.xlsx"""
import json, runpy, io, contextlib, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

with contextlib.redirect_stdout(io.StringIO()):
    rows = runpy.run_path("kamch_status.py")["rows"]
REG = json.load(open("data/kamch/registry.json"))
EPAS = {x["tt"]: x for x in json.load(open("data/kamch/kp/epas_2710.json"))["items"]} if True else {}
N = 101102908; K = N / sum(d["our"] * d["q"] for d in rows)  # НМЦК / наша цена лота
q = [d for d in rows if d["inv"] is not None]
def grp(d):
    if any(f.startswith("BAD") for f in d["flags"]): return 2
    return 1 if d["dev"] <= 0 else 3 if d["dev"] <= 0.10 else 2
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); M = '#,##0" ₽"'
head = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row] if c.value is not None]
TIT = {1: "1. Точно есть — цена КП не выше нашей цены, товар по ТЗ", 2: "2. Выше нашего бюджета (> +10% к нашей цене) или не по ТЗ — заказывать у других / торг",
       3: "3. Дороже нашей цены на 0–10% — можно брать и продавать с наценкой 7–10%"}
COL = {1: "C6EFCE", 2: "F8CBAD", 3: "FFEB9C"}
wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Сводка"
sv.append(["Камчатскэнерго 32616346943 — все полученные счета и КП попозиционно (на 29.10.2026)"]); sv["A1"].font = Font(bold=True, size=12)
sv.append([f"Наша цена — проверенная рыночная цена 2026 (с НДС). Ориентир заказчика за позицию = наша цена × {K:.3f} (НМЦК 101,1 млн распределена пропорционально, оценка — НМЦ по позициям заказчик не раскрыл)."])
sv.append(["Группа", "Позиций", "Сумма по КП, ₽", "Наша цена этих позиций, ₽", "Отклонение КП", "Ориентир заказчика, ₽", "Маржа к ориентиру", "Что делать"]); head(sv)
WHAT = {1: "заказывать; подтвердить № реестра ПП 719 (ДН.РУ — изготовитель!)", 2: "не брать у этого поставщика: заводы (KAMCHATKA_FINAL_ZAVODY) или торг до нашей цены", 3: "брать, если завод не даст дешевле; запас по марже есть"}
for g in (1, 2, 3):
    G = [d for d in q if grp(d) == g]; a = sum(d["inv"] * d["q"] for d in G); b = sum(d["our"] * d["q"] for d in G)
    sv.append([TIT[g], len(G), a, b, a / b - 1, b * K, 1 - a / (b * K), WHAT[g]]); X = sv.max_row
    for j in (3, 4, 6): sv.cell(X, j).number_format = M
    sv.cell(X, 5).number_format = "+0%;-0%"; sv.cell(X, 7).number_format = "0%"; sv.cell(X, 1).fill = PatternFill("solid", fgColor=COL[g])
nq = [d for d in rows if d["inv"] is None]
sv.append(["Без КП (счетов нет)", len(nq), None, sum(d["our"] * d["q"] for d in nq), None, sum(d["our"] * d["q"] for d in nq) * K, None, "запросы заводам — KAMCHATKA_FINAL_ZAVODY"])
for j in (4, 6): sv.cell(sv.max_row, j).number_format = M
sv.append([]); sv.append(["Счета по поставщикам", "Позиций", "Группа 1", "Группа 3", "Группа 2", "Сумма КП, ₽", "", "Условия"]); head(sv)
TERMS = {"ДН.РУ": "счета 819892 / 819896 (26.10), склад Москва; номера реестра не даны", "ТД АДЛ (Гранвэл)": "КП УТД000112759: 100% предоплата, самовывоз Лыткарино, 0–2 нед.",
         "Лидер-Энерго": "КП 2022 / 2035: 100% предоплата, 1–3 раб. дня, гарантия 60 мес.", "КИПИКА (прайс 01.10)": "Пензапромарматура: прайс от 51 шт, ждут артикулы",
         "Тулаэлектропривод": "КП 10/160: 100% предоплата, 75–90 дн."}
for s in sorted({d["sup"] for d in q}):
    G = [d for d in q if d["sup"] == s]
    sv.append([s, len(G), sum(grp(d) == 1 for d in G), sum(grp(d) == 3 for d in G), sum(grp(d) == 2 for d in G), sum(d["inv"] * d["q"] for d in G), "", TERMS.get(s, "")]); sv.cell(sv.max_row, 6).number_format = M
sv.append(["ЭПАС (КП 27.10, ЗДП)", len(EPAS), "", "", "", None, "", "альтернатива Лидер-Энерго — дороже (76 775 против 74 542 по 3 ЗДП), не берём"])
for row in sv.iter_rows(min_row=3):
    for c in row:
        if c.value is not None: c.border = BOX
for i, w in enumerate([62, 9, 16, 18, 11, 18, 11, 60], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
H = ["№", "Табл./№ ТЗ", "Позиция ТЗ", "Кол-во", "Поставщик", "Что в счёте", "Цена КП, ₽/шт", "Наша цена, ₽/шт", "Отклонение", "Сумма КП, ₽", "Ориентир заказчика, ₽/шт", "Маржа к ориентиру",
     "Срок", "№ реестра ПП 719", "Замечания (почему / что подтвердить)"]
for g in (1, 3, 2):
    ws = wb.create_sheet(TIT[g][:2] + {1: " Точно есть", 2: " Выше бюджета", 3: " Продать 7–10%"}[g])
    ws.append([TIT[g]]); ws["A1"].font = Font(bold=True, size=12); ws.append(H); head(ws)
    G = sorted([d for d in q if grp(d) == g], key=lambda d: (d["sup"], -(d["inv"] * d["q"])))
    for i, d in enumerate(G, 1):
        why = "; ".join(f.split(":", 1)[1] for f in d["flags"]) or d["why"]
        ws.append([i, d["tt"], d["name"], d["q"], d["sup"], d["offer"][:140], d["inv"], d["our"], d["dev"], d["inv"] * d["q"], d["our"] * K, 1 - d["inv"] / (d["our"] * K), d["srok"],
                   REG.get(d["tt"], {}).get("num") or "нет", why[:250]])
        X = ws.max_row
        for j in (7, 8, 10, 11): ws.cell(X, j).number_format = M
        ws.cell(X, 9).number_format = "+0%;-0%"; ws.cell(X, 12).number_format = "0%"; ws.cell(X, 9).fill = PatternFill("solid", fgColor=COL[g])
        for c in ws[X]: c.border = BOX
    ws.append(["", "", "ИТОГО", sum(d["q"] for d in G), "", "", None, None, None, sum(d["inv"] * d["q"] for d in G)]); ws.cell(ws.max_row, 10).number_format = M
    for c in ws[ws.max_row]: c.font = Font(bold=True)
    for i, w in enumerate([4, 8, 34, 7, 16, 44, 12, 12, 9, 13, 13, 9, 14, 12, 50], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "D3"; ws.auto_filter.ref = f"A2:O{ws.max_row - 1}"
tidy(wb, row_h=36); wb.save("../KAMCHATKA_SCHETA_2910.xlsx")
for g in (1, 3, 2):
    G = [d for d in q if grp(d) == g]; print(g, len(G), round(sum(d["inv"] * d["q"] for d in G)), round(sum(d["our"] * d["q"] for d in G)))
print("K", round(K, 3), "без КП", len(nq))
