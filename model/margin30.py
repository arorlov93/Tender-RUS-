"""Лоты из 9 выгрузок 19.10 с маржой ≥25% при НМЦК (без аукционов), со ссылками. 20.10.2026."""
import io, contextlib, json, sys, openpyxl
sys.argv = [sys.argv[0], "data/verify_tubes_gorka_2010.json"]  # проверенные цены пробирок и «Горки»
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
with contextlib.redirect_stdout(io.StringIO()):
    import go_list as g

T = {x["num"]: x for x in json.load(open("data/triage_1910.json"))}
L = json.load(open("data/go_links_2010.json"))
pct = lambda v: None if v is None else (v / 100 if v > 1 else v)
NOTE = {"0373100001026000845": "дёшево только б/у/серый Cisco; 44-ФЗ требует новый, нацрежим — ограничение; с новым маржа ~35%",
        "0373200011426001679": "нужен российский поставщик контейнеров с поршнем ~12 ₽ (импорт ограничен)",
        "0318300596226000122": "количество не задано (цены за ед.); проверить крышку с двухэтапным снятием у МиниМед",
        "0318200081326000039": "цены заказчика 39–58 ₽ при рынке 5–10 ₽ — ожидать демпинг",
        "32616465624": "17 типов затворов, цен за единицу нет — себес оценочный",
        "0345100011326000005": "ловушка ТЗ: диски ≥10 ТБ и ≥1000 МБ/с — подходит только SSD, он дороже НМЦК; камеры TRASSIR (реестр РФ)",
        "32616450774": "подача 20.10 10:00 МСК — если не подали, пропускаем; заказчик-банкрот, отгрузка после аванса 50%",
        "0348100043326000102": "лапароскопический инструмент; нацрежим — российские (Эндомедиум, ЭлМед), цены по запросу",
        "32616472923": "лампы, светильники, провод — характеристики не получены",
        "32616477173": "изолирующие соединения — нужна спецификация",
        "10762184": "малая закупка СИЗ — позиции неизвестны",
        "32616477523": "видеонаблюдение — нужна спецификация",
        "588439": "образцы и сертификаты до 21.10 13:00; количество неизвестно"}
R = []
for r in g.ROWS:
    m = 1 - r["C"] / r["nmck"]
    if m < 0.25: continue
    R.append(dict(num=r["num"], cust=r["cust"], subj=r["subj"], end=r["end"], law=r["law"], etp=r["etp"], nmck=r["nmck"], C=r["C"], m=m, d=r["d"],
                  net=g.model(r["nmck"] * (1 - r["d"]), r["C"], r.get("days", 60)), be=g.be(r["nmck"], r["C"], r.get("days", 60)),
                  conf="перепроверено: цены 2026 по позициям", note=r.get("risk") or NOTE.get(r["num"], ""), basis=r["basis"]))
for e in g.EST:
    m = 1 - e["C"] / e["nmck"]
    if m >= 0.25:
        R.append(dict(num=e["num"], cust=e["cust"], subj=e["subj"], end=e["end"], law=e["law"], etp=e["etp"], nmck=e["nmck"], C=e["C"], m=m, d=e["d"],
                      net=g.model(e["nmck"] * (1 - e["d"]), e["C"]), be=g.be(e["nmck"], e["C"]), conf="оценка — нужна спецификация", note="", basis=e["basis"], href=e.get("href")))
for f in ("scan_out_1", "scan_out_2", "scan_out_3"):
    for x in json.load(open(f"data/scan/{f}.json")):
        m = pct(x.get("margin_pct"))
        if m is None or m < 0.25 or str(x.get("decision", "")).startswith("Нет") or not x.get("nmck"): continue
        t = T.get(x["num"], {})
        C = (x.get("cost_total") or 0) + (x.get("logistics") or 0)
        R.append(dict(num=x["num"], cust=t.get("cust") or x.get("cust"), subj=t.get("name") or x.get("name"), end=(t.get("end") or "")[8:10] + "." + (t.get("end") or "")[5:7] + ".2026 " + (t.get("end") or "")[11:16],
                      law=f'{t.get("law", "")} · {t.get("method", "")}', etp=t.get("etp", ""), nmck=x["nmck"], C=C, m=m, d=x.get("exp_disc") or 0.15, net=x.get("net_expected"),
                      be=x.get("breakeven_disc"), conf={"low": "оценка (низкая)", "med": "средняя", "high": "высокая"}.get(x.get("confidence"), x.get("confidence")),
                      note=NOTE.get(x["num"], ""), basis=str(x.get("risks") or "")[:160]))
R.sort(key=lambda r: -(r["net"] or 0))
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; P = "0%"
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Маржа ≥25%"
ws.append([f"Лоты с маржой ≥25% при НМЦК — из 9 выгрузок 19.10, без аукционов, подача с 20.10. Всего {len(R)}"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Маржа = 1 − себестоимость/НМЦК. Прибыль — после УСН 15% при ожидаемом снижении. Зелёные — посчитано по позициям с проверенными ценами 2026; серые — оценка, нужна спецификация/КП поставщика."])
H = ["№", "Подача до (МСК)", "Номер", "Заказчик", "Предмет", "Способ", "Площадка", "НМЦК, ₽", "Себестоимость, ₽", "Маржа при НМЦК", "Ожид. снижение", "Прибыль после УСН, ₽",
     "В ноль при снижении", "Надёжность расчёта", "Важно", "Ссылка ЕИС", "Ссылка площадки"]
ws.append(H)
for c in ws[3]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
sh = lambda t, n: (t or "") if len(t or "") <= n else t[:n - 1] + "…"
for i, r in enumerate(R, 1):
    lk = L.get(r["num"], {})
    href = lk.get("href") or ""; etpu = lk.get("etp_url") or r.get("href") or ""
    if etpu and not etpu.startswith("http"): etpu = "https://" + etpu
    ws.append([i, r["end"], r["num"], sh(r["cust"], 60), sh(r["subj"], 90), sh(r["law"], 40), r["etp"], r["nmck"], r["C"], r["m"], r["d"], r["net"], r["be"], r["conf"],
               sh(r["note"], 150), href if "zakupki" in href else "", etpu])
    X = ws.max_row; ws.cell(X, 3).number_format = "@"
    for j in (8, 9, 12): ws.cell(X, j).number_format = M
    for j in (10, 11, 13): ws.cell(X, j).number_format = P
    for j in (16, 17):
        if ws.cell(X, j).value: ws.cell(X, j).hyperlink = ws.cell(X, j).value; ws.cell(X, j).font = Font(color="0563C1", underline="single")
    ws.cell(X, 14).fill = PatternFill("solid", fgColor="C6EFCE" if r["conf"].startswith("перепроверено") else "F2F2F2")
    for c in ws[X]: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
ws.append(["", "", "", "ИТОГО", "", "", "", sum(r["nmck"] for r in R), sum(r["C"] for r in R), None, None, sum(r["net"] or 0 for r in R)])
for j in (8, 9, 12): ws.cell(ws.max_row, j).number_format = M
for c in ws[ws.max_row]: c.font = Font(bold=True)
for i, w in enumerate([4, 16, 21, 32, 40, 24, 16, 14, 14, 9, 9, 14, 10, 18, 44, 40, 40], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "D4"
tidy(wb, row_h=46)
wb.save("../MARGIN30_2010.xlsx")
for r in R: print(r["num"], r["end"], round(r["nmck"]), round(r["m"], 2), round(r["net"] or 0), r["conf"][:12])
