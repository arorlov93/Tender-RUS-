"""Камчатка 32616346943: что брать у ДН.РУ / Тулаэлектропривода, где торговаться (наш бюджет) и что заказывать у других заводов — по счетам 819892, 819896 и КП 10/160."""
import json, re, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

TORG = 0.10
SUP = json.load(open("data/kamch/suppliers.json"))
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); M = '#,##0" ₽"'; P = "+0%;-0%;0%"
COL = {"Брать": "C6EFCE", "Брать, если": "E2EFDA", "Торг": "FFEB9C", "Заказать": "F8CBAD", "Запросить": "DDEBF7"}
SKIP = re.compile(r"не использовать|не опубликован|только через дилеров|консультация|ДН\.ру|DN\.ru|ТехГидроМаш", re.I)
KEY = [(r"30нж|нерж.*задв", "Задвижки нерж"), (r"30ч|чугун.*задв", "Задвижки чугун"), (r"задвижк|30с", "Задвижки сталь"), (r"обратн|NVD|Dendor|16ч|19[чсн]", "Клапаны обратн"),
       (r"предохран|17[лс]|ПСК|ARI-SAFE|СППК", "Клапаны предохран"), (r"регулир|регулятор|VFG|AFD|РД-НО|АРТ|25ч|25с", "Клапаны регулир"),
       (r"привод|Н-В", "Электроприводы"), (r"клапан|вентил|ВДХ|15", "Клапаны запорн")]


def others(name):
    k = next(v for rx, v in KEY if re.search(rx, name, re.I))
    g = next(g for g in SUP if g["group"].lower().startswith(k.lower()))
    return "; ".join(f'{s["company"].split("(")[0].strip()} {str(s.get("phone") or "").split(";")[0].split(chr(10))[0]}' for s in g["suppliers"] if not SKIP.search(str(s.get("checked")) + s["company"]))[:220]


def load(f, sh, c):
    ws = openpyxl.load_workbook(f)[sh]; out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] is None or not str(r[0]).isdigit(): continue
        out.append(dict(src=sh, z=r[0], tt=r[1], name=r[c["name"]], q=r[c["q"]], our=r[c["our"]], inv=r[c["inv"]], offer=str(r[c["offer"]] or ""), note=str(r[c["note"]] or "")))
    return out


R = [dict(d, grp="Задвижки", sup="ДН.РУ (сч. 819892)") for d in load("../KAMCHATKA_KP_DNRU_819892.xlsx", "Построчно", dict(name=2, q=4, our=5, inv=6, offer=10, note=13))]
R += [dict(d, grp="Клапаны", sup="ДН.РУ (сч. 819896)") for d in load("../KAMCHATKA_KP_DNRU_819896_TULA.xlsx", "Клапаны построчно", dict(name=3, q=4, our=5, inv=6, offer=10, note=12))]
R.append(dict(src="", z="—", tt="1.1 / 24", name="Электропривод общепромышленный Н-В-02 У2", q=2, our=130000, inv=203984, offer="Н-В-02 У2, ТУ 26-07-015-89; 100% предоплата, 75–90 дн.",
              note="", grp="Приводы", sup="Тулаэлектропривод (КП 10/160)"))
for d in R:
    d["budget"] = int(d["our"] * (1 - TORG) // 100 * 100)  # бюджет = наша себестоимость в расчёте (цена −10%), вниз до 100 ₽
    n = d["note"]; dev = d["inv"] / d["our"] - 1 if d["inv"] else None; d["dev"] = dev
    bad = n.startswith("ДРУГОЙ") or "давление ниже" in n or "Ду не" in n
    cond = [t for t in ("АНАЛОГ" in n and "эквивалент (обоснование)", re.search(r"КНР|импорт", n) and "номер в реестре Минпромторга", "КОФ" in n and "докупить КОФ",
                        "CF8" in n and "согласие заказчика на корпус CF8") if t]
    if dev is None: d.update(act="Запросить у других", why="ДН.РУ не котировал")
    elif bad: d.update(act="Заказать у других", why=next(t.strip() for t in n.split(";") if t.strip().startswith(("ДРУГОЙ", "Ду"))))
    elif dev > 0.30: d.update(act="Заказать у других", why=f"дороже нашей цены на {dev:+.0%}")
    elif dev > 0.10: d.update(act="Торг: наш бюджет", why=f"дороже на {dev:+.0%} — предложить бюджет, иначе у других")
    elif cond: d.update(act="Брать, если подтвердят", why=", ".join(cond))
    else: d.update(act="Брать у поставщика", why="цена в пределах нашей (±10%) и по ТЗ" if dev >= -0.1 else f"дешевле нашей на {dev:+.0%}")
    d["alt"] = "; ".join(x for x in others(d["name"]).split("; ") if d["sup"].split(" (")[0] not in x) if d["act"].startswith(("Заказать", "Запросить", "Торг")) else ""
ORDER = ["Брать у поставщика", "Брать, если подтвердят", "Торг: наш бюджет", "Заказать у других", "Запросить у других"]
R.sort(key=lambda d: (ORDER.index(d["act"]), d["grp"], -(d["our"] * d["q"])))

wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Итог"
sv.append(["Камчатскэнерго 32616346943 — что берём у ДН.РУ / Тулы, где торгуемся, что заказываем у других (по счетам на 23.10.2026)"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["Бюджет для торга = наша себестоимость в расчёте лота (проверенная цена − 10%), с НДС за шт. Говорить поставщику: «наш бюджет — X ₽, сможете?»"])
sv.append(["Решение", "Позиций", "Шт", "Наша цена × кол-во, ₽", "Цена поставщика × кол-во, ₽", "Бюджет × кол-во, ₽", "Что делать"])
for c in sv[3]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
WHAT = {"Брать у поставщика": "заказывать у ДН.РУ по счёту (перед оплатой — обновить счёт, он действует 2 дня)",
        "Брать, если подтвердят": "брать у ДН.РУ после письменного подтверждения: реестр МПТ / эквивалент / КОФ / CF8",
        "Торг: наш бюджет": "назвать бюджет из таблицы; не соглашаются — заказать у заводов из колонки «Кому ещё»",
        "Заказать у других": "у ДН.РУ не брать: дороже рынка или не по ТЗ — запрос заводам из колонки «Кому ещё», ориентир — бюджет",
        "Запросить у других": "ДН.РУ не котировал — запрос заводам, ориентир — бюджет"}
for a in ORDER:
    g = [d for d in R if d["act"] == a]
    sv.append([a, len(g), sum(d["q"] for d in g), sum(d["our"] * d["q"] for d in g), sum((d["inv"] or 0) * d["q"] for d in g) or None, sum(d["budget"] * d["q"] for d in g), WHAT[a]])
    for j in (4, 5, 6): sv.cell(sv.max_row, j).number_format = M
    sv.cell(sv.max_row, 1).fill = PatternFill("solid", fgColor=COL[next(k for k in COL if a.startswith(k))])
sv.append(["ИТОГО запрошено", len(R), sum(d["q"] for d in R), sum(d["our"] * d["q"] for d in R), None, sum(d["budget"] * d["q"] for d in R), "из 315 позиций лота; остальные 152 (краны, затворы и пр.) ещё не запрашивались"])
for j in (4, 6): sv.cell(sv.max_row, j).number_format = M
for c in sv[sv.max_row]: c.font = Font(bold=True)
for row in sv.iter_rows(min_row=3):
    for c in row:
        if c.value is not None: c.border = BOX
for i, w in enumerate([26, 9, 8, 18, 18, 18, 80], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws = wb.create_sheet("По позициям")
ws.append(["Решение", "Группа", "№ по заявке", "ТТ / №", "Позиция ТЗ", "Кол-во", "Наша цена, ₽/шт", "Цена поставщика, ₽/шт", "Отклонение", "Бюджет для торга, ₽/шт", "Бюджет × кол-во, ₽",
           "Поставщик", "Что предложено", "Почему", "Кому ещё запросить"])
for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for d in R:
    ws.append([d["act"], d["grp"], d["z"], d["tt"], str(d["name"])[:90], d["q"], d["our"], d["inv"], d["dev"], d["budget"], d["budget"] * d["q"], d["sup"], d["offer"][:120], d["why"], d["alt"]])
    X = ws.max_row
    for j in (7, 8, 10, 11): ws.cell(X, j).number_format = M
    ws.cell(X, 9).number_format = P
    ws.cell(X, 1).fill = PatternFill("solid", fgColor=COL[next(k for k in COL if d["act"].startswith(k))])
    for c in ws[X]: c.border = BOX
for i, w in enumerate([20, 10, 7, 9, 40, 7, 13, 13, 10, 13, 14, 18, 40, 40, 60], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "F2"; ws.auto_filter.ref = f"A1:O{ws.max_row}"

for a in ("Торг: наш бюджет", "Заказать у других", "Запросить у других"):
    t = wb.create_sheet({"Торг: наш бюджет": "Торг (бюджет)", "Заказать у других": "Заказать у других", "Запросить у других": "Не котировано"}[a])
    t.append(["Позиция ТЗ", "ТТ / №", "Кол-во", "Наш бюджет, ₽/шт с НДС", "Бюджет × кол-во, ₽", "Цена поставщика, ₽/шт", "Почему", "Кому запросить"])
    for c in t[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for d in [d for d in R if d["act"] == a]:
        t.append([str(d["name"])[:90], d["tt"], d["q"], d["budget"], d["budget"] * d["q"], d["inv"], d["why"], d["alt"]])
        for j in (4, 5, 6): t.cell(t.max_row, j).number_format = M
        for c in t[t.max_row]: c.border = BOX
    for i, w in enumerate([46, 9, 7, 14, 14, 14, 40, 70], 1): t.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=40)
wb.save("../KAMCHATKA_BUY_PLAN_2310.xlsx")
for row in sv.iter_rows(min_row=4, max_row=sv.max_row, values_only=True): print(row[:6])
