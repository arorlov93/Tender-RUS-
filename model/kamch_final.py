"""Камчатка 32616346943 — ФИНАЛ для заполнения: официальная форма Приложения 1 к ТТ (Лист2, колонки «Предложение участника» 1–18)
с выбранным поставщиком/заводом по каждой из 315 позиций, данные для КП (производитель, страна, № реестровой записи ПП 719 — ТТ п. 1.5),
жёлтым — что ещё заполнить. Выбор: текущий поставщик с номером записи → завод-замена ДН.РУ (dnru_assign) → завод нового запроса (assign),
только если построчная проверка ТЗ (data/kamch/tz_check.json) не дала «не соответствует». → ../KAMCHATKA_FINAL_ZAPOLNENIE.xlsx"""
import json, os, re, runpy, io, contextlib, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

with contextlib.redirect_stdout(io.StringIO()):
    rows = {d["tt"]: d for d in runpy.run_path("kamch_status.py")["rows"]}
TT = {f'{t["table"]}/{t["n_table"]}': t for t in json.load(open("data/kamch/tz_verbatim.json"))}
REG = json.load(open("data/kamch/registry.json")); AS = json.load(open("data/kamch/assign.json"))
DA = {a["tt"]: a for a in json.load(open("data/kamch/dnru_assign.json"))}
CHK = json.load(open("data/kamch/tz_check.json")) if os.path.exists("data/kamch/tz_check.json") else []
FAC = {}
for f in ("rem_krany", "rem_zatvory", "new_sup_krany", "new_sup_other", "dnru_factories"):
    for x in json.load(open(f"data/kamch/{f}.json")): FAC.setdefault(x["company"], x)
nk = lambda s: re.sub(r"[^а-яa-z0-9]", "", re.sub(r"\(.*", "", str(s or "")).lower().replace("ооо", "").replace("оао", "").replace("ао", ""))[:14]
ph = lambda s: str(s or "").split(";")[0].split("\n")[0].strip()


def chk(tt, fac):
    c = [x for x in CHK if x.get("tt") == tt]
    return next((x for x in c if nk(x.get("factory")) and (nk(x["factory"]) in nk(fac) or nk(fac) in nk(x["factory"]))), None)


def contact(name):
    f = FAC.get(name) or next((v for k, v in FAC.items() if nk(k) == nk(name)), {})
    return f"{ph(f.get('phone'))} {ph(f.get('email'))}".strip()


CUR = {"ТД АДЛ (Гранвэл)": "+7 (495) 937-89-68 info@adl.ru", "Лидер-Энерго": "+7 (8482) 21-66-32 trans-energo@bk.ru", "Тулаэлектропривод": "+7 (4872) 72-47-16 td@tulaprivod.ru",
       "КИПИКА (прайс 01.10)": "+7 (8412) 909-200 b18bk@11b18bk.ru", "ДН.РУ": "менеджер по счетам 819892/819896"}
GAR = {"Лидер-Энерго": "60 мес. с ввода в эксплуатацию, не более 66 мес. с даты поставки", "Тулаэлектропривод": "24 мес. с ввода в эксплуатацию, не более 36 мес. с даты отгрузки"}
BAD = "не соответствует"
SEL = {}
for tt, d in rows.items():
    r = REG.get(tt, {}); a = DA.get(tt, {}); s = None
    if d["sup"] and d["act"] in ("Брать", "Брать при условии", "Торг: наш бюджет") and r.get("num"):
        s = dict(src="текущий поставщик", who=d["sup"], contact=CUR.get(d["sup"], ""), product=d["offer"][:150], man=r["man"], num=r["num"], valid=r["valid_to"], tu="",
                 verdict="соответствует (по КП)" if d["act"] == "Брать" else "подтвердить: " + d["why"][:120], ask="" if d["act"] == "Брать" else d["why"][:200], chk=None)
    if not s and a.get("main"):
        for who, reg, val, prod in ((a["main"], a.get("main_reg"), a.get("main_reg_valid"), a.get("main_reg_product")), (a.get("backup"), a.get("backup_reg"), a.get("backup_reg_valid"), a.get("backup_reg_product"))):
            c = chk(tt, who) if who else None
            if who and reg and not (c and c["verdict"].startswith(BAD)):
                s = dict(src="завод вместо ДН.РУ", who=who, contact=contact(who), product=(c or {}).get("product") or prod, man=re.sub(r"\s*\(.*", "", who) + (f' (ИНН {a.get("main_inn")})' if who == a["main"] and a.get("main_inn") else ""),
                         num=reg, valid=val, tu=a.get("main_reg_tu", "") if who == a["main"] else "", verdict=(c or {}).get("verdict", "не проверено построчно"), ask=(c or {}).get("ask", ""), chk=c); break
    if not s and AS.get(tt) and not d["sup"] or (not s and AS.get(tt, {}).get("comp") and d["act"] == "Заказать у других"):
        who = AS[tt]["comp"]; c = chk(tt, who)
        if not (c and c["verdict"].startswith(BAD)):
            s = dict(src="завод (новый запрос КП)", who=who, contact=contact(who), product=(c or {}).get("product", ""), man=r.get("man") or re.sub(r"\s*\(.*", "", who), num=r.get("num", ""),
                     valid=r.get("valid_to", ""), tu="", verdict=(c or {}).get("verdict", "не проверено построчно"), ask=(c or {}).get("ask", "") or ("" if r.get("num") else "№ реестровой записи"), chk=c)
    if not s and d["sup"]:
        s = dict(src="текущий поставщик (номера записи нет)", who=d["sup"], contact=CUR.get(d["sup"], ""), product=d["offer"][:150], man="", num="", valid="", tu="",
                 verdict="нет записи ПП 719 — не проходит нацрежим", ask="изготовитель и № записи ПП 719 или российская замена", chk=None)
    SEL[tt] = s or dict(src="НЕ ЗАКРЫТО", who="", contact="", product="", man="", num="", valid="", tu="", verdict="нет поставщика с записью ПП 719 и соответствием ТЗ",
                        ask=(a.get("why") or "")[:250], chk=None)

wb = openpyxl.load_workbook("data/kamch/forms/pril1.xlsx")
if "Лист1" in wb.sheetnames: del wb["Лист1"]
ws = wb["Лист2"]; ws.title = "Прил.1 — предложение участника"
Y = PatternFill("solid", fgColor="FFF2CC"); G = PatternFill("solid", fgColor="E2EFDA"); O = PatternFill("solid", fgColor="FCE4D6"); RD = PatternFill("solid", fgColor="F8CBAD")
FLD = {21: "Ду, мм", 22: "Проход", 23: "Давление, МПа", 24: "Температура, °С", 25: "Среда", 26: "Материал корпуса", 27: "Климатическое исполнение", 28: "Присоединение",
       29: "Приводное устройство: наименование", 30: "Приводное устройство: технические характеристики", 31: "Установочное положение", 32: "Материал штока, шара, уплотнений", 33: "Габариты"}
from openpyxl.cell.cell import MergedCell
def cell(r, j):
    c = ws.cell(r, j)
    return c if not isinstance(c, MergedCell) else type("N", (), {"value": None, "fill": None})()
tab = None; done = 0
for rr in range(5, ws.max_row + 1):
    a0 = ws.cell(rr, 1).value
    if isinstance(a0, str) and "Таблица" in a0: tab = re.search(r"Таблица\s*(\d\.\d)", a0).group(1); continue
    if a0 is None or not re.match(r"^\d+$", str(a0).strip()) or not ws.cell(rr, 2).value: continue
    tt = f"{tab}/{int(str(a0).strip())}"; s = SEL[tt]; t = TT[tt]; done += 1
    for j in range(19, 36): cell(rr, j).value = None
    ok = s["verdict"].startswith(("соответствует", "эквивалент"))
    if s["product"]: cell(rr, 19).value = s["product"]
    if s["tu"]: cell(rr, 20).value = s["tu"]
    diffs = {x.get("field", ""): x.get("factory") for x in ((s["chk"] or {}).get("diffs") or [])}
    for j, f in FLD.items():
        if ok:
            v = next((dv for k, dv in diffs.items() if k and (k.lower() in f.lower() or f.lower().startswith(k.lower()[:8]))), None)
            cell(rr, j).value = v if v else (t.get(f) or None); cell(rr, j).fill = O if v else G
        else: cell(rr, j).fill = Y
    cell(rr, 34).value = (f'Производитель: {s["man"]}; страна происхождения: Россия; № реестровой записи (ПП 719): {s["num"]}' + (f', действует до {s["valid"]}' if s["valid"] else "")) if s["num"] else None
    cell(rr, 35).value = GAR.get(s["who"])
    for j in (19, 20, 34, 35):
        if not cell(rr, j).value: cell(rr, j).fill = Y
    if s["src"] == "НЕ ЗАКРЫТО" or s["verdict"].startswith("нет записи"): cell(rr, 19).fill = RD
ws.cell(1, 18).value = "Предложение участника по характеристикам (заполнено ООО «ГИДРОПАС»: зелёное — по ТЗ, подтверждено; оранжевое — значение завода; жёлтое — заполнить по ответу завода; красное — не закрыто)"

thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
def sheet(name, H, data, widths):
    w = wb.create_sheet(name); w.append(H)
    for c in w[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496"); c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in data:
        w.append(r)
        for c in w[w.max_row]: c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, x in enumerate(widths, 1): w.column_dimensions[openpyxl.utils.get_column_letter(i)].width = x
    w.freeze_panes = "C2"; w.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(H))}{w.max_row}"
    return w
order = sorted(SEL, key=lambda t: (t.split("/")[0], int(t.split("/")[1])))
sheet("Выбор поставщика", ["Табл./№", "Позиция ТЗ", "Кол-во", "Источник", "Поставщик / завод", "Контакт", "Изделие", "Изготовитель (как в реестре)", "№ записи ПП 719", "Действует до",
                           "Проверка ТЗ", "Что запросить / согласовать"],
      [[t, TT[t]["Наименование"], TT[t]["qty"], SEL[t]["src"], SEL[t]["who"], SEL[t]["contact"], SEL[t]["product"], SEL[t]["man"], SEL[t]["num"], SEL[t]["valid"], SEL[t]["verdict"], SEL[t]["ask"]] for t in order],
      [8, 36, 7, 18, 30, 26, 36, 30, 12, 11, 24, 40])
sheet("Для КП (ТТ п. 1.5)", ["Табл./№", "Наименование", "Ед.", "Кол-во", "Производитель", "Страна происхождения", "№ реестровой записи ПП 719 / ЕАЭС", "Цена за ед. без НДС, ₽ (заполнить)", "Сумма, ₽"],
      [[t, TT[t]["Наименование"], TT[t].get("unit", "шт"), TT[t]["qty"], SEL[t]["man"], "Россия" if SEL[t]["num"] else "", SEL[t]["num"], None, None] for t in order], [8, 44, 5, 7, 34, 12, 18, 16, 14])
kp = wb["Для КП (ТТ п. 1.5)"]
for r in range(2, kp.max_row + 1):
    kp.cell(r, 9).value = f"=D{r}*H{r}"; kp.cell(r, 8).fill = Y
    for j in (5, 6, 7):
        if not kp.cell(r, j).value: kp.cell(r, j).fill = Y
from collections import Counter
cnt = Counter(SEL[t]["src"] for t in order); ver = Counter(("ок" if SEL[t]["verdict"].startswith(("соответствует", "эквивалент")) else SEL[t]["verdict"][:40]) for t in order)
sv = wb.create_sheet("Итог", 0)
sv.append(["Камчатскэнерго 32616346943 — финальная таблица для заполнения заявки"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["Лист «Прил.1 — предложение участника» — официальная форма ТТ (столбцы 1–18 «Предложение участника»): зелёное — по ТЗ, подтверждено; оранжевое — значение завода (эквивалент); жёлтое — заполнить по ответу завода; красное — не закрыто."])
sv.append(["Лист «Для КП» — производитель, страна, № реестровой записи по каждой позиции (ТТ п. 1.5.1–1.5.3); цену вносить после ответов. Номера — открытое зеркало реестра, перед подачей сверить с выписками ГИСП."])
sv.append([]); sv.append(["Источник", "Позиций"])
for k, v in cnt.most_common(): sv.append([k, v])
sv.append([]); sv.append(["Проверка ТЗ", "Позиций"])
for k, v in ver.most_common(): sv.append([k, v])
sv.column_dimensions["A"].width = 60; sv.column_dimensions["B"].width = 10
wb.save("../KAMCHATKA_FINAL_ZAPOLNENIE.xlsx")
print("строк формы", done, dict(cnt), dict(ver))
