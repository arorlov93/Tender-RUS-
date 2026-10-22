"""Камчатскэнерго 32616346943 — построчная сверка 315 позиций с реальными ценами 2026; итоговая себестоимость и маржа."""
import json, os, re, sys, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
from kamch_groups import group

DIR = sys.argv[1]
NMCK = 101102908; TORG = 0.10; FREIGHT = 55
Q = json.load(open("data/kamch/kam_qty.json"))
P = {}
for g in "ABC":
    for x in json.load(open(f"data/kamch/price_{g}.json")):
        P[(re.search(r"1\.\d", x["spec"]).group(0), x["n"])] = x
V = {}
for b in (1, 2, 3, 4):
    f = f"{DIR}/ver_kam_b{b}.json"
    if os.path.exists(f):
        for x in json.load(open(f)): V[(str(x["table"]), int(x["n"]))] = x
FILL = {"ok": "C6EFCE", "занижена": "FCE4D6", "завышена": "DDEBF7", "не тот товар": "FFC7CE", "нет данных": "F2F2F2"}
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; PCT = "0.0%"


def head(ws):
    for c in ws[ws.max_row]:
        if c.value is not None: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")


rows = []
for q in Q:
    k = (q["table"], q["n"]); p = P.get(k, {}); v = V.get(k)
    cur = p.get("price") or 0; rec = (v or {}).get("rec_price") or cur
    rows.append(dict(t=q["table"], n=q["n"], g=group(q["name"]), name=q["name"], qty=q["qty"], cur=cur, rec=rec, med=(v or {}).get("median_price"),
                     verdict=(v or {}).get("verdict") or "в работе", conf=(v or {}).get("confidence") or "", spec=str((v or {}).get("spec_match") or ""),
                     pp=str((v or {}).get("pp1875") or ""), src=(v or {}).get("sources") or [], note=str((v or {}).get("note") or ""), w=p.get("weight_kg") or 0))

wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Итог"
ws = wb.create_sheet("Построчно")
ws.append(["Таблица", "№", "Группа", "Позиция", "Кол-во", "Было, ₽/ед.", "Медиана источников, ₽", "Итоговая цена, ₽/ед.", "Изменение суммы, ₽", "Вердикт", "Надёжность",
           "Соответствие ТЗ", "Нацрежим (ПП 1875)", "Источник 1", "Источник 2", "Источник 3", "Комментарий"]); head(ws)
sh = lambda t, n: t if len(t) <= n else t[:n - 1] + "…"
for r in rows:
    srcs = [f'{s.get("company", "")[:30]} — {s.get("price")} ₽ ({s.get("evidence", "")[:12]})' for s in r["src"][:3]] + ["", "", ""]
    ws.append([r["t"], r["n"], r["g"].split(" (")[0], sh(r["name"], 80), r["qty"], r["cur"], r["med"], r["rec"], (r["rec"] - r["cur"]) * r["qty"], r["verdict"], r["conf"],
               sh(r["spec"], 70), sh(r["pp"], 40), srcs[0], srcs[1], srcs[2], sh(r["note"], 140)])
    X = ws.max_row
    for j in (6, 7, 8, 9): ws.cell(X, j).number_format = M
    ws.cell(X, 10).fill = PatternFill("solid", fgColor=FILL.get(r["verdict"], "FFF2CC"))
    for i, s in enumerate(r["src"][:3]):
        if s.get("url"): ws.cell(X, 14 + i).hyperlink = s["url"]
for row in ws.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([7, 5, 18, 44, 7, 12, 12, 12, 13, 12, 9, 30, 20, 30, 30, 30, 46], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "E2"; ws.auto_filter.ref = f"A1:Q{ws.max_row}"

# итог
G = {}
for r in rows:
    d = G.setdefault(r["g"], dict(pos=0, chk=0, cur=0.0, rec=0.0, w=0.0, cnt={}))
    d["pos"] += 1; d["cur"] += r["cur"] * r["qty"]; d["rec"] += r["rec"] * r["qty"]; d["w"] += r["w"] * r["qty"]
    if r["verdict"] != "в работе": d["chk"] += 1
    d["cnt"][r["verdict"]] = d["cnt"].get(r["verdict"], 0) + 1
sv.append(["Камчатскэнерго 32616346943 — построчная сверка цен 2026 (315 позиций)"]); sv["A1"].font = Font(bold=True, size=13)
sv.append(["Итоговая цена = рекомендованная после сверки с 2–3 источниками 2026 (с НДС, до торга). Себестоимость = итоговая цена × кол-во − 10% торг + доставка 55 ₽/кг до Петропавловска-Камчатского."])
sv.append(["Группа", "Позиций", "Проверено", "ok", "занижена", "завышена", "не тот товар", "нет данных", "Было (рынок), ₽", "Стало (рынок), ₽", "Разница, ₽"]); head(sv)
for g, d in sorted(G.items(), key=lambda kv: -kv[1]["rec"]):
    c = d["cnt"]
    sv.append([g, d["pos"], d["chk"], c.get("ok", 0), c.get("занижена", 0), c.get("завышена", 0), c.get("не тот товар", 0), c.get("нет данных", 0), d["cur"], d["rec"], d["rec"] - d["cur"]])
    for j in (9, 10, 11): sv.cell(sv.max_row, j).number_format = M
T = lambda k: sum(d[k] for d in G.values())
sv.append(["ИТОГО", T("pos"), T("chk"), "", "", "", "", "", T("cur"), T("rec"), T("rec") - T("cur")])
for j in (9, 10, 11): sv.cell(sv.max_row, j).number_format = M
for c in sv[sv.max_row]: c.font = Font(bold=True)
FR = T("w") * FREIGHT; COST = T("rec") * (1 - TORG) + FR


def model(P_, C, days=120):
    pre = P_ - C - 0.008 * P_ - C * 0.2 * days / 365
    return pre - 0.15 * max(pre, 0)


lo, hi = 0.0, 0.95
for _ in range(60):
    m = (lo + hi) / 2
    lo, hi = (m, hi) if model(NMCK * (1 - m), COST) > 0 else (lo, m)
sv.append([])
sv.append(["Показатель", "Значение"]); head(sv)
for a, b, f in [("НМЦК", NMCK, M), ("Закупка по рынку после сверки", T("rec"), M), ("Торг −10%", -T("rec") * TORG, M), ("Доставка (55 ₽/кг)", FR, M),
                ("Себестоимость лота", COST, M), ("Маржа при НМЦК", 1 - COST / NMCK, PCT), ("Безубыточное снижение", lo, PCT)]:
    sv.append([a, b]); sv.cell(sv.max_row, 2).number_format = f
sv.append([])
sv.append(["Снижение", "Цена заявки, ₽", "Прибыль после УСН, ₽"]); head(sv)
for d in (0, 0.07, 0.15, 0.23, 0.30):
    sv.append([d, NMCK * (1 - d), model(NMCK * (1 - d), COST)])
    sv.cell(sv.max_row, 1).number_format = "0%"; sv.cell(sv.max_row, 2).number_format = M; sv.cell(sv.max_row, 3).number_format = M
for row in sv.iter_rows(min_row=3):
    for c in row:
        if c.value is not None: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([46, 18, 18, 7, 9, 9, 11, 10, 16, 16, 14], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

kp = wb.create_sheet("Нужен КП")
kp.append(["Позиции, где цена не подтверждена или сильно изменилась — запросить КП у завода в первую очередь"]); kp["A1"].font = Font(bold=True)
kp.append(["Таблица", "№", "Группа", "Позиция", "Кол-во", "Итоговая цена, ₽/ед.", "Сумма, ₽", "Вердикт", "Комментарий"]); head(kp)
for r in sorted(rows, key=lambda r: -r["rec"] * r["qty"]):
    if r["verdict"] in ("нет данных", "не тот товар") or r["conf"] == "low" or (r["cur"] and abs(r["rec"] / r["cur"] - 1) > 0.3):
        kp.append([r["t"], r["n"], r["g"].split(" (")[0], sh(r["name"], 80), r["qty"], r["rec"], r["rec"] * r["qty"], r["verdict"], sh(r["note"], 140)])
        kp.cell(kp.max_row, 6).number_format = M; kp.cell(kp.max_row, 7).number_format = M
for row in kp.iter_rows(min_row=3):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([7, 5, 18, 50, 7, 14, 14, 12, 60], 1): kp.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
# --- цены для заявки (−15% от НМЦК), распределены пропорционально себестоимости позиции
BID = NMCK * 0.85
cp = {(r["t"], r["n"]): r["rec"] * (1 - TORG) + r["w"] * FREIGHT for r in rows}
k_ = BID / sum(cp[(r["t"], r["n"])] * r["qty"] for r in rows)
bp = wb.create_sheet("Цены для заявки")
bp.append(["Цены для коммерческого предложения: −15% от НМЦК, распределены пропорционально себестоимости позиции (доставка включена). Округление — до рубля."]); bp["A1"].font = Font(bold=True)
bp.append(["Таблица", "№", "Позиция", "Кол-во", "Себес за ед., ₽", "Цена за ед. в заявке, ₽", "Сумма, ₽", "Наценка"]); head(bp)
tot = 0
for r in rows:
    c_ = cp[(r["t"], r["n"])]; u = round(c_ * k_); tot += u * r["qty"]
    bp.append([r["t"], r["n"], sh(r["name"], 90), r["qty"], c_, u, u * r["qty"], u / c_ - 1 if c_ else None])
    for j in (5, 6, 7): bp.cell(bp.max_row, j).number_format = M
    bp.cell(bp.max_row, 8).number_format = PCT
bp.append(["", "", "ИТОГО", "", "", "", tot, None]); bp.cell(bp.max_row, 7).number_format = M
for c in bp[bp.max_row]: c.font = Font(bold=True)
bp.append(["", "", f"Снижение от НМЦК {1 - tot / NMCK:.1%}; прибыль после УСН ≈ {model(tot, COST):,.0f} ₽".replace(",", " ")])
for row in bp.iter_rows(min_row=3):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([8, 5, 60, 7, 14, 16, 15, 9], 1): bp.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
bp.freeze_panes = "D3"
# --- заводы для КП
SUP = {x["group"]: x["suppliers"] for x in json.load(open("data/kamch/suppliers.json"))}
zs = wb.create_sheet("Заводы (КП)")
zs.append(["Группа", "Сумма после сверки, ₽", "Поставщик", "Телефон", "E-mail", "Сайт", "Что закрывает", "Звонок: результат", "Цена КП", "Срок", "Реестр МПТ №"]); head(zs)
for g_, d in sorted(G.items(), key=lambda kv: -kv[1]["rec"]):
    for i, sp in enumerate(SUP.get(g_, [])):
        zs.append([g_ if i == 0 else "", d["rec"] if i == 0 else "", sp.get("company"), (sp.get("phone") + (f"\n[{sp['checked']}]" if sp.get("checked") else "")), sp.get("email"), sp.get("site"), str(sp.get("covers") or "")[:110], "", "", "", ""])
        if i == 0: zs.cell(zs.max_row, 2).number_format = M
for row in zs.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([30, 15, 30, 22, 26, 20, 46, 20, 12, 12, 14], 1): zs.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
zs.freeze_panes = "C2"
# --- условия
us = wb.create_sheet("Условия")
us.append(["Условие", "Что в документации"]); head(us)
for a_, b_ in [("Закупка", "223-ФЗ 32616346943, конкурс в эл. форме, Росэлторг (секция РусГидро), лот 0200-РЕМ ПРОД-2027-КамчЭн"),
               ("Ссылка ЕИС", "https://zakupki.gov.ru/epz/order/notice/notice223/common-info.html?regNumber=32616346943"),
               ("Подача", "до 29.10.2026 10:00 МСК"), ("НМЦК", "101 102 908 ₽ с НДС"),
               ("Поставка", "120 календарных дней с даты договора; Петропавловск-Камчатский (ул. Пограничная 75А, ул. Степная 50, пр. Победы 24), Елизовский р-н, пос. Светлый"),
               ("Обеспечение заявки", "1 005 499,42 ₽ (БГ или деньги)"), ("Аванс", "есть; обеспечение возврата аванса — 100% суммы аванса (БГ)"),
               ("Нацрежим", "ПП 1875 — ограничение и преимущество; номер записи в реестре Минпромторга/ЕАЭС по каждой позиции"),
               ("Документы", "сертификаты/декларации ТР ТС 032/2013 и 010/2011 (схемы 1с/5Д) с протоколами испытаний; техпредложение по Прил. 1"),
               ("Аккредитация", "аккредитация в РусГидро (Прил. 10)"), ("Лот", "неделимый; эквиваленты допускаются с обоснованием"),
               ("Гарантия", "по заводу, не менее 12 мес.; товар изготовлен в год поставки или предыдущий"),
               ("Импорт — заменить", "Ридан NVD (КНР), Regada (Словакия), ESBE (Швеция), Dungs (Германия), Zenova/КШХ, 10нж45фт (ТУ Украины) → российские аналоги из реестра")]:
    us.append([a_, b_])
us.cell(3, 2).hyperlink = "https://zakupki.gov.ru/epz/order/notice/notice223/common-info.html?regNumber=32616346943"
for row in us.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
us.column_dimensions["A"].width = 24; us.column_dimensions["B"].width = 130
tidy(wb, row_h=40)
wb.save("../KAMCHATKA_VERIFY_2010.xlsx")
print("заявка", round(tot), "прибыль@заявка", round(model(tot, COST)), "| проверено", T("chk"), "из", T("pos"), "| рынок было", round(T("cur")), "стало", round(T("rec")), "| себес", round(COST), "маржа", round(1 - COST / NMCK, 3), "| безуб", round(lo, 3),
      "| прибыль −15%", round(model(NMCK * 0.85, COST)))
