"""Из выгрузок 12.10: коммерческие без НМЦК / одежда / топливо — отдельными листами."""
import json, re, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
d=json.load(open("data/triage_1210.json"))
TODAY=dt.datetime(2026,10,13)
CL=re.compile(r"одежд|костюм|куртк|обмундир|форменн|термобел|обув|ботин|берц|СИЗ|средств индивидуальной защиты|мягкого инвентаря",re.I)
FUEL=re.compile(r"дизельн\w* топлив|ГСМ|бензин|мазут|керосин|нефтепродукт",re.I)
JUNK=re.compile(r"мониторинг цен|ремонт|монтаж|работ|услуг|строительств|проектир|аренд|перевооружен|Мебель для склада",re.I)
def end_dt(s):
    try: return dt.datetime.strptime(s[:16],"%Y-%m-%d %H:%M")
    except Exception: return None
def status(e):
    if e is None: return "Срок не указан"
    if e<TODAY: return "Срок прошёл"
    if e.date()==TODAY.date(): return "Подача сегодня"
    return "Актуален"
def cl_type(n):
    if re.search(r"форменн|обмундир",n,re.I): return "Форменная / обмундирование"
    if re.search(r"термобел",n,re.I): return "Термобельё"
    if re.search(r"мягкого инвентаря|воспитанник|студент|Закупка одежды|ГАЛАНТЕРЕЯ",n,re.I): return "Повседневная (соцучреждения)"
    if re.search(r"обув|ботин|берц",n,re.I) and not re.search(r"одежд",n,re.I): return "Спецобувь"
    return "Спецодежда / СИЗ"
def fuel_type(n):
    if re.search(r"дизельн",n,re.I): return "Дизельное топливо"
    return "ГСМ / масла"
REL={str(r[1]) for r in openpyxl.load_workbook("../TENDERS_1210.xlsx",read_only=True)["Коммерческие без НМЦК"].iter_rows(min_row=2,values_only=True) if r[1]}
com=[];cl=[];fuel=[];seen=set()
for x in d:
    n=x["name"]; key=(x["num"],n); e=end_dt(x["end"])
    if key in seen or status(e)=="Срок прошёл": continue
    if FUEL.search(n) and not re.search(r"ремонт|инжектор|автоматик",n,re.I):
        seen.add(key); fuel.append((fuel_type(n),x)); continue
    if (x["cat"] in("Одежда","СИЗ") or CL.search(n)) and not re.search(r"спортивн|Мебель",n,re.I):
        seen.add(key); cl.append((cl_type(n),x)); continue
    if x["law"]=="Коммерческие" and not x["nmck"] and x["num"] in REL and not JUNK.search(n):
        seen.add(key); com.append((x["cat"],x))
H=["№","Тип / категория","Статус","Номер закупки","Предмет","Заказчик","Регион","Закон","Способ","Площадка","НМЦК, ₽","Окончание подачи","Только СМП","Обеспечение заявки","Обеспечение контракта","Аванс","Ссылка","Примечание"]
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
FILL={"Подача сегодня":"FCE4D6"}
wb=openpyxl.Workbook(); wb.remove(wb.active)
def sheet(title,items):
    ws=wb.create_sheet(title); ws.append(H)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True,vertical="center")
    rank={"Подача сегодня":0,"Актуален":1,"Срок не указан":2}
    rows=sorted(((t,x,end_dt(x["end"])) for t,x in items),key=lambda r:(rank[status(r[2])],r[0],-(r[1]["nmck"] or 0),r[2] or TODAY))
    for i,(t,x,e) in enumerate(rows,1):
        st=status(e); note=[]
        if not x["nmck"]: note.append("НМЦК не указана — запросить спецификацию, выставить КП")
        if x["nmck"] and x["nmck"]>50e6: note.append("крупный лот >50 млн")
        if re.search(r"единственн",x["method"] or "",re.I): note.append("способ «у единственного поставщика» — участие только по приглашению")
        link=x["link"] or ""
        ws.append([i,t,st,x["num"],x["name"],x["cust"],x["region"],x["law"],x["method"],x["etp"],x["nmck"],e,
                   x["smp"] or "",x["sec_app"] or "",x["sec_contr"] or "",x["advance"] if x["advance"] not in(None,"") else "",
                   "открыть" if link else "","; ".join(note)])
        r=ws.max_row
        if link: ws.cell(r,17).hyperlink=link; ws.cell(r,17).font=Font(color="0563C1",underline="single")
        ws.cell(r,11).number_format=MONEY; ws.cell(r,12).number_format="DD.MM.YYYY HH:MM"; ws.cell(r,4).number_format="@"
        for c in ws[r]:
            c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
            if st in FILL: c.fill=PatternFill("solid",fgColor=FILL[st])
    for i,w in enumerate([5,24,15,22,50,34,20,13,18,16,16,17,10,14,14,10,9,34],1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes="E2"; ws.auto_filter.ref=f"A1:R{ws.max_row}"
    return rows
import collections
for t,items in(("Коммерческие без НМЦК",com),("Одежда",cl),("Топливо",fuel)):
    rows=sheet(t,items)
    print(t,len(rows),collections.Counter(r[0] for r in rows),"сегодня:",sum(status(r[2])=="Подача сегодня" for r in rows),"с НМЦК:",sum(1 for r in rows if r[1]["nmck"]),round(sum(r[1]["nmck"] or 0 for r in rows)))
wb.save("../COMMERCIAL_CLOTHES_FUEL_1210.xlsx")
