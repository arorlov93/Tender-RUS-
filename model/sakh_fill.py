"""Сахалинэнерго (РусГидро) лоты 11201 и 11202: запрос КП на производстве — по каждому заводу отдельный xlsx с его позициями
(наименование дословно из ТТ, Прил. 1, табл. 3.1) и графами для заполнения (марка, завод, страна, № реестра ПП 719, цена, срок). Без наших цен.
Назначение позиций — data/sakh/assign.json, контакты — data/sakh/factories.json (+ проверенные заводы Камчатки). → ../SAKHALIN_ZAPROS/*.xlsx"""
import json, os, re, shutil, openpyxl
from collections import defaultdict
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

P = json.load(open("data/sakh/tz_verbatim.json")); GEN = json.load(open("data/sakh/tt_general.json"))
AS = json.load(open("data/sakh/assign.json"))
FAC = {}
for f in ("data/kamch/rem_krany.json", "data/kamch/rem_zatvory.json", "data/kamch/new_sup_krany.json", "data/kamch/new_sup_other.json", "data/sakh/factories.json"):
    if os.path.exists(f):
        for x in json.load(open(f)): FAC[x["company"]] = x
ph = lambda s: str(s or "").split(";")[0].split("\n")[0].strip()
key = lambda s: re.sub(r"[^а-яa-z0-9]", "", str(s or "").lower())[:18]


def fac(name):
    if not name: return None
    x = FAC.get(name) or next((v for k, v in FAC.items() if key(k) == key(name) or key(name) in key(k) or key(k) in key(name)), None)
    return x


LOT = {"11201": "лот 1 № 11201-РЕМ ПРОД-2027-СЭ (ОП «Сахалинская ГРЭС»), поставка до 230 кал. дн.", "11202": "лот 2 № 11202-РЕМ ПРОД-2027-СЭ (ОП «Южно-Сахалинская ТЭЦ-1»), поставка до 220 кал. дн."}
A = defaultdict(list); NONE = []
LOTS = ("11202",)  # 28.10: идём только на лот 2 (лот 1 в минусе)
for L in LOTS:
    a = {(x["part"], x["n"]): x for x in AS[L]}
    KP = {(str(y["part"]), y["n"]) for y in json.load(open("data/sakh/sakh_prices.json"))[L] if y.get("kp")}  # КП уже получено — не запрашиваем
    for p in P[L]:
        x = a.get((p["part"], p["n"]), {})
        if (p["part"][-1], p["n"]) in KP: continue
        if x.get("main"): A[x["main"]].append((L, p, x, "основной"))
        else: NONE.append((L, p, x))
        if x.get("backup"): A[x["backup"]].append((L, p, x, "запасной"))
FILL = ["Предлагаемая марка / обозначение", "Производитель (завод)", "Страна происхождения", "№ реестровой записи (ПП 719) или евразийского реестра", "Срок действия записи",
        "Соответствие ТТ: «да» или отличия", "Цена за ед. с НДС, ₽", "Сумма с НДС, ₽", "Срок изготовления / отгрузки", "Вес ед., кг / габариты", "Сертификат / декларация ТР ТС №"]
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); Y = PatternFill("solid", fgColor="FFF2CC")
OUT = "../SAKHALIN_ZAPROS"; shutil.rmtree(OUT, ignore_errors=True); os.makedirs(OUT)
summ = []
for k, (comp, L) in enumerate(sorted(A.items(), key=lambda x: -len(x[1])), 1):
    f = fac(comp) or {}
    L.sort(key=lambda x: (x[0], x[1]["part"], x[1]["n"]))
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Заполнить"
    ws.append([f"Запрос КП — ООО «ГИДРОПАС» (ИНН 9729231965) для ПАО «Сахалинэнерго» (группа РусГидро). Получатель: {comp}"]); ws["A1"].font = Font(bold=True, size=12)
    ws.append(["Просим заполнить жёлтые графы. Наименование и характеристики — дословно из технических требований заказчика (марки носят описательный характер, эквивалент допускается с подробным описанием)."])
    ws.append(["Нацрежим ПП 1875 (ограничение): обязательны производитель и ДЕЙСТВУЮЩИЙ номер реестровой записи ПП 719 / евразийского реестра, производитель должен совпадать с записью. Приложите выписку из ГИСП."])
    ws.append(["Товар новый 2026–2027 гг. выпуска; сертификаты/декларации ТР ТС 010/2011 и 032/2013; гарантия по паспорту завода; упаковка для морской перевозки. Доставка — до г. Южно-Сахалинск (укажите вес и габариты)."])
    ws.append(["Лот", "Партия", "№ п/п в ТТ", "Наименование продукции и требования заказчика (дословно из ТТ)", "Кол-во", "Ед.", *FILL])
    for c in ws[5]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for j in range(7, 7 + len(FILL)): ws.cell(5, j).fill = PatternFill("solid", fgColor="BF8F00")
    for Lt, p, x, role in L:
        pre = [None] * len(FILL)
        if role == "основной" and x.get("offer_mark"): pre[0] = x["offer_mark"]
        if role == "основной":
            nm = (p["name"] + " " + str(x.get("offer_mark") or "")).lower(); rg = None
            if "ЧелябинскСпецГражданСтрой" in comp: rg = ("ООО «ЛД Прайд»", "10826766", "04.05.2031") if re.search(r"латун|11б27", nm) else (None, "10655491", "30.06.2028")
            elif "АЛСО" in comp: rg = (None, "10647852", "")
            elif "Барнаульский котельный" in comp: rg = (None, "10769412", "")
            elif "АДЛ" in comp and re.search(r"гранвэл|зпвс|зптл", nm): rg = (None, "10837159", "13.07.2029")
            elif "ЗЭиМ" in comp and re.search(r"пэм|привод", nm): rg = (None, "10382108", "14.07.2029")
            if f.get("type") == "завод" or (f.get("is_manufacturer")) or rg: pre[1] = (rg and rg[0]) or comp; pre[2] = "Россия"
            if rg: pre[3], pre[4] = rg[1], rg[2] or None
        ws.append([Lt, p["part"].replace("ПАРТИЯ ", ""), p["n"], ILLEGAL_CHARACTERS_RE.sub("", p["name"]), p["qty"], p["unit"]] + pre)
        X = ws.max_row; ws.cell(X, 14).value = f"=M{X}*E{X}"
        for j in range(7, 7 + len(FILL)): ws.cell(X, j).fill = Y
        for j in (13, 14): ws.cell(X, j).number_format = '#,##0.00" ₽"'
        for c in ws[X]: c.border = BOX
    X = ws.max_row; ws.append([""] * 12 + ["ИТОГО", f"=SUM(N6:N{X})"]); ws.cell(ws.max_row, 14).number_format = '#,##0.00" ₽"'
    for i, w in enumerate([7, 7, 6, 60, 7, 5, 26, 26, 14, 26, 13, 24, 13, 14, 14, 14, 20], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "E6"
    tidy(wb, row_h=48)
    fn = f"{k:02d}_" + re.sub(r"[^\wА-Яа-яЁё-]+", "_", re.sub(r"[«»\"()]", "", comp.split("(")[0]))[:40].strip("_") + ".xlsx"
    wb.save(f"{OUT}/{fn}")
    main = [y for y in L if y[3] == "основной"]
    summ.append([k, comp, ph(f.get("phone")), ph(f.get("email")), "да" if f.get("verified") else "нет — уточнить звонком", len(main), len(L) - len(main),
                 ", ".join(f'{y[0][-1]}:{y[1]["part"][-1]}/{y[1]["n"]}' for y in main)[:400], fn])
wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Кому что отправить"
sv.append(["Сахалинэнерго 11201 / 11202 — запросы на производство: по каждому заводу файл в папке SAKHALIN_ZAPROS. Подача заявки 05.11.2026 10:00 МСК — ответы нужны до 31.10"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["№", "Завод / поставщик", "Телефон", "E-mail", "Контакты проверены (2 источника)", "Позиций (основной)", "Позиций (запасной)", "Позиции (лот:партия/№)", "Файл"])
for c in sv[2]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for r in summ:
    sv.append(r)
    for c in sv[sv.max_row]: c.border = BOX
sv.append(["", "ИТОГО", "", "", "", sum(r[5] for r in summ), sum(r[6] for r in summ)]); [setattr(c, "font", Font(bold=True)) for c in sv[sv.max_row]]
for i, w in enumerate([4, 40, 22, 28, 16, 10, 10, 60, 40], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
if NONE:
    nn = wb.create_sheet("Без завода")
    nn.append(["Лот", "Партия", "№", "Наименование (ТТ)", "Кол-во", "Почему не найден"])
    for c in nn[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    for Lt, p, x in NONE: nn.append([Lt, p["part"], p["n"], p["name"], p["qty"], x.get("why", "")])
    for i, w in enumerate([7, 12, 5, 70, 7, 60], 1): nn.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=44)
wb.save(f"{OUT}/00_СВОДКА_кому_что_отправить.xlsx")
for r in summ: print(r[0], r[1][:45], r[2], r[3], r[4], r[5], r[6])
print("без завода", len(NONE), "основных", sum(r[5] for r in summ))
