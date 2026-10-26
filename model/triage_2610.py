"""Отбор лотов из выгрузок 26.10.2026: без аукционов, без работ/услуг, подача с 27.10; наш профиль — арматура, ППУ, пробирки, кабель. → data/triage_2610.json, TRIAGE_2610.xlsx"""
import openpyxl, glob, re, json, datetime as dt, sys, collections
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

SRC = sys.argv[1]
CAT = {"_2.xlsx": "Соль", "_3.xlsx": "Одежда/обувь/СИЗ", "_4.xlsx": "Форма/снаряжение", "_5.xlsx": "СИЗ/спецодежда", "_6.xlsx": "Вентиляция/фильтры", "_7.xlsx": "Арматура/трубы ППУ", "_8_1.xlsx": "Медрасходники (пробирки)", "_9.xlsx": "Кабель (ЭОМ)", "_10.xlsx": "СС/IT"}
BAD = re.compile(r"^(?!.*(поставк|закупк|приобрет)).*(работ|услуг|ремонт|монтаж)|выполнени\w+ работ|оказани\w+ услуг|строительно-монтаж|СМР|проектир|аренд|мониторинг цен|аутсорсинг", re.I)
AUC = re.compile(r"аукцион", re.I)
PROF = [(r"соль|реагент|галит|противогололед", "Соль/реагент"), (r"одежд|обув|ботин|берц|костюм|куртк|СИЗ|форм|спальн|жилет|термобел|бронежилет", "Одежда/СИЗ — 3500/5000 ₽ комплект"), (r"ППУ|пенополиуретан|предизолир", "ППУ — есть заводы (НЗТИ, ПКФ ТЕПЛО, Урал)"), (r"задвиж|затвор|кран|клапан|арматур|вентил", "Арматура — есть цены Камчатки и заводы"),
        (r"пробирк|вакуумн", "Пробирки — МиниМед (себес 10–15 ₽/шт)"), (r"кабел|провод", "Кабель"), (r"фильтр", "Фильтры")]
FROM = dt.datetime(2026, 10, 27)
res, seen = [], set()
for f in sorted(glob.glob(SRC + "/*.xlsx")):
    cat = next(v for k, v in CAT.items() if f.endswith(k))
    rows = list(openpyxl.load_workbook(f, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)); H = rows[1]
    for r in rows[2:]:
        if not r or not r[0]: continue
        d = dict(zip(H, r)); name, num = str(r[2] or ""), str(r[0])
        if (num, name) in seen: continue
        seen.add((num, name))
        try: nm = float(d.get("НМЦ"))
        except Exception: nm = None
        end = d.get("Окончание приема заявок")
        if isinstance(end, str):
            try: end = dt.datetime.strptime(end[:16], "%Y-%m-%d %H:%M")
            except Exception: end = None
        meth = str(d.get("Способ отбора") or "")
        why = []
        if AUC.search(meth): why.append("аукцион")
        if BAD.search(name): why.append("работы/услуги")
        if isinstance(end, dt.datetime) and end < FROM: why.append("срок до 27.10")
        if end is None: why.append("срок не указан")
        prof = next((p for rx, p in PROF if re.search(rx, name, re.I)), "")
        res.append(dict(cat=cat, num=num, name=name, nmck=nm, end=str(end)[:16] if end else "", law=d.get("Тип торгов"), method=meth, etp=d.get("ЭТП"), smp=d.get("СМП, СОНО"),
                        link=d.get("Ссылка на ЕИС"), region=r[24], cust=r[25], sec_app=d.get("Обеспечение заявки"), sec_contr=d.get("Обеспечение контракта"), advance=d.get("Аванс"),
                        place=r[28] if len(r) > 28 else None, prof=prof, ok=not why, why=", ".join(why)))
json.dump(res, open("data/triage_2610.json", "w"), ensure_ascii=False, indent=1, default=str)

MONEY = '#,##0" ₽"'; thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
wb = openpyxl.Workbook(); S = wb.active; S.title = "Сводка"
S.append(["Выгрузки 26.10.2026: отбор без аукционов и работ/услуг, подача с 27.10"]); S["A1"].font = Font(bold=True, size=13)
S.append(["Категория", "Всего строк", "Прошло отбор", "С НМЦК", "Сумма НМЦК прошедших, ₽", "Отсеяно: аукцион", "Отсеяно: работы/услуги", "Отсеяно: срок"])
for c in S[2]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
cats = list(dict.fromkeys(x["cat"] for x in res))
for cat in cats:
    a = [x for x in res if x["cat"] == cat]; ok = [x for x in a if x["ok"]]
    S.append([cat, len(a), len(ok), sum(1 for x in ok if x["nmck"]), sum(x["nmck"] or 0 for x in ok), sum("аукцион" in x["why"] for x in a), sum("работы" in x["why"] for x in a), sum("срок" in x["why"] for x in a)])
    S.cell(S.max_row, 5).number_format = MONEY
for row in S.iter_rows(min_row=2):
    for c in row: c.border = BOX
for i, w in enumerate([26, 11, 12, 9, 20, 14, 18, 12], 1): S.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
H = ["№", "Номер закупки", "Предмет", "Наш профиль", "Заказчик", "Регион", "Закон", "Способ", "Площадка", "НМЦК, ₽", "Подача до", "СМП", "Обеспечение заявки", "Обеспечение контракта", "Аванс", "Ссылка"]
for cat in cats:
    ws = wb.create_sheet(cat.replace("/", " ")[:30]); ws.append(H)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
    ok = sorted([x for x in res if x["cat"] == cat and x["ok"]], key=lambda x: (x["prof"] == "", x["nmck"] is None, -(x["nmck"] or 0)))
    for i, x in enumerate(ok, 1):
        ws.append([i, x["num"], x["name"], x["prof"], x["cust"], x["region"], x["law"], x["method"], x["etp"], x["nmck"], x["end"][:16].replace("T", " "), x["smp"], x["sec_app"], x["sec_contr"], x["advance"], x["link"]])
        R = ws.max_row; ws.cell(R, 2).number_format = "@"; ws.cell(R, 10).number_format = MONEY
        if x["link"]: ws.cell(R, 16).hyperlink = x["link"]; ws.cell(R, 16).font = Font(color="0563C1", underline="single")
        if x["prof"]: ws.cell(R, 4).fill = PatternFill("solid", fgColor="C6EFCE")
        for cc in ws[R]: cc.border = BOX
    for i, w in enumerate([4, 20, 56, 26, 30, 20, 12, 22, 16, 14, 15, 12, 13, 13, 9, 30], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:P{ws.max_row}"
tidy(wb, row_h=36)
wb.save("../TRIAGE_2610.xlsx")
c = collections.Counter((x["cat"], x["ok"]) for x in res)
for cat in cats: print(f"{cat}: всего {c[(cat, True)] + c[(cat, False)]}, прошло {c[(cat, True)]}")
print(collections.Counter(x["why"] for x in res if not x["ok"]).most_common(6))
