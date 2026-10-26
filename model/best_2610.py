"""Выгрузки 26.10.2026: сводка построчных разборов (data/2610/deep_*.json) — лучшие закупки без аукционов и план подготовки к подаче. → BEST_2610.xlsx"""
import json, re, openpyxl, datetime as dt
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

T = {x["num"]: x for x in json.load(open("data/triage_2610.json"))}
D = []
for k in ("A1", "A2", "B", "C", "D"):
    for x in json.load(open(f"data/2610/deep_{k}.json")): x["batch"] = k; D.append(x)
pct = lambda v: None if v is None else (v / 100 if abs(v) > 1.5 else v)
ORD = {"Участвовать": 0, "Осторожно": 1, "Запросить спецификацию": 2, "Нет": 3}
dec = lambda s: next((v for k, v in ORD.items() if str(s).startswith(k)), 3)
GRP = {"A1": "Арматура/ППУ крупные", "A2": "Арматура", "B": "Одежда/СИЗ/соль", "C": "Пробирки", "D": "Видеонаблюдение/IT"}
FILL = {0: "C6EFCE", 1: "FFEB9C", 2: "DDEBF7", 3: "F2F2F2"}
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin); M = '#,##0" ₽"'


def prep_txt(p):
    if isinstance(p, dict):
        return "; ".join(f"{k}: {v if not isinstance(v, list) else ', '.join(map(str, v))}" for k, v in p.items())
    return str(p or "")


def end_of(x):
    e = str(x.get("end") or T.get(x["num"], {}).get("end") or "")
    m = re.search(r"(\d{4})-(\d{2})-(\d{2})[ T](\d{2}):(\d{2})", e) or re.search(r"(\d{2})\.(\d{2})\.(\d{4})\D+(\d{1,2}):(\d{2})", e)
    if not m: return None
    g = m.groups()
    return dt.datetime(int(g[0]), int(g[1]), int(g[2]), int(g[3]), int(g[4])) if len(g[0]) == 4 else dt.datetime(int(g[2]), int(g[1]), int(g[0]), int(g[3]), int(g[4]))


for x in D:
    x["_d"] = dec(x.get("decision")); x["_m"] = pct(x.get("margin_pct")); x["_end"] = end_of(x)
    x["_link"] = T.get(x["num"], {}).get("link") or ""
D.sort(key=lambda x: (x["_d"], -(x.get("net_expected") or 0)))

wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Лучшие — план подачи"
ws.append(["Выгрузки 26.10.2026 — построчный разбор 51 закупки (без аукционов), лучшие и план подготовки к подаче"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Прибыль — после УСН 15% при ожидаемом снижении. Решение «Запросить спецификацию» — документы за логином площадки, расчёт по оценке. Время подачи — МСК."])
H = ["№", "Решение", "Подача до", "Группа", "Номер", "Предмет", "Заказчик", "Способ", "НМЦК, ₽", "Себестоимость + логистика, ₽", "Маржа при НМЦК", "Ожидаемая прибыль, ₽",
     "Безубыточное снижение", "Надёжность", "Что сделать для подачи", "Риски", "Ссылка"]
ws.append(H)
for c in ws[3]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for i, x in enumerate([x for x in D if x["_d"] <= 1], 1):
    ws.append([i, x.get("decision"), x["_end"].strftime("%d.%m.%Y %H:%M") if x["_end"] else x.get("end"), GRP[x["batch"]], x["num"], str(x.get("name"))[:90], str(x.get("cust") or "")[:50],
               str(x.get("method") or T.get(x["num"], {}).get("method") or "")[:30], x.get("nmck"), (x.get("cost_total") or 0) + (x.get("logistics") or 0) or None, x["_m"], x.get("net_expected"),
               pct(x.get("breakeven_disc")), x.get("confidence"), prep_txt(x.get("prep"))[:400], str(x.get("risks") or "")[:300], x["_link"]])
    R = ws.max_row; ws.cell(R, 5).number_format = "@"
    for j in (9, 10, 12): ws.cell(R, j).number_format = M
    for j in (11, 13): ws.cell(R, j).number_format = "0%"
    ws.cell(R, 2).fill = PatternFill("solid", fgColor=FILL[x["_d"]])
    if x["_link"]: ws.cell(R, 17).hyperlink = x["_link"]
    for c in ws[R]: c.border = BOX
for i, w in enumerate([4, 14, 15, 16, 20, 40, 28, 18, 14, 15, 9, 14, 10, 9, 60, 50, 28], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "F4"

al = wb.create_sheet("Все разобранные (51)")
al.append(H)
for c in al[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for i, x in enumerate(D, 1):
    al.append([i, x.get("decision"), x["_end"].strftime("%d.%m.%Y %H:%M") if x["_end"] else x.get("end"), GRP[x["batch"]], x["num"], str(x.get("name"))[:90], str(x.get("cust") or "")[:50],
               str(x.get("method") or "")[:30], x.get("nmck"), (x.get("cost_total") or 0) + (x.get("logistics") or 0) or None, x["_m"], x.get("net_expected"), pct(x.get("breakeven_disc")),
               x.get("confidence"), prep_txt(x.get("prep"))[:400], str(x.get("risks") or "")[:300], x["_link"]])
    R = al.max_row; al.cell(R, 5).number_format = "@"
    for j in (9, 10, 12): al.cell(R, j).number_format = M
    for j in (11, 13): al.cell(R, j).number_format = "0%"
    al.cell(R, 2).fill = PatternFill("solid", fgColor=FILL[x["_d"]])
    for c in al[R]: c.border = BOX
for i, w in enumerate([4, 14, 15, 16, 20, 40, 28, 18, 14, 15, 9, 14, 10, 9, 60, 50, 28], 1): al.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
al.freeze_panes = "F2"; al.auto_filter.ref = f"A1:Q{al.max_row}"

ln = wb.create_sheet("Построчно")
ln.append(["Номер", "Решение по лоту", "Позиция", "Кол-во", "Ед.", "НМЦК за ед., ₽", "Себестоимость за ед., ₽", "Маржа позиции", "Источник цены"])
for c in ln[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for x in D:
    for l in (x.get("lines") or x.get("positions") or []):
        if not isinstance(l, dict): continue
        nu, cu = l.get("nmck_unit"), l.get("cost_unit")
        ln.append([x["num"], x.get("decision"), str(l.get("item"))[:100], l.get("qty"), l.get("unit"), nu, cu, (1 - cu / nu) if nu and cu else None, str(l.get("cost_src") or "")[:150]])
        R = ln.max_row; ln.cell(R, 1).number_format = "@"
        for j in (6, 7): ln.cell(R, j).number_format = M
        ln.cell(R, 8).number_format = "0%"
for i, w in enumerate([20, 14, 56, 8, 6, 13, 13, 9, 60], 1): ln.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ln.freeze_panes = "C2"; ln.auto_filter.ref = f"A1:I{ln.max_row}"
tidy(wb, row_h=48)
wb.save("../BEST_2610.xlsx")
for x in D:
    if x["_d"] <= 1: print(x.get("decision")[:11], x["_end"], x["num"], round(x.get("nmck") or 0), x["_m"], round(x.get("net_expected") or 0), str(x.get("name"))[:45])
print(len(D), "лотов; строк построчно", ln.max_row - 1)
