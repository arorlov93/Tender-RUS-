"""Одежда, спецодежда, СИЗ, форма — будущие торги (с 20.10, без аукционов), от самой большой НМЦК к меньшей."""
import io, contextlib, re, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
with contextlib.redirect_stdout(io.StringIO()):
    import future_1910 as f

NOTCL = re.compile(r"палатк|огнев|спальн|ОБЗР|игруш|питбайк|оборудовани|шатер|тамбур|Снаряжение тактическое|палаток", re.I)
C = [r for r in f.out if re.search(r"одежд|СИЗ|Форма", r["cat"], re.I) and (not NOTCL.search(r["name"]) or re.search(r"костюм|одежд", r["name"], re.I))]
C.sort(key=lambda r: (r["nmck"] is None, -(r["nmck"] or 0)))
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'
FILL = {"ПОДАЁМ": "C6EFCE", "Осторожно": "FFEB9C", "НЕТ": "FFC7CE", "КП": "DDEBF7", "Посчитать": "F2F2F2"}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Одежда и СИЗ"
ws.append([f"Одежда, спецодежда, СИЗ, форма — подача с 20.10.2026, без аукционов. От большей НМЦК к меньшей. Всего {len(C)}"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Себестоимость одежды: летний комплект 3 500 ₽, зимний 5 000 ₽; прочее — справочник цен 2026. Прибыль — после УСН 15% при ожидаемом снижении. Без НМЦК — внизу (коммерческие запросы КП)."])
ws.append(["№", "НМЦК, ₽", "Подача до", "Номер", "Предмет", "Заказчик", "Закон / способ", "Себестоимость, ₽", "Маржа при НМЦК", "Прибыль после УСН, ₽", "Решение", "Условия"])
for c in ws[3]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
sh = lambda t, n: t if len(t) <= n else t[:n - 1] + "…"
for i, r in enumerate(C, 1):
    ws.append([i, r["nmck"], r["e"], r["num"], sh(r["name"], 110), sh(r["cust"], 60), f'{r["law"]} · {r["meth"]}', r["C"], r["gm"], r["net"], r["dec"], sh(r["cond"], 140)])
    R = ws.max_row
    for j in (2, 8, 10): ws.cell(R, j).number_format = M
    ws.cell(R, 9).number_format = "0%"; ws.cell(R, 3).number_format = "DD.MM HH:MM"; ws.cell(R, 4).number_format = "@"
    ws.cell(R, 11).fill = PatternFill("solid", fgColor=next(v for k, v in FILL.items() if r["dec"].startswith(k)))
    for c in ws[R]: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
ws.append(["", sum(r["nmck"] or 0 for r in C), "", "", "ИТОГО НМЦК"]); ws.cell(ws.max_row, 2).number_format = M
for c in ws[ws.max_row]: c.font = Font(bold=True)
for i, w in enumerate([4, 15, 11, 20, 50, 34, 26, 14, 9, 14, 24, 44], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "E4"; ws.auto_filter.ref = f"A3:L{ws.max_row - 1}"
tidy(wb, row_h=40)
wb.save("../CLOTHES_BY_NMCK_1910.xlsx")
for r in C: print(r["num"], r["nmck"], r["dec"][:20], r["name"][:50])
