"""Все лоты с подачей с 25.10.2026 из выгрузок 19.10 — без аукционов и работ/услуг."""
import json, re, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
d=json.load(open("data/triage_1910.json"))
F=dt.datetime(2026,10,25)
NOT=re.compile(r"канцеляр|хозяйств|игрушк|реактив|лаб\. посуд|стоматолог|газон|офисн|учебн|театральн|сварк|мотопомп|инструмент|тара |укрывн|СМР|переустройств|реализаци|невостребован|запчаст|специнвентар|оснащени|пневматик|ЗИП|компьютерн|оргтехник|стройматериал",re.I)
sel=[]
for x in d:
    if "аукцион" in (x["method"] or "").lower(): continue
    if re.search(r"работ|услуг|ремонт|монтаж",x["name"],re.I): continue
    try: e=dt.datetime.strptime(x["end"],"%Y-%m-%d %H:%M")
    except Exception: continue
    if e<F: continue
    x=dict(x); x["e"]=e; sel.append(x)
seen={}
for x in sel:
    k=(x["name"].strip().lower()[:60],(x["cust"] or "")[:30],x["nmck"])
    x["dup"]="дубль: "+seen[k] if k in seen else ""
    seen.setdefault(k,x["num"])
    x["prof"]="не наш профиль" if NOT.search(x["name"]) else "наш профиль"
ORDER=["Соль","Одежда/СИЗ","Форма/снаряжение","СИЗ (респираторы)","Медрасходники (пробирки)","Сантехника/арматура","Кабель (ЭОМ)","СС/видеонаблюдение","Вентиляция/фильтры"]
sel.sort(key=lambda x:(x["prof"]!="наш профиль",ORDER.index(x["cat"]),x["e"],-(x["nmck"] or 0)))
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Подача с 25.10"
ws.append(["Лоты с подачей с 25.10.2026 — без аукционов (выгрузки 19.10)"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["Отбор: способ не аукцион; не работы/услуги; окончание подачи ≥ 25.10. Сверху — наш профиль, ниже — прочее. Дубли (одна закупка на разных площадках) помечены."])
H=["№","Категория","Профиль","Номер закупки","Предмет","Заказчик","Регион","Закон","Способ","Площадка","НМЦК, ₽","Окончание подачи","СМП","Обеспечение заявки","Обеспечение контракта","Аванс","Дубль"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
for i,x in enumerate(sel,1):
    ws.append([i,x["cat"],x["prof"],x["num"],x["name"],x["cust"],x["region"],x["law"],x["method"],x["etp"],x["nmck"],x["e"],x["smp"],x["sec_app"],x["sec_contr"],x["advance"],x["dup"]])
    R=ws.max_row; ws.cell(R,4).number_format="@"; ws.cell(R,11).number_format=MONEY; ws.cell(R,12).number_format="DD.MM.YYYY HH:MM"
    ws.cell(R,3).fill=PatternFill("solid",fgColor="C6EFCE" if x["prof"]=="наш профиль" else "F2F2F2")
    if x["dup"]:
        for cc in ws[R]: cc.font=Font(color="808080",italic=True)
    for cc in ws[R]: cc.border=Border(top=thin,bottom=thin,left=thin,right=thin)
for i,w in enumerate([4,20,13,22,56,32,20,12,24,18,15,15,10,13,13,9,22],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="E4"; ws.auto_filter.ref=f"A3:Q{ws.max_row}"
from tidy import tidy; tidy(wb,row_h=30)
wb.save("../FROM_2510_1910.xlsx")
import collections
ours=[x for x in sel if x["prof"]=="наш профиль" and not x["dup"]]
print(len(sel),"всего;",len(ours),"наш профиль без дублей")
for x in ours: print(x["cat"][:14],'|',x["num"],'|',x["name"][:60],'|',x["nmck"],'|',x["e"].strftime('%d.%m %H:%M'),'|',x["method"][:24],'|',(x["cust"] or "")[:28])
