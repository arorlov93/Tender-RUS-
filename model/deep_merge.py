"""Сводная таблица глубокого анализа цен и маржи (одежда, коммерческие без НМЦК, топливо)."""
import json, os, shutil, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
SRC="/tmp/claude-0/-home-user-Tender-RUS-/e981d7a6-b89f-5df6-824c-c1907918e938/scratchpad/deep"
DST="data/deep"
for f in ["out_clothes_comm.json","out_comm_arm.json","out_comm_other.json","out_fuel.json","bench_arm.json","bench_clothes.json","fuel_note.md"]:
    if os.path.exists(f"{SRC}/{f}"): shutil.copy(f"{SRC}/{f}",f"{DST}/{f}")
def load(f):
    if os.path.exists(f"{DST}/{f}"): return json.load(open(f"{DST}/{f}"))
    if f=="out_clothes_comm.json" and os.path.exists(f"{SRC}/clothes_comm.json"):
        out=[]
        for x in json.load(open(f"{SRC}/clothes_comm.json")):
            exp=str(x.get("end",""))[:10]<"2026-10-14"
            out.append(dict(num=x["num"],name=x["name"],cust=x["cust"],law=x["law"],method=x["method"],end=x["end"],nmck=None,
                decision="Нет" if exp else "Запросить спецификацию",confidence="low",positions=[],
                risks=("Срок подачи истёк. " if exp else "")+"Спецификация закрыта логином площадки. Ориентир для КП: себестоимость по правилу клиента (летний 3 500 / зимний 5 000) или по листу «Цены — одежда (рынок)» + наценка 25–40% → маржа 20–28%. В коммерческих закупках часто сравнивают цены без НДС — для УСН это минус ~18%."))
        return out
    return []
def pct(v):
    if v is None or isinstance(v,str): return None
    return v/100 if abs(v)>1.5 else v
SHEETS=[("Одежда с НМЦК","out_clothes_nmck.json"),("Одежда коммерческая","out_clothes_comm.json"),
        ("Арматура (коммерч.)","out_comm_arm.json"),("Прочие коммерч.","out_comm_other.json"),("Топливо","out_fuel.json")]
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
DFILL={"Участвовать":"C6EFCE","Осторожно":"FFEB9C","Нет":"F2F2F2","Запросить спецификацию":"DDEBF7"}
H=["№","Решение","Номер","Предмет","Заказчик","Способ","Подача","НМЦК, ₽","Себестоимость, ₽","Логистика, ₽","Маржа при НМЦК","Чистая прибыль при НМЦК, ₽","Ожид. снижение","Чистая прибыль при ожид. снижении, ₽","Безубыточное снижение","Рекоменд. цена КП, ₽","Маржа КП","Достоверность","Риски / комментарий","Источники"]
def hdr(ws,h):
    ws.append(h)
    for c in ws[ws.max_row]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True,vertical="center")
wb=openpyxl.Workbook(); summ=wb.active; summ.title="Сводка"
ORDER={"Участвовать":0,"Осторожно":1,"Запросить спецификацию":2,"Нет":3}
allpos=[]; stats=[]
for title,f in SHEETS:
    d=load(f)
    if not d: continue
    ws=wb.create_sheet(title); hdr(ws,H)
    d.sort(key=lambda x:(ORDER.get(x.get("decision"),9),-(x.get("net_expected") or 0)))
    for i,x in enumerate(d,1):
        conf={"high":"высокая","med":"средняя","low":"низкая"}.get(x.get("confidence"),x.get("confidence"))
        ws.append([i,x.get("decision"),str(x.get("num")),x.get("name"),x.get("cust"),f'{x.get("law") or ""} {x.get("method") or ""}'.strip(),str(x.get("end") or "")[:16],
                   x.get("nmck"),x.get("cost_total"),x.get("logistics"),pct(x.get("margin_pct")),x.get("net_at_nmck"),pct(x.get("exp_disc")),x.get("net_expected"),pct(x.get("breakeven_disc")),
                   x.get("kp_price"),pct(x.get("kp_margin_pct")),conf,x.get("risks"),"\n".join((x.get("sources") or [])[:6])])
        r=ws.max_row; ws.cell(r,3).number_format="@"
        for k in (8,9,10,12,14,16): ws.cell(r,k).number_format=MONEY
        for k in (11,13,15,17): ws.cell(r,k).number_format="0%"
        f_=DFILL.get(x.get("decision"))
        for c in ws[r]:
            c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
            if f_: c.fill=PatternFill("solid",fgColor=f_)
        for p in x.get("positions") or []:
            allpos.append([title,str(x.get("num")),x.get("cust"),p.get("item"),p.get("qty"),p.get("unit"),p.get("nmck_unit"),p.get("cost_unit"),
                           (p["nmck_unit"]-p["cost_unit"])/p["nmck_unit"] if p.get("nmck_unit") and p.get("cost_unit") and isinstance(p["nmck_unit"],(int,float)) and isinstance(p["cost_unit"],(int,float)) and p["nmck_unit"] else None,p.get("cost_src")])
    for i,w in enumerate([5,14,20,40,30,22,16,15,15,12,10,16,10,18,11,15,9,11,60,40],1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes="D2"; ws.auto_filter.ref=f"A1:T{ws.max_row}"
    for dec in ORDER:
        s=[x for x in d if x.get("decision")==dec]
        stats.append([title,dec,len(s),sum(x.get("nmck") or 0 for x in s),sum(x.get("net_expected") or 0 for x in s)])
hdr(summ,["Лист","Решение","Лотов","НМЦК, ₽","Чистая прибыль при ожид. снижении, ₽"])
for s in stats:
    if s[2]: summ.append(s)
for row in summ.iter_rows(min_row=2):
    row[3].number_format=MONEY; row[4].number_format=MONEY
    f_=DFILL.get(row[1].value)
    if f_:
        for c in row: c.fill=PatternFill("solid",fgColor=f_)
for i,w in enumerate([24,24,8,20,26],1): summ.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
summ.auto_filter.ref=f"A1:E{summ.max_row}"
ws=wb.create_sheet("Позиции"); hdr(ws,["Лист","Номер","Заказчик","Позиция","Кол-во","Ед.","Цена за ед. в НМЦК, ₽","Наша закупка за ед., ₽","Маржа позиции","Источник цены"])
for p in allpos:
    ws.append(p); r=ws.max_row; ws.cell(r,2).number_format="@"
    for k in (7,8): ws.cell(r,k).number_format='#,##0.00" ₽"'
    ws.cell(r,9).number_format="0%"
for i,w in enumerate([18,20,30,50,9,8,14,14,10,60],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="E2"; ws.auto_filter.ref=f"A1:J{ws.max_row}"
for title,f in (("Цены — одежда (рынок)","bench_clothes_market.json"),("Цены — арматура","bench_arm.json"),("Цены — одежда (тендеры)","bench_clothes.json")):
    b=load(f)
    if isinstance(b,dict): b=b.get('items') or next((v for v in b.values() if isinstance(v,list)),[])
    if not b: continue
    ws=wb.create_sheet(title); keys=list(dict.fromkeys(k for x in b for k in x.keys()))
    hdr(ws,keys)
    for x in b: ws.append([("\n".join(map(str,v)) if isinstance(v,list) else (json.dumps(v,ensure_ascii=False) if isinstance(v,dict) else v)) for v in (x.get(k) for k in keys)])
    for i in range(1,len(keys)+1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=28
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical="top")
    ws.auto_filter.ref=f"A1:{openpyxl.utils.get_column_letter(len(keys))}{ws.max_row}"
wb.save("../DEEP_MARGIN_1210.xlsx")
for s in stats:
    if s[2]: print(s)
print("positions",len(allpos))
