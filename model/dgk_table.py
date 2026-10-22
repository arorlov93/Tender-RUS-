"""АО «ДГК» 32616403815 — арматура среднего и низкого давления (1467 позиций): цены 2026, себестоимость, маржа, заводы для КП."""
import json, os, re, sys, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
from kamch_groups import group

PR = sys.argv[1] if len(sys.argv) > 1 else ""
NMCK = 186531454.93; TORG = 0.10; FREIGHT = 45  # ₽/кг до Хабаровска/Владивостока (ж/д контейнер, оценка)
Q = json.load(open("data/dgk/dgk_positions.json"))
U = {u["key"]: u for u in json.load(open("data/dgk/dgk_unique.json"))}
PRICE = {}
for f in ("price_zadv", "price_kran", "price_klap", "price_other"):
    p = os.path.join(PR, f + ".json")
    if PR and os.path.exists(p):
        for x in json.load(open(p)):
            if x.get("price"): PRICE[x["key"]] = x
KEYOF = {}
for k, u in U.items():
    for n in u["positions"]: KEYOF[n] = k
SUP = {x["group"]: x["suppliers"] for x in json.load(open("data/kamch/suppliers.json"))}
BIG = lambda g: next((k for k in ("Краны", "Задвижки", "Затворы", "Клапаны") if g.startswith(k)), "Приводы и прочее")
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; P = "0.0%"


def head(ws):
    for c in ws[ws.max_row]:
        if c.value is not None: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")


def wt(name, dn):
    # оценка веса, если поставщик не дал: по DN
    d = int(dn) if str(dn).isdigit() else 50
    return max(0.3, 0.0009 * d ** 1.9) if not name.lower().startswith("кран") else max(0.2, 0.0005 * d ** 1.9)


rows = []
for q in Q:
    k = KEYOF.get(q["n"]); u = U.get(k, {}); pr = PRICE.get(k)
    price = (pr or {}).get("price") or u.get("kam_price")
    src = (pr or {}).get("supplier") or (u.get("kam_src") and f'Камчатка-сверка: {u["kam_src"]}') or ""
    dn = k.split("|")[2] if k else ""
    w = (pr or {}).get("weight_kg") or wt(q["name"], dn)
    g = group(q["name"])
    rows.append(dict(n=q["n"], name=q["name"], qty=q["qty"], year=q["year"], g=g, big=BIG(g), price=price, src=src, url=(pr or {}).get("url", ""),
                     match=(pr or {}).get("match", "по Камчатке" if u.get("kam_price") else ""), pp=(pr or {}).get("pp1875", ""), w=w))
have = [r for r in rows if r["price"]]
mkt = sum(r["price"] * r["qty"] for r in have)
cov = len(have) / len(rows)
# себестоимость: найденные — цена×(1−торг)+доставка; ненайденные — экстраполяция по средней доле в группе
G = {}
for r in rows:
    d = G.setdefault(r["big"], dict(pos=0, priced=0, qty=0, mkt=0.0, fr=0.0))
    d["pos"] += 1; d["qty"] += r["qty"]
    if r["price"]: d["priced"] += 1; d["mkt"] += r["price"] * r["qty"]; d["fr"] += r["w"] * r["qty"] * FREIGHT
for d in G.values():
    d["cost"] = d["mkt"] * (1 - TORG) + d["fr"]
    d["est"] = d["cost"] / d["priced"] * (d["pos"] - d["priced"]) if 0 < d["priced"] < d["pos"] else 0
COST = sum(d["cost"] + d["est"] for d in G.values())


def model(P_, C, days=120):
    pre = P_ - C - 0.008 * P_ - C * 0.2 * days / 365
    return pre - 0.15 * max(pre, 0)


lo, hi = 0.0, 0.95
for _ in range(60):
    m = (lo + hi) / 2
    lo, hi = (m, hi) if model(NMCK * (1 - m), COST) > 0 else (lo, m)

wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Итог"
sv.append(["АО «ДГК» 32616403815 — арматура среднего и низкого давления (Хабаровский, Приморский края, Амурская обл., ЕАО). Подача до 30.10.2026"]); sv["A1"].font = Font(bold=True, size=12)
sv.append([f"1467 позиций, 9 422 ед., 413 уникальных товаров. Цены 2026 с НДС; себестоимость = цена × кол-во − 10% торг + доставка {FREIGHT} ₽/кг. Позиций с ценой: {len(have)} ({cov:.0%})."])
sv.append(["Группа", "Позиций", "С ценой", "Единиц", "Закупка по рынку (найденные), ₽", "Себестоимость найденных, ₽", "Оценка ненайденных, ₽", "Себестоимость группы, ₽"]); head(sv)
for gname, d in sorted(G.items(), key=lambda kv: -(kv[1]["cost"] + kv[1]["est"])):
    sv.append([gname, d["pos"], d["priced"], d["qty"], d["mkt"], d["cost"], d["est"], d["cost"] + d["est"]])
    for j in (5, 6, 7, 8): sv.cell(sv.max_row, j).number_format = M
sv.append(["ИТОГО", len(rows), len(have), sum(r["qty"] for r in rows), mkt, sum(d["cost"] for d in G.values()), sum(d["est"] for d in G.values()), COST])
for j in (5, 6, 7, 8): sv.cell(sv.max_row, j).number_format = M
for c in sv[sv.max_row]: c.font = Font(bold=True)
sv.append([])
sv.append(["Показатель", "Значение"]); head(sv)
for a, b, f in [("НМЦК", NMCK, M), ("Себестоимость лота", COST, M), ("Маржа при НМЦК", 1 - COST / NMCK, P), ("Безубыточное снижение", lo, P)]:
    sv.append([a, b]); sv.cell(sv.max_row, 2).number_format = f
sv.append([])
sv.append(["Снижение", "Цена заявки, ₽", "Прибыль после УСН, ₽"]); head(sv)
for d in (0, 0.07, 0.10, 0.15, 0.20, 0.25):
    sv.append([d, NMCK * (1 - d), model(NMCK * (1 - d), COST)])
    sv.cell(sv.max_row, 1).number_format = "0%"; sv.cell(sv.max_row, 2).number_format = M; sv.cell(sv.max_row, 3).number_format = M
for row in sv.iter_rows(min_row=3):
    for c in row:
        if c.value is not None: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([30, 18, 18, 10, 20, 20, 18, 20], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ws = wb.create_sheet("Позиции (1467)")
ws.append(["№ ТТ", "Позиция", "Год", "Кол-во", "Группа", "Цена за ед., ₽", "Сумма, ₽", "Поставщик / источник", "Соответствие", "Нацрежим", "Ссылка"]); head(ws)
for r in rows:
    ws.append([r["n"], r["name"][:110], r["year"], r["qty"], r["g"].split(" (")[0], r["price"], (r["price"] or 0) * r["qty"] if r["price"] else None, str(r["src"])[:60],
               str(r["match"])[:40], str(r["pp"])[:30], r["url"]])
    for j in (6, 7): ws.cell(ws.max_row, j).number_format = M
    if not r["price"]: ws.cell(ws.max_row, 6).fill = PatternFill("solid", fgColor="FFF2CC")
    if r["url"]: ws.cell(ws.max_row, 11).hyperlink = r["url"]
for row in ws.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([6, 60, 6, 7, 22, 13, 14, 30, 22, 18, 36], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:K{ws.max_row}"

# заводы для КП по укрупнённым группам
zs = wb.create_sheet("Заводы (КП)")
zs.append(["Группа", "Себестоимость группы, ₽", "Поставщик", "Телефон", "E-mail", "Сайт", "Что закрывает", "Результат / цена КП / срок"]); head(zs)
SB = {}
for sg, lst in SUP.items():
    for sp in lst:
        key = re.sub(r"[^a-zа-я0-9]", "", str(sp.get("site") or sp.get("company") or "").lower())[:20]
        SB.setdefault(BIG(sg), {}).setdefault(key, sp)
for gname, d in sorted(G.items(), key=lambda kv: -(kv[1]["cost"] + kv[1]["est"])):
    for i, sp in enumerate(SB.get(gname, {}).values()):
        zs.append([gname if i == 0 else "", (d["cost"] + d["est"]) if i == 0 else "", sp.get("company"), (sp.get("phone") + (f"\n[{sp['checked']}]" if sp.get("checked") else "")), sp.get("email"), sp.get("site"), str(sp.get("covers") or "")[:100], ""])
        if i == 0: zs.cell(zs.max_row, 2).number_format = M
for row in zs.iter_rows(min_row=2):
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([18, 18, 32, 22, 26, 20, 46, 26], 1): zs.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=30)
wb.save("../DGK_32616403815.xlsx")
print("с ценой", len(have), "из", len(rows), "| рынок", round(mkt), "| себес", round(COST), "| маржа", round(1 - COST / NMCK, 3), "| безуб", round(lo, 3), "| прибыль −10%", round(model(NMCK * 0.9, COST)))
