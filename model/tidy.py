"""Единая вёрстка xlsx: одинаковая высота строк, перенос в шапке и текстовых колонках, печать по ширине."""
from openpyxl.styles import Alignment
def tidy(wb, row_h=32, head_h=48):
    for ws in wb:
        hr=next((c.row for row in ws.iter_rows(max_row=6) for c in row
                 if c.fill is not None and c.fill.fgColor is not None and str(c.fill.fgColor.rgb).endswith("305496")),1)
        from openpyxl.utils import get_column_letter
        import math
        hh=head_h
        for c in ws[hr]:
            if isinstance(c.value,str):
                w=ws.column_dimensions[get_column_letter(c.column)].width or 10
                hh=max(hh,min(120,15*math.ceil(len(c.value)/max(1,w*0.95))+6))
        for row in ws.iter_rows():
            r=row[0].row
            for c in row:
                if r<hr: c.alignment=Alignment(wrap_text=False,vertical="center")
                elif r==hr: c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
                else: c.alignment=Alignment(wrap_text=isinstance(c.value,str),vertical="center")
            ws.row_dimensions[r].height=hh if r==hr else (18 if r<hr else row_h)
        ws.page_setup.orientation="landscape"; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        ws.sheet_properties.pageSetUpPr.fitToPage=True
