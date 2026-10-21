"""Все 9 выгрузок 19.10: закупки с подачей не раньше 25.10.2026, без аукционов и работ — с маржой (расчёт или ориентир), решением и ссылками."""
import io, contextlib, json, os, re, sys, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy
with contextlib.redirect_stdout(io.StringIO()):
    import eval_all as e

SP = sys.argv[1] if len(sys.argv) > 1 else ""
EV = {r["num"]: r for r in e.rows}
T = json.load(open("data/triage_1910.json"))
L = json.load(open("data/go_links_2010.json"))
EXTRA = {}
for fn in ("missed_out.json", "rest_out.json"):
    p = os.path.join(SP, fn)
    if SP and os.path.exists(p):
        for x in json.load(open(p)): EXTRA[x["num"]] = x
pct = lambda v: None if v is None else (v / 100 if v > 1 else v)
WORK = re.compile(r"^(?!.*(поставк|закупк|приобрет)).*(работ|услуг|ремонт|монтаж)|выполнени\w+ работ|оказани\w+ услуг|реализаци|невостребован", re.I)
NOTPROF = re.compile(r"канцеляр|стройматериал|компьютерн|монтажных нож|игруш|хозяйств", re.I)
CAT_M = {"Одежда/СИЗ": "25–35%", "Форма/снаряжение": "25–35%", "СИЗ (респираторы)": "20–30%", "Медрасходники (пробирки)": "30–50%",
         "Сантехника/арматура": "15–25%", "Кабель (ЭОМ)": "10–15%", "СС/видеонаблюдение": "15–20%", "Вентиляция/фильтры": "15–20%", "Соль": "20–40%"}
SKIP = {"32616453117", "32616452870"}
rows, seen = [], set()
for x in T:
    if "аукцион" in (x["method"] or "").lower(): continue
    try: en = dt.datetime.strptime(x["end"], "%Y-%m-%d %H:%M")
    except Exception: continue
    exe = str((EXTRA.get(x["num"]) or {}).get("end") or "")  # перенос срока по изменениям извещения
    m_ = re.search(r"(\d{2})\.(\d{2})\.(\d{4})\D+(\d{1,2}):(\d{2})", exe) or re.search(r"(\d{4})-(\d{2})-(\d{2})\D+(\d{1,2}):(\d{2})", exe)
    if m_:
        g_ = m_.groups()
        en = dt.datetime(int(g_[2]), int(g_[1]), int(g_[0]), int(g_[3]), int(g_[4])) if len(g_[0]) == 2 else dt.datetime(int(g_[0]), int(g_[1]), int(g_[2]), int(g_[3]), int(g_[4]))
    if en < dt.datetime(2026, 10, 25) or WORK.search(x["name"]) or "СМР" in x["num"] or x["num"] in SKIP: continue
    k = (x["name"][:60].lower(), (x["cust"] or "")[:30], x["nmck"])
    if k in seen: continue
    seen.add(k)
    r = dict(num=x["num"], end=en, cat=x["cat"], cust=x["cust"] or "", subj=x["name"], law=f'{x["law"]} · {x["method"]}', etp=x["etp"] or "", nmck=x["nmck"],
             m=None, net=None, conf="", dec="", mtxt="")
    ev = EV.get(x["num"]); ex = EXTRA.get(x["num"])
    if ex and ex.get("nmck"):
        C = (ex.get("cost_total") or 0) + (ex.get("logistics") or 0)
        r.update(m=pct(ex.get("margin_pct")) if ex.get("margin_pct") is not None else (1 - C / ex["nmck"] if C else None), net=ex.get("net_expected"),
                 conf={"low": "оценка", "med": "средняя", "high": "высокая"}.get(ex.get("confidence"), ""), dec=ex.get("decision", ""), note=str(ex.get("risks") or "")[:160])
    elif ev and ev.get("m") is not None:
        r.update(m=ev["m"], net=ev.get("net"), conf=ev["conf"], dec=ev["dec"], note=str(ev.get("risk") or "")[:160])
    elif not x["nmck"]:
        r.update(mtxt=f'ориентир {CAT_M.get(x["cat"], "—")}', conf="нет НМЦК", dec="Запросить спецификацию → КП", note="коммерческий запрос / без НМЦК — цена = наше КП")
    else:
        r.update(mtxt=f'ориентир {CAT_M.get(x["cat"], "—")}', conf="не посчитано", dec="Посчитать по позициям", note="")
    if "запрос цен" in x["method"].lower() or "мониторинг" in x["method"].lower():
        r.update(dec="КП без обязательств (сбор цен)", note="не закупка — участие в формировании НМЦК")
    if NOTPROF.search(x["name"]): r.update(dec="Не наш профиль")
    if r["m"] is not None and not r["dec"].startswith(("Не ", "Нет")):
        if r["net"] is not None and r["net"] <= 0 or r["m"] < 0.1: r["dec"] = "Не идём — нет маржи"
        elif r["conf"].startswith(("оценка", "низкая")): r["dec"] = r["dec"] or "Запросить спецификацию / КП"
    rows.append(r)


def rank(r):
    if r["dec"].startswith(("ИДЁМ",)): return 0
    if r["m"] is not None and r["m"] >= 0.25 and not r["dec"].startswith(("Не", "Нет")): return 1
    if r["m"] is not None and not r["dec"].startswith(("Не", "Нет")): return 2
    if r["dec"].startswith(("Запросить", "Посчитать", "КП")): return 3
    return 4


rows.sort(key=lambda r: (rank(r), -(r["net"] or 0), -(r["nmck"] or 0)))
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'
FILL = {0: "C6EFCE", 1: "E2EFDA", 2: "FFEB9C", 3: "DDEBF7", 4: "F2F2F2"}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "С 25.10"
ws.append([f"Все 9 выгрузок 19.10 — закупки с подачей с 25.10.2026, без аукционов и работ. Всего {len(rows)}"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Маржа = 1 − себестоимость/НМЦК (расчёт по позициям с ценами 2026) или «ориентир» по категории, если позиций/НМЦК нет. Прибыль — после УСН 15% при ожидаемом снижении. Время — МСК."])
H = ["№", "Решение", "Подача до", "Категория", "Номер", "Заказчик", "Предмет", "Закон / способ", "Площадка", "НМЦК, ₽", "Маржа при НМЦК", "Прибыль после УСН, ₽", "Надёжность", "Комментарий", "Ссылка ЕИС", "Ссылка площадки"]
ws.append(H)
for c in ws[3]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
sh = lambda t, n: (t or "") if len(t or "") <= n else t[:n - 1] + "…"
for i, r in enumerate(rows, 1):
    lk = L.get(r["num"], {}); href = lk.get("href") if "zakupki" in str(lk.get("href")) else ""
    if not href and re.fullmatch(r"3\d{10}", r["num"]): href = f"https://zakupki.gov.ru/epz/order/notice/notice223/common-info.html?regNumber={r['num']}"
    if not href and re.fullmatch(r"0\d{18}", r["num"]): href = f"https://zakupki.gov.ru/epz/order/extendedsearch/results.html?searchString={r['num']}"
    etpu = lk.get("etp_url") or ""
    if etpu and not str(etpu).startswith("http"): etpu = "https://" + etpu
    ws.append([i, r["dec"] or "—", r["end"].strftime("%d.%m.%Y %H:%M"), r["cat"], r["num"], sh(r["cust"], 50), sh(r["subj"], 90), sh(r["law"], 34), r["etp"], r["nmck"],
               r["m"] if r["m"] is not None else r["mtxt"], r["net"], r["conf"], sh(r.get("note"), 150), href, etpu])
    X = ws.max_row; ws.cell(X, 5).number_format = "@"; ws.cell(X, 10).number_format = M; ws.cell(X, 12).number_format = M
    if r["m"] is not None: ws.cell(X, 11).number_format = "0%"
    ws.cell(X, 2).fill = PatternFill("solid", fgColor=FILL[rank(r)])
    for j in (15, 16):
        v = ws.cell(X, j).value
        if v: ws.cell(X, j).hyperlink = v; ws.cell(X, j).font = Font(color="0563C1", underline="single")
    for c in ws[X]: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([4, 26, 15, 14, 20, 30, 42, 24, 16, 14, 12, 14, 12, 44, 34, 34], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "F4"; ws.auto_filter.ref = f"A3:P{ws.max_row}"
tidy(wb, row_h=44)
wb.save("../FROM25_MARGIN_2110.xlsx")
import collections
print(len(rows), collections.Counter(rank(r) for r in rows))
for r in rows:
    if rank(r) <= 2: print(rank(r), r["num"], r["end"].strftime("%d.%m"), r["nmck"], round((r["m"] or 0) * 100), round(r["net"] or 0), r["dec"][:35])
