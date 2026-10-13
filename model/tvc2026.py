"""Наши лоты против себестоимости по ценам 2026 (проверенные ссылки). Пишет TENDERS_VS_COST_1410.xlsx и data/tvc2026.json."""
import json, importlib.util, sys, io, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
import ref2026 as R
def load(n,f):
    s=importlib.util.spec_from_file_location(n,f); m=importlib.util.module_from_spec(s); o=sys.stdout; sys.stdout=io.StringIO(); s.loader.exec_module(m); sys.stdout=o; return m
ss=load("ss","salt_strategy_1410.py"); P,be=ss.P,ss.be
def group(item,north=False):
    s=item.lower(); h=s.split()[0] if s.split() else ""
    cold=("зим" in s or "утепл" in s or "пониж" in s)
    if "нефт" in s: return "Костюм от нефтепродуктов зимний" if "зим" in s else "Костюм от нефтепродуктов летний"
    if "сигнал" in s and not h.startswith("жилет"): return "Костюм сигнальный зимний" if cold else "Костюм сигнальный летний"
    if "кислот" in s or "кщс" in s: return "Костюм КЩС (кислоты/щелочи)"
    if "энц" in s: return "Костюм противоэнцефалитный"
    if h.startswith("плащ"): return "Плащ влагозащитный"
    if "дуг" in s: return "Костюм от электродуги зимний" if "зим" in s else "Костюм от электродуги летний"
    if "сварщ" in s: return "Костюм сварщика"
    if h.startswith("костюм"):
        if cold: return "Костюм зимний северный (IV/особый пояс)" if north else "Костюм зимний рабочий (I–III пояс)"
        return "Костюм летний спасателя/МЧС" if "спасат" in s else "Костюм летний рабочий"
    if h.startswith("зимн") and ("одежд" in s or "форм" in s): return "Костюм зимний северный (IV/особый пояс)" if north else "Костюм зимний рабочий (I–III пояс)"
    if h.startswith("комбинезон"): return "Комбинезон зимний" if cold else "Комбинезон летний"
    if h.startswith("полукомбинезон") or (h.startswith("брюки") and cold): return "Полукомбинезон / брюки зимние"
    if h.startswith("куртк"): return "Куртка утепленная"
    if h.startswith("жилет"): return "Жилет сигнальный" if "сигнал" in s else "Жилет утепленный"
    if h.startswith("футболк"): return "Футболка"
    if h.startswith("халат"): return "Халат рабочий"
    if h.startswith("термобел"): return "Термобельё (комплект)"
    if h.startswith("ботин"): return "Ботинки утепленные" if cold else "Ботинки кожаные с МП (летние)"
    if h.startswith("полусапог"): return "Полусапоги"
    if h.startswith("сапог"): return "Сапоги утепленные / валенки" if cold else "Сапоги ПВХ"
    if h.startswith("перчат"): return "Перчатки термостойкие/краги" if "термо" in s else "Перчатки х/б с ПВХ (пара)"
    if h.startswith("подшлем"): return "Подшлемник"
    if h.startswith("каска"): return "Каска"
    if h.startswith("щиток"): return "Щиток защитный лицевой"
    return None
DLV=dict(центр=15,регион=35,север=120)
def srcinfo(g):
    S=R.sources(g,3)
    txt=("ваша себестоимость; рынок 2026: " if g in R.CLIENT else "медиана опт/прайс 2026: ")+"; ".join(f'{x["site"]} {x["price"]:,.0f}'.replace(","," ") for x in S)
    return txt,[x["url"] for x in S]
deep={}
for f in ("out_clothes_nmck.json","out_clothes_comm.json"):
    for x in json.load(open("data/deep/"+f)): deep[x["num"]]=x
CL=[("0358200022926000030","Ростов, форма спасателей","19.10 09:00","регион",0.17),
    ("32616431081","ФБУ ЦСМ Башкортостан, спецодежда 18 поз.","19.10 08:00","регион",0.10),
    ("32616453117","Салаватводоканал, спецодежда","21.10 09:00","регион",0.10),
    ("32616452870","Салаватводоканал, 2-й лот","21.10 09:00","регион",0.10),
    ("32616443256","Якутия ГПС, форма (подача 15.10!)","15.10 05:00","север",0.15),
    ("4634447","Биннофарм, пул поставщиков (корзина — оценка)","26.10 12:00","центр",0.07),
    ("372-446","Благо-Армавир (корзина — оценка)","19.10 17:00","регион",0.07)]
OUT={}; rows=[]
for num,name,end,reg,dsc in CL:
    x=deep[num]; north=(reg=="север"); pos=[]; C=0; sn_u=sc_u=0
    for p in x["positions"]:
        g=group(p["item"],north)
        if not g or R.base(g) is None: continue
        cu=R.base(g)+R.weight(g)*DLV[reg]; txt,U=srcinfo(g)
        q=p.get("qty"); nu=p.get("nmck_unit") if num!="32616443256" else None
        if num in("4634447","372-446"): nu=None
        pos.append(dict(item=p["item"],qty=q,unit=p.get("unit") or "шт",nmck_unit=nu,cost_unit=cu,group=g,src=txt,urls=U))
        if q: C+=cu*q
        if nu: sn_u+=nu; sc_u+=cu
    nm=x.get("nmck")
    if num in("32616453117","32616452870"): m=(sn_u-sc_u)/sn_u; C=nm*(1-m)
    if num in("4634447","372-446"):
        kp=C/(1-0.25); nm=None; pe=P(kp,C,dsc); m=0.25; b=be(kp,C)
    else:
        m=(nm-C)/nm; pe=P(nm,C,dsc); b=be(nm,C)
    OUT[num]=dict(name=name,positions=pos,C=C,m=m,pe=pe,be=b,nm=nm); rows.append(("Одежда",name,num,end,nm,C,m,dsc,pe,b))
for L in ss.LOTS:
    num=L[1]; cpt,txt,U=R.SALT2026[num]; nm,t=L[3],L[4]; C=cpt*t
    OUT[num]=dict(name=L[0],positions=[dict(item="Соль техническая",qty=t,unit="т",nmck_unit=nm/t,cost_unit=cpt,group="Соль",src=txt,urls=U)],C=C,m=(nm-C)/nm,pe=P(nm,C,0.15),be=be(nm,C),nm=nm)
    rows.append(("Соль",L[0],num,L[2],nm,C,(nm-C)/nm,0.15,P(nm,C,0.15),be(nm,C)))
json.dump(OUT,open("data/tvc2026.json","w"),ensure_ascii=False,indent=1)
json.dump([dict(num=r[2],C=r[5],m=r[6],pe=r[8],be=r[9]) for r in rows],open("data/tvc_rows.json","w"),ensure_ascii=False)
# xlsx
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); S=wb.active; S.title="Сводка по лотам"
S.append(["Наши лоты против себестоимости по ценам 2026 (проверенные прайсы и каталоги, 14.10.2026)"]); S["A1"].font=Font(bold=True,size=13)
S.append(["Одежда/СИЗ — медиана опт/прайс 2026 (Тракт 01.10.2026, Спецобъединение 14.10.2026, Техноавиа 09.10.2026, Фабрикант 07.10.2026, ivuniforma и др.) + доставка; соль — проверенные цены 2026 с доставкой."])
H=["Категория","Лот","№","Подача","НМЦК, ₽","Себестоимость с доставкой, ₽","Маржинальность","Ожид. снижение","Чистая прибыль, ₽","Безубыток до снижения"]
S.append(H)
for c in S[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
rows.sort(key=lambda r:-r[8])
for r in rows:
    S.append(list(r)); Rr=S.max_row; S.cell(Rr,3).number_format="@"
    for j in (5,6,9): S.cell(Rr,j).number_format=MONEY
    for j in (7,8,10): S.cell(Rr,j).number_format="0%"
    S.cell(Rr,9).fill=PatternFill("solid",fgColor="C6EFCE" if r[8]>300000 else ("FFEB9C" if r[8]>0 else "FFC7CE"))
for i,w in enumerate([10,44,21,12,15,16,11,10,16,11],1): S.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
W=wb.create_sheet("Позиции")
W.append(["Лот","Позиция","Кол-во","Ед.","Цена за ед. в НМЦК, ₽","Себестоимость за ед. с доставкой, ₽","Маржа на ед.","Группа","Источники цены 2026","Ссылка"])
for c in W[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
for num,o in OUT.items():
    for p in o["positions"]:
        W.append([o["name"],p["item"],p["qty"],p["unit"],p["nmck_unit"],p["cost_unit"],(p["nmck_unit"]-p["cost_unit"])/p["nmck_unit"] if p["nmck_unit"] else None,p["group"],p["src"],p["urls"][0] if p["urls"] else ""])
        Rr=W.max_row
        for j in (5,6): W.cell(Rr,j).number_format=MONEY
        W.cell(Rr,7).number_format="0%"
        if p["urls"]: W.cell(Rr,10).hyperlink=p["urls"][0]; W.cell(Rr,10).font=Font(color="0563C1",underline="single")
for i,w in enumerate([30,52,8,8,14,15,10,30,60,40],1): W.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
from tidy import tidy; tidy(wb,row_h=30)
wb.save("../TENDERS_VS_COST_1410.xlsx")
for r in rows: print(r[0],r[1][:32],round(r[4] or 0),round(r[5]),f"{r[6]:.0%}",round(r[8]))
