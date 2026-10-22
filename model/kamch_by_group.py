"""Камчатскэнерго 32616346943 — по листу на каждую группу арматуры: заводы для КП сверху, позиции с количеством и проверенной ценой ниже."""
import json, re, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from tidy import tidy
from kamch_groups import group

Q = json.load(open("data/kamch/kam_qty.json"))
S = {(re.search(r"Таблица\s*(1\.\d)", x["spec"]).group(1), x["n"]): x for x in json.load(open("data/kamch/kam_spec.json"))["sheet2"]}
V = {}
for b in (1, 2, 3, 4):
    for x in json.load(open(f"data/kamch/verify/ver_kam_b{b}.json")): V[(str(x["table"]), int(x["n"]))] = x
SUP = {x["group"]: x["suppliers"] for x in json.load(open("data/kamch/suppliers.json"))}
SHORT = {"Краны шаровые стальные": "Краны стальные", "Задвижки стальные": "Задвижки стальные", "Затворы дисковые": "Затворы дисковые",
         "Клапаны регулирующие и регуляторы": "Клапаны регулирующие", "Задвижки нержавеющие": "Задвижки нержавеющие", "Клапаны запорные / вентили": "Клапаны запорные",
         "Клапаны обратные": "Клапаны обратные", "Задвижки чугунные": "Задвижки чугунные", "Клапаны предохранительные": "Клапаны предохранительные",
         "Краны шаровые латунные": "Краны латунные", "Краны шаровые нержавеющие": "Краны нержавеющие", "Электроприводы и редукторы": "Электроприводы",
         "Краны трёхходовые/манометровые": "Краны манометровые", "Запчасти, уплотнения, рыбозащита, прочее": "Рыбозащита и прочее"}
def big(g):
    for k_ in ("Краны", "Задвижки", "Затворы", "Клапаны"):
        if g.startswith(k_): return k_
    return "Приводы и прочее"
G = {}; SUPB = {}
for q in Q:
    k = (q["table"], q["n"]); sg = group(q["name"]); g = big(sg); v = V.get(k, {}); s = S[k]
    G.setdefault(g, []).append(dict(t=q["table"], n=q["n"], sub=SHORT.get(sg.split(" (")[0], sg[:28]), name=q["name"], qty=q["qty"], unit=q["unit"], dn=s.get("dn"), pn=s.get("pn"), mat=s.get("mat"), conn=s.get("conn"),
                                    drive=s.get("drive"), price=v.get("rec_price"), verdict=v.get("verdict", ""), pp=str(v.get("pp1875") or "")))
for sg, lst in SUP.items():
    g = big(sg); d = SUPB.setdefault(g, {})
    for sp in lst:
        m_ = re.search(r"([a-z0-9\-]+\.(?:ru|com|рф))", str(sp.get("site") or "") + " " + str(sp.get("email") or ""))
        key = m_.group(1) if m_ else re.sub(r"[^А-Яа-яA-Za-z]", "", sp.get("company", ""))[:12]
        if key in d:
            d[key]["subs"].append(SHORT.get(sg.split(" (")[0], sg[:28])); d[key]["covers"] = str(d[key].get("covers") or "")[:80] + " | " + str(sp.get("covers") or "")[:80]
        else: d[key] = dict(sp, subs=[SHORT.get(sg.split(" (")[0], sg[:28])])
SUP = {g: list(d.values()) for g, d in SUPB.items()}
thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'
def head(ws):
    for c in ws[ws.max_row]:
        if c.value is not None: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
def box(ws, r0, r1):
    for row in ws.iter_rows(min_row=r0, max_row=r1):
        for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
sh = lambda t, n: (str(t or "")) if len(str(t or "")) <= n else str(t)[:n - 1] + "…"
order = sorted(G, key=lambda g: -sum((p["price"] or 0) * p["qty"] for p in G[g]))
wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Оглавление"
sv.append(["Камчатскэнерго 32616346943 — позиции по группам и заводы для запроса КП. Подача 29.10.2026 10:00 МСК"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["В каждом запросе КП: цена за ед. с НДС, срок изготовления и доставки до Петропавловска-Камчатского, номер записи в реестре Минпромторга (ПП 719), сертификаты ТР ТС 032/2013 и 010/2011"])
sv.append(["№", "Группа (лист)", "Позиций", "Единиц", "Сумма по проверенным ценам, ₽", "Кому запрашивать КП (первые)"]); head(sv)
for i, g in enumerate(order, 1):
    ps = G[g]; tot = sum((p["price"] or 0) * p["qty"] for p in ps)
    sv.append([i, g, len(ps), sum(p["qty"] for p in ps), tot, "; ".join(s.get("company", "")[:40] for s in SUP.get(g, [])[:5])])
    sv.cell(sv.max_row, 5).number_format = M
    sv.cell(sv.max_row, 2).hyperlink = f"#'{g}'!A1"; sv.cell(sv.max_row, 2).font = Font(color="0563C1", underline="single")
sv.append(["", "ИТОГО", sum(len(G[g]) for g in G), sum(p["qty"] for g in G for p in G[g]), sum((p["price"] or 0) * p["qty"] for g in G for p in G[g])])
sv.cell(sv.max_row, 5).number_format = M
for c in sv[sv.max_row]: c.font = Font(bold=True)
box(sv, 3, sv.max_row)
for i, w in enumerate([4, 26, 9, 9, 18, 90], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

for g in order:
    ws = wb.create_sheet(g)
    ps = G[g]; tot = sum((p["price"] or 0) * p["qty"] for p in ps)
    ws.append([f"{g} — {len(ps)} позиций, {sum(p['qty'] for p in ps):,.0f} ед., ≈ {tot:,.0f} ₽ по проверенным ценам 2026".replace(",", " ")]); ws["A1"].font = Font(bold=True, size=12)
    ws.append(["ЗАВОДЫ / ПОСТАВЩИКИ ДЛЯ ЗАПРОСА КП"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["№", "Поставщик", "Что закрывает (подгруппы)", "", "", "Телефон", "E-mail / сайт", "Город", "Реестр МПТ", "Результат звонка / цена КП / срок", ""]); head(ws)
    ws.merge_cells(start_row=ws.max_row, start_column=3, end_row=ws.max_row, end_column=5); ws.merge_cells(start_row=ws.max_row, start_column=10, end_row=ws.max_row, end_column=11)
    r0 = ws.max_row
    for i, s in enumerate(SUP.get(g, []), 1):
        ws.append([i, s.get("company"), sh(", ".join(s["subs"]) + ": " + str(s.get("covers") or ""), 140), "", "", (s.get("phone") + (f"\n[{s['checked']}]" if s.get("checked") else "")), f'{s.get("email") or ""}\n{s.get("site") or ""}', s.get("city"), sh(str(s.get("registry") or "").split(";")[0].split("(")[0], 30), "", ""])
        ws.merge_cells(start_row=ws.max_row, start_column=3, end_row=ws.max_row, end_column=5); ws.merge_cells(start_row=ws.max_row, start_column=10, end_row=ws.max_row, end_column=11)
    box(ws, r0, ws.max_row)
    ws.append([]); ws.append(["ПОЗИЦИИ ДЛЯ ЗАПРОСА"]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.append(["№", "Таблица ТТ / №", "Подгруппа · Позиция", "DN", "PN, МПа", "Материал корпуса", "Присоединение / привод", "Кол-во", "Проверенная цена, ₽/ед.", "Сумма, ₽", "Нацрежим / замечание"]); head(ws)
    r1 = ws.max_row
    for i, p in enumerate(sorted(ps, key=lambda p: (p["sub"], p["t"], p["n"])), 1):
        ws.append([i, f'{p["t"]} / {p["n"]}', sh(p["sub"] + " · " + p["name"], 100), p["dn"], p["pn"], sh(p["mat"], 30), sh(f'{p["conn"] or ""}; {p["drive"] or ""}', 40), p["qty"], p["price"],
                   (p["price"] or 0) * p["qty"], sh(p["pp"], 40)])
        ws.cell(ws.max_row, 9).number_format = M; ws.cell(ws.max_row, 10).number_format = M
        if "импорт" in p["pp"].lower(): ws.cell(ws.max_row, 11).fill = PatternFill("solid", fgColor="FCE4D6")
    ws.append(["", "", "ИТОГО", "", "", "", "", sum(p["qty"] for p in ps), "", tot]); ws.cell(ws.max_row, 10).number_format = M
    for c in ws[ws.max_row]: c.font = Font(bold=True)
    box(ws, r1, ws.max_row)
    for i, w in enumerate([4, 30, 46, 7, 8, 22, 46, 14, 16, 14, 30], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=36)
for ws in wb.worksheets[1:]:
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, str) and len(v) > 10:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=11); ws.cell(r, 1).alignment = Alignment(wrap_text=False, vertical="center"); ws.row_dimensions[r].height = 22
        if ws.cell(r, 7).value and "\n" in str(ws.cell(r, 7).value): ws.row_dimensions[r].height = 52
wb.save("../KAMCHATKA_BY_GROUP_KP.xlsx")
for g in order: print(g, len(G[g]), round(sum((p["price"] or 0) * p["qty"] for p in G[g])), len(SUP.get(g, [])))
