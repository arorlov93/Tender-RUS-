"""Таблица для обзвона поставщиков (лоты 1–9): контакты, что спросить, ориентир цены, ссылки."""
import json, collections, importlib.util, sys, io, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
import ref2026 as R
def load(n,f):
    s=importlib.util.spec_from_file_location(n,f); m=importlib.util.module_from_spec(s); o=sys.stdout; sys.stdout=io.StringIO(); s.loader.exec_module(m); sys.stdout=o; return m
tt=load("tt","tenderer_1_9.py")
TV=json.load(open("data/tvc2026.json"))
LC={x["url"].split("//")[1].split("/")[0].replace("www.",""):x for x in json.load(open("data/link_check.json"))}
lots=[(i+1,r[1],r[2],r[0]) for i,r in enumerate(tt.ROWS)]   # (№, номер, заказчик, категория)
need=collections.defaultdict(list)   # group -> [(lot#, qty, unit)]
for k,n,cust,cat in lots:
    if n in TV:
        for p in TV[n]["positions"]: need[p["group"]].append((k,p["qty"],p["unit"]))
C={}   # site -> record
def add(x,what,lotlabel,price_txt):
    site=x.get("site") or x["url"].split("//")[1].split("/")[0].replace("www.","")
    r=C.setdefault(site,dict(company=x.get("company") or site,city=x.get("city",""),phone=x.get("phone",""),email=x.get("email",""),url=x["url"],ev=x.get("date_evidence",""),what=[],lots=set(),prices=[]))
    if what not in r["what"]: r["what"].append(what)
    r["lots"].add(lotlabel); r["prices"].append(price_txt)
    if x.get("date_evidence","").startswith("A"): r["ev"]=x["date_evidence"]; r["url"]=x["url"]
for g,lst in need.items():
    if g=="Соль": continue
    qty=sum(q for _,q,_ in lst if q) ; lotsg=sorted({k for k,_,_ in lst})
    what=f"{g}" + (f" — ~{qty:g} шт" if qty else " — объём по заявкам")
    for x in R.sources(g,4): add(x,what,", ".join(map(str,lotsg)),f'{g}: {x["price"]:,.0f} ₽'.replace(","," "))
# соль — по лотам
SR={x["url"]:x for x in R.SALTROWS}
for k,n,cust,cat in lots:
    if n in R.SALT2026:
        cpt,txt,U=R.SALT2026[n]; t=TV[n]["positions"][0]["qty"]
        for u in U:
            x=SR.get(u) or next((y for y in R.SALTROWS if y["url"].split("?")[0].rstrip("/")==u.rstrip("/") or y.get("site","") in u),None)
            if x: add(x,f"Соль техническая {t:g} т с доставкой ({cust[:30]})",str(k),f'соль {x["price"]:,.0f}–{(x.get("price_max") or x["price"]):,.0f} ₽/т'.replace(","," "))
# ПГР (лот Мосгортур)
for x in [y for y in R.SALTROWS if y["group"]=="ПГР"]:
    k=next(k for k,n,_,_ in lots if n=="32616449090")
    add(x,"ПГР твёрдый −20…−30 °C, 23,6 т, с записью ГИСП",str(k),f'ПГР {x["price"]:,.0f} ₽/т'.replace(","," "))
# заводы соли — КП франко-вагон
for x in [y for y in R.SALTROWS if y["group"].startswith("Соль: цена завода")]:
    add(x,"КП франко-вагон / с доставкой: соль 400–1 000 т (ДГХ, Новосибирск, Струнино)","1, 5, 8",f'{x["price"]:,.0f}–{(x.get("price_max") or x["price"]):,.0f} {x.get("currency","₽")}/т'.replace(","," "))
# арматура
for site,what in (("specarmatura.ru","Запорная арматура: 19 позиций Т Плюс (ВТ-5, ВИГ, 15с27нж, 15с22нж, 19лс53нж…) — КП и аналоги"),("dn.ru","Запорная арматура (краны LD, задвижки) — КП по спецификации Т Плюс")):
    x=LC.get(site)
    if x: add(dict(site=site,company=site,phone=x["phone"],email=x["email"],url=x["url"],date_evidence="B: © 2026, живой каталог"),what,str(next(k for k,n,_,_ in lots if n=="32616442935")),"см. сайт")
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Обзвон поставщиков"
ws.append(["Обзвон поставщиков на 15.10.2026 — лоты 1–9 (для внутреннего пользования)"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["Лоты: "+"; ".join(f"{k} — {cat} ({cust[:28]})" for k,n,cust,cat in lots)])
H=["№","Приоритет","Компания","Город","Телефон","E-mail","Сайт / прайс (ссылка)","Что спросить (товар — объём)","Для лотов №","Цена на сайте (ориентир)","Подтверждение даты цены","Дозвонился (да/нет)","Цена по звонку, ₽","НДС / доставка / сроки","Комментарий"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
def prio(r):
    s=0
    if r["ev"].startswith("A"): s+=3
    s+=len(r["what"])
    if any(l in r["lots"] for l in ("1","5","8")) or any("Соль" in w for w in r["what"]): s+=2
    return s
recs=sorted(C.items(),key=lambda kv:-prio(kv[1]))
for i,(site,r) in enumerate(recs,1):
    p=prio(r); pl="высокий" if p>=6 else ("средний" if p>=3 else "низкий")
    ws.append([i,pl,r["company"],r["city"],r["phone"],r["email"],site,"\n".join(r["what"]),", ".join(sorted(set(", ".join(r["lots"]).split(", ")),key=lambda z:int(z) if z.isdigit() else 99)),"\n".join(dict.fromkeys(r["prices"])),r["ev"],"","","",""])
    R_=ws.max_row; ws.cell(R_,7).hyperlink=r["url"]; ws.cell(R_,7).font=Font(color="0563C1",underline="single")
    ws.cell(R_,2).fill=PatternFill("solid",fgColor={"высокий":"C6EFCE","средний":"FFEB9C","низкий":"F2F2F2"}[pl])
    for cc in ws[R_]: cc.border=Border(top=thin,bottom=thin,left=thin,right=thin)
for i,w in enumerate([4,10,30,14,20,26,22,52,10,34,34,12,13,20,24],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="D4"; ws.auto_filter.ref=f"A3:O{ws.max_row}"
from tidy import tidy; tidy(wb,row_h=60)
wb.save("../CALL_LIST_1510.xlsx")
for i,(site,r) in enumerate(recs,1): print(i,site,r["phone"],r["email"],len(r["what"]),sorted(r["lots"]))
