"""Одна итоговая таблица: все лоты с подачей с 16.10 — себестоимость, маржа, прибыль."""
import importlib.util, sys, io, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
spec=importlib.util.spec_from_file_location("ft","final_top_1610.py"); ft=importlib.util.module_from_spec(spec); o=sys.stdout; sys.stdout=io.StringIO(); spec.loader.exec_module(ft); sys.stdout=o
E=ft.E
BASIS={"Соль":"средняя рыночная с доставкой (завод вагоном / склад авто)","Спецодежда":"средние цены фабрик РФ + доставка ТК","Форма спасателей":"средние цены фабрик РФ + доставка ТК",
       "ПГР":"прайс поставщика ПГР","Пробирки":"UNIVAC/МиниМед ~9–9,6 ₽/шт","Арматура":"оценка, цен по позициям нет","Трубы ППУ":"оценка по аналогам","Спецодежда (коммерч.)":"типовая корзина, без НМЦК"}
def basis(cat): return next((v for k,v in BASIS.items() if cat.startswith(k)),"")
rows=[(r,"обычные") for r in ft.MAIN]+[(r,"ЖЁСТКИЕ: "+r[11]) for r in ft.BIG]
rows.sort(key=lambda x:(x[1]!="обычные",-x[0][8]))
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Итог"
ws.append(["Итоговая таблица лотов (подача с 16.10.2026) — себестоимость, маржинальность, прибыль"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["Прибыль — чистая после УСН 15% при ожидаемом снижении; маржа = (НМЦК − себестоимость) / НМЦК. Зелёный ≥40%, жёлтый 15–40%, красный <15%."])
H=["№","Категория","№ закупки","Ссылка","Заказчик","Способ","Подача (МСК)","НМЦК, ₽","Себестоимость с доставкой, ₽","Маржинальность","Валовая маржа, ₽","Ожид. снижение","Чистая прибыль, ₽","Безубыток до снижения","Обеспечение заявки / исполнения","Условия","Основа себестоимости","Что важно","Достоверность"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
DISC={"Соль":.15,"ПГР":.15,"Пробирки":.25,"Форма":.17,"Спецодежда (коммерч":.07,"Спецодежда":.10,"Арматура":.14,"Трубы":.10}
for i,(r,cond) in enumerate(rows,1):
    cat,n,cust,meth,end,nm,c,m,pe,be,sec,note,conf=r
    d=next((v for k,v in DISC.items() if cat.startswith(k)),None)
    ws.append([i,cat,n,"открыть",cust,meth,end,nm,c,m,(nm-c) if nm and c else None,d,pe,be,sec,"обычные" if cond=="обычные" else "жёсткие",basis(cat),note,conf])
    R=ws.max_row; ws.cell(R,4).hyperlink=E(n); ws.cell(R,4).font=Font(color="0563C1",underline="single"); ws.cell(R,3).number_format="@"
    for k in (8,9,11,13): ws.cell(R,k).number_format=MONEY
    for k in (10,12,14): ws.cell(R,k).number_format="0%"
    if m is not None: ws.cell(R,10).fill=PatternFill("solid",fgColor="C6EFCE" if m>=.4 else ("FFEB9C" if m>=.15 else "FFC7CE"))
    ws.cell(R,13).fill=PatternFill("solid",fgColor="C6EFCE" if pe>300000 else ("FFEB9C" if pe>0 else "FFC7CE"))
    if cond!="обычные":
        for k in (2,16): ws.cell(R,k).font=Font(bold=True,color="C00000")
    for cc in ws[R]: cc.border=Border(top=thin,bottom=thin,left=thin,right=thin)
def total(label,flt):
    s=[r for r,c in rows if flt(c)]
    nm=sum(r[5] or 0 for r in s); c=sum(r[6] or 0 for r in s if r[5]); pe=sum(r[8] for r in s)
    ws.append(["","","","",label,"","",nm,c,(nm-c)/nm if nm else None,nm-c,None,pe]); R=ws.max_row
    for k in (8,9,11,13): ws.cell(R,k).number_format=MONEY
    ws.cell(R,10).number_format="0%"
    for cc in ws[R]: cc.font=Font(bold=True)
total("ИТОГО — обычные условия",lambda c:c=="обычные")
total("ИТОГО — с крупными (жёсткие условия)",lambda c:True)
for i,w in enumerate([4,18,21,9,32,22,12,15,16,11,15,9,16,10,15,10,30,44,11],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="D4"; ws.auto_filter.ref=f"A3:S{ws.max_row-2}"
from tidy import tidy; tidy(wb)
wb.save("../ITOG_LOTS_1410.xlsx")
for r,c in rows: print(r[0][:12],r[1],round(r[5] or 0),r[7] and f"{r[7]:.0%}",round(r[8]),c[:10])
