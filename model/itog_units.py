"""Одна таблица без аукционов: лоты и позиции — цена за ед. в НМЦК, себестоимость за ед. с доставкой, маржа."""
import json, importlib.util, sys, io, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
def load(n,f):
    s=importlib.util.spec_from_file_location(n,f); m=importlib.util.module_from_spec(s); o=sys.stdout; sys.stdout=io.StringIO(); s.loader.exec_module(m); sys.stdout=o; return m
it=load("it","itog_one.py")   # строки лотов без аукционов: it.rows
E=it.E
TVC=openpyxl.load_workbook("../TENDERS_VS_COST_1410.xlsx")["Одежда — по позициям"]
pos_tvc={}
LOTNAME={"ФБУ ЦСМ Башкортостан, спецодежда 18 поз.":"32616431081","Салаватводоканал, спецодежда":"32616453117"}
for r in TVC.iter_rows(min_row=2,values_only=True):
    n=LOTNAME.get(r[0])
    if n: pos_tvc.setdefault(n,[]).append((r[1],r[2],"шт/компл",r[3],r[4],r[8]))
deep={}
for f in ("out_clothes_nmck.json","out_clothes_comm.json"):
    for x in json.load(open("data/deep/"+f)): deep[x["num"]]=x
def dpos(n,dl=35):
    out=[]
    for p in deep[n]["positions"]:
        c=p.get("cost_unit")
        out.append((p["item"],p.get("qty"),p.get("unit") or "шт",p.get("nmck_unit"),(c+dl*1.5) if c else None,p.get("cost_src","")))
    return out
SALT={"32616451106":(400,4870),"0851300006226000106":(400,4310),"0128600006826000407":(150,5060)}
SRC_SALT="средняя рыночная с доставкой: завод вагоном / склад авто (SALT_STRATEGY_1410)"
def positions(r):
    n=r[1]; nm=r[5]
    if n in SALT: t,c=SALT[n]; return [("Соль техническая",t,"т",nm/t,c,SRC_SALT)]
    if n=="32616449090": return [("ПГР твёрдый (реагент)",23.6,"т",54490,553080/23.6,"прайс «А-Стандарт» −25 °C 22 800 ₽/т + доставка")]
    if n in pos_tvc: return pos_tvc[n]
    if n=="32616442935": return [("Запорная арматура стальная, 19 поз. (среднее за ед.)",127,"шт",nm/127,r[6]/127,"оценка по дилерским прайсам (цен по позициям нет)")]
    if n=="32616450774": return [("Отводы в ППУ-изоляции Ду108–530 (среднее за ед.)",9,"шт",nm/9,r[6]/9,"оценка по аналогам")]
    if n in deep: return dpos(n)
    return []
import re
cr=load("cr","cost_reference_1410.py")
SPECURL={"Тракт 6 731":["https://www.trakt.ru/product/12492/","https://www.spets.ru/products/spetsodezhda/odezhda_spetsialnaya/"],
 "Тракт 11 253":["https://www.trakt.ru/product/12720/","https://www.spets.ru/products/spetsodezhda/odezhda_spetsialnaya/"],
 "Спецобъединение «Рефлект-2»":["https://www.spets.ru/products/spetsodezhda/odezhda_spetsialnaya/"],
 "Тракт, костюм КЩС":["https://www.trakt.ru/product/133819/"],
 "Фабрикант 975":["https://fabrikant-textile.ru/catalog/spetsodezhda/","https://ivsoteks.ru/specodejda"],
 "Тракт 971":["https://www.trakt.ru/product/11992/","https://fabrikant-textile.ru/catalog/spetsodezhda/"],
 "SALT":["https://www.esolk.ru/produktsiya/kontsentrat-mineralnyy-galit/mkr-1-tonna/","https://nsmgr.ru/reagenty/tekhnicheskaya-sol/big-begi","https://sknerud.ru/reagent/sol-texnicheskaya/","https://vektor-e.ru/sol/"],
 "прайс «А-Стандарт»":["https://www.bicotender.ru/tender342335220.html"],
 "оценка по дилерским":["https://www.specarmatura.ru/catalog/zadvizhka_30s41nzh_ru_16/zadvizhka_30s41nzh_du_100/","https://dn.ru/sharovyi-kran/teplosnabzhenie/flantcevyi-kran/standartnyi/ld/st20/dn100"],
 "price_Technoavia":["https://files.technoavia.ru/files/price/11/price_Technoavia_moscow_2025_06_27_08_46_34.pdf"]}
def urls(src):
    src=src or ""
    if src.startswith("средняя рыночная с доставкой"): return SPECURL["SALT"]
    found=re.findall(r"https?://\S+",src)
    if found: return [u.rstrip(").,;") for u in found]
    m=re.search(r"справочник: (.+)$",src)
    if m and m.group(1).strip() in cr.G:
        rows=sorted(cr.G[m.group(1).strip()][1],key=lambda r:(r[4]!="опт",r[2]))
        return list(dict.fromkeys(r[5] for r in rows))[:3]
    for k,v in SPECURL.items():
        if k in src: return v
    if "правило клиента" in src or "клиента" in src: return [r[5] for r in sorted(cr.G["Костюм зимний рабочий (I–III пояс)"][1],key=lambda r:r[2])[:2]]
    return []
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Итог без аукционов"
ws.append(["Лоты БЕЗ АУКЦИОНОВ (подача с 16.10.2026): цена за единицу в НМЦК против нашей себестоимости за единицу с доставкой"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["Строка лота (жирная) — итог по лоту; строки ниже — позиции. Себестоимость за ед. — средняя по открытым ценам фабрик/оптовиков + доставка. Прибыль чистая (УСН 15%) при ожидаемом снижении."])
H=["№","Лот / позиция","№ закупки","Ссылка","Способ","Подача","Кол-во","Ед.","Цена за ед. в НМЦК, ₽","Себестоимость за ед. с доставкой, ₽","Маржа на ед.","НМЦК / сумма, ₽","Себестоимость / сумма, ₽","Маржинальность","Чистая прибыль, ₽","Безубыток до снижения","Обеспечение","Откуда себестоимость","Где цена (ссылка)","Ещё источники"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
LOTF=PatternFill("solid",fgColor="DDEBF7")
k=0
for r,cond in it.rows:
    k+=1; cat,n,cust,meth,end,nm,c,m,pe,be,sec,note,conf=r
    ps=positions(r)
    qty=ps[0][1] if len(ps)==1 else None; unit=ps[0][2] if len(ps)==1 else ""
    nu=ps[0][3] if len(ps)==1 else None; cu=ps[0][4] if len(ps)==1 else None
    title=f"{cat}: {cust}"+(" — КРУПНЫЙ, жёсткие условия" if cond!="обычные" else "")
    ws.append([k,title,n,"открыть",meth,end,qty,unit,nu,cu,(nu-cu)/nu if nu and cu else None,nm,c,m,pe,be,sec,ps[0][5] if len(ps)==1 else "по позициям ниже"])
    R=ws.max_row
    if len(ps)==1:
        U=urls(ps[0][5])
        if U: ws.cell(R,19).value=re.sub(r"^https?://(www\.)?","",U[0]).split("/")[0]; ws.cell(R,19).hyperlink=U[0]; ws.cell(R,20).value="\n".join(U[1:])
        else: ws.cell(R,19).value="нет открытой цены — оценка"
    ws.cell(R,4).hyperlink=E(n); ws.cell(R,4).font=Font(color="0563C1",underline="single"); ws.cell(R,3).number_format="@"
    for cc in ws[R]: cc.fill=LOTF; cc.border=Border(top=thin,bottom=thin,left=thin,right=thin)
    for j in (2,15): ws.cell(R,j).font=Font(bold=True)
    for j in (9,10,12,13,15): ws.cell(R,j).number_format=MONEY
    for j in (11,14,16): ws.cell(R,j).number_format="0%"
    if m is not None: ws.cell(R,14).fill=PatternFill("solid",fgColor="C6EFCE" if m>=.4 else ("FFEB9C" if m>=.15 else "FFC7CE"))
    if len(ps)>1:
        for it_,q,u,nu,cu,src in ps:
            ws.append(["","   "+it_,"","","","",q,u,nu,cu,(nu-cu)/nu if nu and cu else None,(nu*q if nu and q else None),(cu*q if cu and q else None),None,None,None,"",src])
            R=ws.max_row
            U=urls(src)
            if U: ws.cell(R,19).value=re.sub(r"^https?://(www\.)?","",U[0]).split("/")[0]; ws.cell(R,19).hyperlink=U[0]; ws.cell(R,20).value="\n".join(U[1:])
            elif cu: ws.cell(R,19).value="нет открытой цены — оценка"
            for j in (9,10,12,13): ws.cell(R,j).number_format=MONEY
            ws.cell(R,11).number_format="0%"
            if nu and cu: ws.cell(R,11).fill=PatternFill("solid",fgColor="C6EFCE" if (nu-cu)/nu>=.4 else ("FFEB9C" if nu>cu else "FFC7CE"))
            for cc in ws[R]: cc.border=Border(top=thin,bottom=thin,left=thin,right=thin)
L=[r for r,c in it.rows if c=="обычные"]
ws.append(["","ИТОГО — обычные условия","","","","","","","","","",sum(r[5] or 0 for r in L),sum(r[6] or 0 for r in L if r[5]),None,sum(r[8] for r in L)])
R=ws.max_row
for j in (12,13,15): ws.cell(R,j).number_format=MONEY
for cc in ws[R]: cc.font=Font(bold=True)
for i,w in enumerate([4,62,21,9,22,12,8,10,13,14,9,15,15,10,15,10,14,45,18,40],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="C4"
for row in ws.iter_rows(min_row=4):
    if row[18].hyperlink: row[18].font=Font(color="0563C1",underline="single")
from tidy import tidy; tidy(wb, row_h=30)
wb.save("../ITOG_LOTS_1410.xlsx")
print("ok", ws.max_row)
