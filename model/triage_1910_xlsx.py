"""TRIAGE_1910.xlsx: все прошедшие отбор лоты из выгрузок 19.10 (без аукционов, подача с 20.10), по листам-категориям."""
import json, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
d=json.load(open("data/triage_1910.json"))
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); S=wb.active; S.title="Сводка"
S.append(["Выгрузки 19.10.2026: отбор без аукционов и работ, подача с 20.10"]); S["A1"].font=Font(bold=True,size=13)
S.append(["Категория","Всего строк","Прошло отбор","С НМЦК","Сумма НМЦК прошедших, ₽","Отсеяно: аукцион","Отсеяно: работы/услуги","Отсеяно: срок"])
for c in S[2]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
cats=list(dict.fromkeys(x["cat"] for x in d))
for cat in cats:
    a=[x for x in d if x["cat"]==cat]; ok=[x for x in a if x["ok"]]
    S.append([cat,len(a),len(ok),sum(1 for x in ok if x["nmck"]),sum(x["nmck"] or 0 for x in ok),sum("аукцион" in x["why"] for x in a),sum("работы" in x["why"] for x in a),sum("срок" in x["why"] for x in a)])
    S.cell(S.max_row,5).number_format=MONEY
for i,w in enumerate([24,11,12,9,20,14,18,12],1): S.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
H=["№","Номер закупки","Предмет","Заказчик","Регион","Закон","Способ","Площадка","НМЦК, ₽","Окончание подачи","СМП","Обеспечение заявки","Обеспечение контракта","Аванс"]
for cat in cats:
    ws=wb.create_sheet(cat.replace("/"," ")[:30]); ws.append(H)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
    ok=sorted([x for x in d if x["cat"]==cat and x["ok"]],key=lambda x:(x["nmck"] is None,-(x["nmck"] or 0)))
    for i,x in enumerate(ok,1):
        try: e=dt.datetime.strptime(x["end"],"%Y-%m-%d %H:%M")
        except Exception: e=x["end"]
        ws.append([i,x["num"],x["name"],x["cust"],x["region"],x["law"],x["method"],x["etp"],x["nmck"],e,x["smp"],x["sec_app"],x["sec_contr"],x["advance"]])
        R=ws.max_row; ws.cell(R,2).number_format="@"; ws.cell(R,9).number_format=MONEY; ws.cell(R,10).number_format="DD.MM.YYYY HH:MM"
        for cc in ws[R]: cc.border=Border(top=thin,bottom=thin,left=thin,right=thin)
    for i,w in enumerate([4,22,60,34,22,12,24,18,15,16,12,14,14,10],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes="C2"; ws.auto_filter.ref=f"A1:N{ws.max_row}"
from tidy import tidy; tidy(wb,row_h=30)
wb.save("../TRIAGE_1910.xlsx"); print("ok")
