"""Из выгрузок клиента 12.10.2026 — только соль и одежда, с фильтрами."""
import json, re, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
d=json.load(open("data/triage_1210.json"))
TODAY=dt.datetime(2026,10,13)
SALT=re.compile(r"сол[ьи]\b|галит|пескосол|противогол|ПГМ|ПГР",re.I)
CL=re.compile(r"одежд|костюм|куртк|обмундир|форменн|термобел|обув|ботин|берц|СИЗ|средств индивидуальной защиты|мягкого инвентаря|плащ",re.I)
def salt_type(n):
    if re.search(r"пескосол",n,re.I): return "Пескосоляная смесь + соль"
    if re.search(r"галит|сол[ьи]",n,re.I): return "Соль техническая"
    return "Сыпучие/хозтовары (соль может быть в спецификации — проверить)"
def cl_type(n):
    if re.search(r"форменн|обмундир",n,re.I): return "Форменная одежда / обмундирование"
    if re.search(r"термобел",n,re.I): return "Термобельё"
    if re.search(r"мягкого инвентаря|воспитанник|студент|одежды$|Закупка одежды|ГАЛАНТЕРЕЯ",n,re.I): return "Повседневная одежда (соцучреждения)"
    if re.search(r"обув|ботин|берц",n,re.I) and not re.search(r"одежд",n,re.I): return "Спецобувь"
    if re.search(r"плащ",n,re.I): return "Снаряжение (плащ-палатки)"
    return "Спецодежда / СИЗ"
def end_dt(s):
    try: return dt.datetime.strptime(s[:16],"%Y-%m-%d %H:%M")
    except Exception: return None
def status(e):
    if e is None: return "Срок не указан"
    if e<TODAY: return "Срок прошёл"
    if e.date()==TODAY.date(): return "Подача сегодня"
    return "Актуален"
import sys
ONLY=len(sys.argv)>1 and sys.argv[1]=="konkurs"
OKM=re.compile(r"конкурс|котиров|запрос цен|запрос предложен|иного способа",re.I)
FROM=dt.datetime(2026,10,16)
salt=[];cl=[];seen=set()
for x in d:
    n=x["name"]; key=(x["num"],n)
    if key in seen: continue
    if ONLY:
        e=end_dt(x["end"])
        if not OKM.search(x["method"] or "") or e is None or e<FROM: continue
        if re.search(r"спортивн",n,re.I): continue
    if x["cat"]=="Соль и сыпучие" or SALT.search(n):
        if re.search(r"дерматолог|реактив|лаб\. посуд",n,re.I): continue
        seen.add(key); salt.append((salt_type(n),x))
    elif x["cat"] in("Одежда","СИЗ") or CL.search(n):
        if re.search(r"Мебель для склада",n,re.I): continue
        seen.add(key); cl.append((cl_type(n),x))
H=["№","Тип","Статус","Номер закупки","Предмет","Заказчик","Регион","Закон","Способ","Площадка","НМЦК, ₽","Окончание подачи","Только СМП","Обеспечение заявки","Обеспечение контракта","Аванс","Ссылка","Примечание"]
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
FILL={"Подача сегодня":"FCE4D6","Срок прошёл":"EDEDED","Актуален":None,"Срок не указан":None}
wb=openpyxl.Workbook(); wb.remove(wb.active)
def sheet(title,items):
    ws=wb.create_sheet(title)
    ws.append(H)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True,vertical="center")
    rank={"Подача сегодня":0,"Актуален":1,"Срок не указан":2,"Срок прошёл":3}
    rows=[(t,x,end_dt(x["end"])) for t,x in items]
    rows.sort(key=lambda r:(rank[status(r[2])],-(r[1]["nmck"] or 0)))
    for i,(t,x,e) in enumerate(rows,1):
        st=status(e)
        note=[]
        if not x["nmck"]: note.append("НМЦК не указана — цену запрашивать")
        if x["nmck"] and x["nmck"]>50e6: note.append("крупный лот >50 млн")
        if re.search(r"иного способа",x["method"] or "",re.I): note.append("в ЕИС «иной способ» — уточнить в извещении (не аукцион)")
        link=x["link"] or ""
        ws.append([i,t,st,x["num"],x["name"],x["cust"],x["region"],x["law"],x["method"],x["etp"],x["nmck"],e,
                   x["smp"] or "",x["sec_app"] or "",x["sec_contr"] or "",x["advance"] if x["advance"] not in (None,"") else "",
                   "открыть" if link else "","; ".join(note)])
        r=ws.max_row
        if link: ws.cell(r,17).hyperlink=link; ws.cell(r,17).font=Font(color="0563C1",underline="single")
        ws.cell(r,11).number_format=MONEY; ws.cell(r,12).number_format="DD.MM.YYYY HH:MM"; ws.cell(r,4).number_format="@"
        f=FILL[st]
        for c in ws[r]:
            c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
            if f: c.fill=PatternFill("solid",fgColor=f)
    for i,w in enumerate([5,26,15,22,50,34,20,13,16,16,16,17,10,14,14,10,9,30],1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.freeze_panes="E2"; ws.auto_filter.ref=f"A1:R{ws.max_row}"
    return rows
rs=sheet("Соль",salt); rc=sheet("Одежда",cl)
wb.save("../SALT_CLOTHES_KONKURS_1210.xlsx" if ONLY else "../SALT_CLOTHES_1210.xlsx")
import collections
for nm,rows in(("Соль",rs),("Одежда",rc)):
    print(nm,len(rows),collections.Counter(status(e) for _,_,e in rows),collections.Counter(t for t,_,_ in rows))
    act=[x for t,x,e in rows if status(e) in("Актуален","Подача сегодня")]
    print(" актуальных с НМЦК:",sum(1 for x in act if x["nmck"]),"сумма",sum(x["nmck"] or 0 for x in act))
