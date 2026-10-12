"""Внутренняя таблица маржи по 30 лотам из TOP30_TENDERER_1210.xlsx (не для передачи тендерщику)."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
src=open("tenderer_top30.py").read()
ns={}; exec(src[:src.index("assert len(L)")],ns); L=ns["L"]
wb=openpyxl.load_workbook("../TENDERS_1210.xlsx",data_only=True)
INT={r[1]:r for r in wb["Тендеры 12.10"].iter_rows(min_row=2,values_only=True) if r[1]}
def P(nm,c,d):
    p=nm*(1-d); pre=p-c-p*0.008-c*0.2*60/365; return pre-max(pre,0)*0.15
def be(nm,c):
    lo,hi=0.0,1.0
    for _ in range(60):
        m=(lo+hi)/2
        if P(nm,c,m)>0: lo=m
        else: hi=m
    return lo
SALT=2_500
# себестоимость, основа, ожидаемое снижение — для лотов не из выгрузки 12.10 и для формы
OVR={
 "32616441686":(None,"Нет цены за т в извещении — запросить документацию",None),
 "0372200193426000027":(6.5*SALT+5_000,"6,5 т × 2 500 ₽ + фасовка в мешки 5 000 ₽ [оценка]",0.15),
 "32616449090":(553_080,"ПГМ А-Стандарт 23,6 т × 22 800 ₽ + доставка 15 000 ₽",0.15),
 "0164300016526000498":(966_000,"ПГРт 30 т × 30 000 ₽ [оценка] + доставка 66 000 ₽",0.15),
 "0164300016526000499":(325_080,"ПГРт −30 °C 13,6 т, по прайсу поставщика",0.15),
 "0372200258526000596":(8*22_800+20_000,"ПГРт с крошкой 8 т × 22 800 ₽ + доставка 20 000 ₽ [оценка]",0.15),
 "0373100068726000393":(70_714,"ПГР 1,5 т + щебень 1 т, по прайсам",0.15),
 "0358200022926000030":(168*5000+241*3500+208*5000+239*3500+268*350+30_000,
   "Зимн. костюм 168 × 5 000 + летн. костюм 241 × 3 500 + зимн. комбинезон 208 × 5 000 + летн. комбинезон 239 × 3 500 + футболка 268 × 350 [оценка] + доставка 30 000 [оценка]",0.15),
 "32616443256":(405*3500+150*5000+150_000,"Летний комплект 405 × 3 500 + зимний 150 × 5 000 + доставка в Якутию 150 000 [оценка]",0.15),
}
NOTE={"0142200001326018965":"Нужен опыт: исполненный контракт от 24,2 млн ₽; поставка 11 000 т",
      "0358200022926000030":"Пошив с символикой ПСС РО — проверить, входит ли в себестоимость",
      "32616443256":"Подача 15.10 05:00 МСК; доставка в Якутию — уточнить",
      "32616442935":"Себестоимость — оценка, цен по позициям нет",
      "32616441686":"Не считать до получения документации"}
rows=[]
for i,r in enumerate(L,1):
    cat,n,_,_,cust,reg,subj,q,u,up,nm=r[:11]
    if n in OVR: c,basis,d=OVR[n]
    elif cat=="Соль": c,basis,d=q*SALT,f"{q:g} т × 2 500 ₽ (Казахстан, с доставкой)",0.15
    else:
        x=INT[n]; c=x[9]; d=x[11]; basis="Себестоимость из внутреннего расчёта (TENDERS_1210)"
    if c is None: rows.append([i,cat,n,cust,reg,r[11],nm,None,basis,None,None,None,None,None,NOTE.get(n,"")]); continue
    rows.append([i,cat,n,cust,reg,r[11],nm,round(c),basis,(nm-c)/nm,round(P(nm,c,0)),d,round(P(nm,c,d)),be(nm,c),NOTE.get(n,"")])
out=openpyxl.Workbook(); ws=out.active; ws.title="Маржа топ-30"
ws.append(["Маржа по топ-30 лотам — ВНУТРЕННЯЯ (не передавать тендерщику)"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["УСН 15%: прибыль = цена − себестоимость − 0,8% прочих − стоимость денег 20% годовых × 60 дн., минус налог 15% с прибыли. Себестоимость соли 2 500 ₽/т с доставкой (без последней мили)."])
H=["№","Категория","№ закупки","Заказчик","Регион","Подача (МСК)","НМЦК, ₽","Себестоимость, ₽","Из чего себестоимость","Валовая маржа при НМЦК","Чистая прибыль при НМЦК, ₽","Ожидаемое снижение","Чистая прибыль при ожид. снижении, ₽","Снижение до нуля (безубыток)","Комментарий"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496"); c.alignment=Alignment(wrap_text=True,vertical="center")
for r in rows: ws.append(r)
def tot(label,flt):
    s=[r for r in rows if r[7] is not None and flt(r)]
    ws.append(["","",label,"","","",sum(r[6] for r in s),sum(r[7] for r in s),"",None,sum(r[10] for r in s),"",sum(r[12] for r in s),"",""])
    for c in ws[ws.max_row]: c.font=Font(bold=True)
ws.append([])
for cat in dict.fromkeys(r[1] for r in rows): tot(f"Итого: {cat}",lambda r,c=cat:r[1]==c)
tot("ИТОГО без Самары",lambda r:r[2]!="0142200001326018965")
tot("ИТОГО с Самарой",lambda r:True)
thin=Side(style="thin",color="BFBFBF")
fills={"Соль":"E2EFDA","ПГР/ПГМ":"FFF2CC","Медрасходники":"DDEBF7","Форменная одежда":"FCE4D6"}
for row in ws.iter_rows(min_row=4):
    f=fills.get(row[1].value)
    for c in row:
        c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=Border(top=thin,bottom=thin,left=thin,right=thin)
        if f: c.fill=PatternFill("solid",fgColor=f)
    for k in (6,7,10,12): row[k].number_format='# ##0'
    for k in (9,11,13): row[k].number_format='0%'
for i,w in enumerate([4,14,21,34,18,15,14,14,44,11,15,11,16,12,36],1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="D4"
out.save("../MARGIN_TOP30_1210.xlsx")
for r in rows: print(r[0],r[1],r[2],r[6],r[7],None if r[9] is None else f"{r[9]:.0%}",r[10],r[11],r[12],None if r[13] is None else f"{r[13]:.0%}")
