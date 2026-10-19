"""Будущие торги (подача с 20.10.2026) — без аукционов: сроки, условия, себестоимость, маржа, решение.
Собирает: выгрузки 19.10 (triage_1910), прежние расчёты (itog_one), позиции (units_1910), условия (cond_2510), ППУ Уфа (ppu_ufa)."""
import json, re, io, contextlib, datetime as dt, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from tidy import tidy
import ref2026 as ref

src = open("tvc2026.py").read()
exec(src[src.index("def group("):src.index("DLV=")])  # group(item, north) без побочных эффектов модуля

with contextlib.redirect_stdout(io.StringIO()):
    import itog_one as it
    import ppu_ufa as pu

T = json.load(open("data/triage_1910.json"))
U = json.load(open("data/units_1910.json"))
COND = {x["num"]: x for x in json.load(open("data/cond_2510.json"))}
F = dt.datetime(2026, 10, 20)
NOT = re.compile(r"канцеляр|хозяйств|игрушк|реактив|лаб\. посуд|стоматолог|газон|офисн|учебн|театральн|сварк|мотопомп|инструмент|тара |укрывн|СМР|переустройств|реализаци|невостребован|запчаст|специнвентар|оснащени|пневматик|ЗИП|компьютерн|оргтехник|стройматериал", re.I)
TUBE = 9.9  # ₽/шт пробирка вакуумная, с доставкой (расчёт 19.10)
DISC = {"Медрасходники (пробирки)": 0.25, "Одежда/СИЗ": 0.10, "Форма/снаряжение": 0.10}


def model(P, C):
    pre = P - C - 0.008 * P - C * 0.2 * 60 / 365
    return pre - 0.15 * max(pre, 0)


def be(nm, C):
    lo, hi = 0.0, 0.95
    for _ in range(60):
        m = (lo + hi) / 2
        lo, hi = (m, hi) if model(nm * (1 - m), C) > 0 else (lo, m)
    return lo


def cost_units(num, cat):
    """Себестоимость по позициям извещения: пробирки — 9,9 ₽/шт, одежда — справочник 2026 + доставка 35 ₽/кг."""
    pos = U.get(num)
    if not pos: return None, None
    C = nm = 0; miss = []
    for p in pos:
        q, pr, n = p.get("qty") or 0, p.get("price") or 0, p.get("name") or ""
        if "пробирк" in n.lower(): c = TUBE
        elif cat in ("Одежда/СИЗ", "Форма/снаряжение") and (g := group(n)) and ref.base(g):
            c = ref.base(g) + 35 * ref.weight(g)
        else: miss.append(n[:30]); continue
        C += c * q; nm += pr * q
    if miss or not C: return None, "нет себеса: " + "; ".join(miss[:2])
    return C, f"по {len(pos)} поз.: " + ("пробирка 9,9 ₽/шт" if cat.startswith("Мед") else "справочник цен 2026 + доставка")


rows, seen = {}, set()
for x in T:
    if "аукцион" in (x["method"] or "").lower(): continue
    if re.search(r"работ|услуг|ремонт|монтаж", x["name"], re.I) or NOT.search(x["name"]): continue
    try: e = dt.datetime.strptime(x["end"], "%Y-%m-%d %H:%M")
    except Exception: continue
    if e < F: continue
    k = (x["name"].strip().lower()[:60], (x["cust"] or "")[:30], x["nmck"])
    if k in seen: continue
    seen.add(k)
    r = dict(cat=x["cat"], num=x["num"], name=x["name"], cust=x["cust"] or "", law=x["law"], meth=x["method"], nmck=x["nmck"], e=e,
             C=None, basis="", d=DISC.get(x["cat"], 0.10), cond="", src="выгрузка 19.10")
    c = COND.get(x["num"])
    if c: r["cond"] = f"заявка: {c['app']}; контракт: {c['contr']}; аванс: {c['adv']}; оплата: {c['pay']}"
    else:
        sa, sc = x["sec_app"] or "—", x["sec_contr"] or "—"
        r["cond"] = f"заявка: {sa}; контракт: {sc}; аванс: {x['advance'] or 'не указан'}"
    if x["nmck"]:
        C, b = cost_units(x["num"], x["cat"])
        r["C"], r["basis"] = C, b or ""
    rows[x["num"]] = r

# прежние расчёты (лоты 1–9 и др.) — если подача ещё впереди
for (rr, _c) in it.rows:
    cat, num, cust, meth, end, nm, C = rr[0], rr[1], rr[2], rr[3], rr[4], rr[5], rr[6]
    m = re.match(r"(\d\d)\.(\d\d)(?: (\d\d):(\d\d))?", end or "")
    if not m: continue
    e = dt.datetime(2026, int(m.group(2)), int(m.group(1)), int(m.group(3) or 0), int(m.group(4) or 0))
    if e < F and num not in rows: continue
    r = rows.get(num) or dict(cat=cat, num=num, name=cat, cust=cust, law=meth.split(",")[0].split(" ")[0], meth=meth, nmck=nm, e=e, cond=f"обеспечение: {rr[10]}", src="расчёт 14–16.10")
    if _c.startswith("ЖЁСТК"): r["hard"] = _c
    r.update(C=C, basis=rr[11] if len(rr) > 11 else "расчёт 14–16.10 (цены 2026)", d=0.10 if nm else 0.07)
    if not nm: r["kp"] = C / (1 - 0.25)  # коммерческие без НМЦК — КП с наценкой ~33%
    rows[num] = r

# ППУ Уфа — свежий расчёт по заводам
p = rows["32616450774"]
p.update(C=pu.C_med + pu.FREIGHT, basis="4 завода, медиана сайтов 2026 + доставка 50 тыс.", d=0.10,
         cond="обеспечений нет; аванс 50%; остаток 7 раб. дн. после приёмки; ⚠ заказчик в конкурсном производстве (А07-13106/2018)")

out = []
for r in rows.values():
    nm, C = r["nmck"], r["C"]
    if nm and C:
        P = nm * (1 - r["d"]); r["gm"] = 1 - C / nm; r["P"] = P; r["net"] = model(P, C); r["be"] = be(nm, C)
    else: r["gm"] = r["P"] = r["net"] = r["be"] = None
    q = r["meth"].lower()
    if re.search(r"запрос цен|мониторинг", q): r["dec"] = "КП (сбор цен, без обязательств)"
    elif r.get("kp"): r["dec"] = f"КП ≈ {r['kp']:,.0f} ₽ (без НМЦК)".replace(",", " ")
    elif not nm: r["dec"] = "КП — запросить спецификацию"
    elif r["net"] is None: r["dec"] = "Посчитать: нужны позиции/цены"
    elif r["net"] <= 0: r["dec"] = "НЕТ — убыток"
    elif r["gm"] >= 0.2 and r["net"] >= 30000: r["dec"] = "ПОДАЁМ"
    else: r["dec"] = "Осторожно — малая маржа"
    if nm and C and C > 2 * nm: r["dec"] = "Посчитать: себес > НМЦК — сверить кол-во/ед. изм."
    if r.get("hard"): r["dec"] = "Осторожно — " + r["hard"].lower()
    if r["num"] == "32616450774": r["dec"] = "ПОДАЁМ (только после аванса)"
    out.append(r)
ORD = {"ПОДАЁМ": 0, "Осторожно": 1, "КП": 2, "Посчитать": 3, "НЕТ": 4}
out.sort(key=lambda r: (r["e"].date(), next(v for k, v in ORD.items() if r["dec"].startswith(k)), -(r["net"] or 0)))

thin = Side(style="thin", color="BFBFBF"); M = '#,##0" ₽"'; PCT = "0%"
FILL = {"ПОДАЁМ": "C6EFCE", "Осторожно": "FFEB9C", "НЕТ": "FFC7CE", "КП": "DDEBF7", "Посчитать": "F2F2F2"}
wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Будущие торги"
ws.append([f"Будущие торги — подача с 20.10.2026, без аукционов (наш профиль, без дублей). Всего {len(out)}"]); ws["A1"].font = Font(bold=True, size=13)
ws.append(["Себестоимость — по ценам 2026 (заводы/прайсы, с НДС) с доставкой; одежда — 3 500 / 5 000 ₽ за комплект; пробирка 9,9 ₽. Прибыль — после УСН 15% при ожидаемом снижении. Решение: ПОДАЁМ = маржа ≥20% и прибыль ≥30 тыс."])
H = ["№", "Подача до", "Категория", "Номер", "Предмет", "Заказчик", "Закон / способ", "НМЦК, ₽", "Себестоимость, ₽", "Маржа при НМЦК", "Ожид. снижение", "Цена заявки, ₽",
     "Прибыль после УСН, ₽", "Безубыточное снижение", "Решение", "Условия (обеспечение / аванс / оплата)", "Основа себестоимости"]
ws.append(H)
for c in ws[3]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
short = lambda t, n: t if len(t) <= n else t[:n - 1] + "…"
for i, r in enumerate(out, 1):
    ws.append([i, r["e"], r["cat"], r["num"], short(r["name"], 95), short(r["cust"], 60), f'{r["law"]} · {r["meth"]}', r["nmck"], r["C"], r["gm"],
               r["d"] if r["net"] is not None else None, r["P"], r["net"], r["be"], r["dec"], short(r["cond"], 170), short(r["basis"], 90)])
    R = ws.max_row
    ws.cell(R, 2).number_format = "DD.MM HH:MM"; ws.cell(R, 4).number_format = "@"
    for j in (8, 9, 12, 13): ws.cell(R, j).number_format = M
    for j in (10, 11, 14): ws.cell(R, j).number_format = PCT
    ws.cell(R, 15).fill = PatternFill("solid", fgColor=next(v for k, v in FILL.items() if r["dec"].startswith(k)))
    ws.cell(R, 15).font = Font(bold=r["dec"].startswith("ПОДАЁМ"))
    for c in ws[R]: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
go = [r for r in out if r["dec"].startswith("ПОДАЁМ")]
ws.append(["", "", "", "", f"ИТОГО «ПОДАЁМ»: {len(go)} лотов", "", "", sum(r["nmck"] for r in go), sum(r["C"] for r in go), None, None, sum(r["P"] for r in go), sum(r["net"] for r in go)])
R = ws.max_row
for j in (8, 9, 12, 13): ws.cell(R, j).number_format = M
for c in ws[R]: c.font = Font(bold=True)
for i, w in enumerate([4, 11, 15, 21, 40, 28, 24, 13, 13, 9, 9, 13, 13, 11, 22, 46, 30], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ws.freeze_panes = "F4"; ws.auto_filter.ref = f"A3:Q{ws.max_row - 1}"

# сводка по дням
s2 = wb.create_sheet("Календарь")
s2.append(["Дата подачи", "Всего лотов", "ПОДАЁМ", "Прибыль по «ПОДАЁМ», ₽", "Лоты «ПОДАЁМ»"])
for c in s2[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="305496")
for day in sorted({r["e"].date() for r in out}):
    dd = [r for r in out if r["e"].date() == day]; g = [r for r in dd if r["dec"].startswith("ПОДАЁМ")]
    s2.append([dt.datetime.combine(day, dt.time()), len(dd), len(g), sum(r["net"] for r in g) or None, "; ".join(f'{r["num"]} ({short(r["cust"], 25)})' for r in g)])
    s2.cell(s2.max_row, 1).number_format = "DD.MM.YYYY"; s2.cell(s2.max_row, 4).number_format = M
for row in s2.iter_rows():
    for c in row: c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
for i, w in enumerate([13, 11, 10, 18, 90], 1): s2.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=46)
wb.save("../FUTURE_TENDERS_1910.xlsx")
print(len(out), "лотов;", len(go), "ПОДАЁМ; прибыль", round(sum(r["net"] for r in go)))
for r in out: print(r["e"].strftime("%d.%m %H:%M"), r["num"], r["cat"][:12], r["nmck"], round(r["C"] or 0), round((r["gm"] or 0) * 100), round(r["net"] or 0), r["dec"])
