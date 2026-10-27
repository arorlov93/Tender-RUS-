"""Камчатка 32616346943: что не закрыто (data/kamch/open_2710.json) — позиции дословно из ТЗ и кому запросить КП:
уже запрашивали / известные проверенные заводы (запрос не отправлен) / новые компании из исследования рынка 27.10. Без наших цен. → KAMCHATKA_OPEN_ZAPROS_2710.xlsx"""
import json, os, re, openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from collections import defaultdict
from openpyxl.styles import Font, PatternFill, Border, Side
from tidy import tidy

O = json.load(open("data/kamch/open_2710.json"))
thin = Side(style="thin", color="BFBFBF"); BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
ph = lambda s: str(s or "").split(";")[0].split("\n")[0].strip()
# кто уже получил запрос / ответил
ASKED = [("ДН.РУ", "счета 819892, 819896; подбирает клапаны 1.3/2–5, 1.6/92, задвижку 1.1/6", "ответил частично"),
         ("АО «Тулаэлектропривод»", "КП 10/160 на 1.1/24", "ответил"), ("ООО «Лидер-Энерго» / «Транс-Энерго»", "КП на ЗДП 1.3/151–153", "ответил"),
         ("ООО «ТД АДЛ» (Гранвэл)", "КП УТД000112759 на затворы", "ответил"), ("ООО «ЭПАС»", "КП на ЗДП", "ответил (дороже Лидер-Энерго)"),
         ("ООО «Пензапромарматура» (КИПИКА)", "краны для манометров 1.3/39, 1.3/100, 1.6/16, 1.6/50", "ответил, ждёт артикулы")]
ASKED_RX = re.compile(r"ДН\.ру|DN\.ru|Рашворк|Тулаэлектропривод|Лидер-Энерго|Транс-Энерго|ЭПАС|ЭнергоПромАтомСервис|Пензапромарматура|ТД АДЛ|Торговый Дом АДЛ", re.I)
open_tt = {o["tt"] for o in O}
C, U = [], []
for f, src in (("data/kamch/rem_krany.json", "известный"), ("data/kamch/rem_zatvory.json", "известный"),
               ("data/kamch/new_sup_krany.json", "новый (исследование 27.10)"), ("data/kamch/new_sup_other.json", "новый (исследование 27.10)")):
    if not os.path.exists(f): continue
    for x in json.load(open(f)):
        cov = [t for t in (x.get("covers_tt") or []) if t in open_tt]
        if not cov and not src.startswith("новый") or ASKED_RX.search(x["company"]): continue
        (C if x.get("verified") and cov else U).append(dict(src=src, company=x["company"], type=x.get("type") or ("завод" if x.get("is_manufacturer") else ""), city=x.get("city") or x.get("plant_city") or "",
                      site=x.get("site") or "", phone=ph(x.get("phone")), email=ph(x.get("email")), cov=cov, note=str(x.get("covers_note") or "")[:200],
                      reg=str(x.get("registry_mpt") or x.get("gisp") or "")[:80], srcs=" ; ".join(x.get("sources") or [x.get("contact_url") or ""])[:200]))
by_tt = defaultdict(list)
for c in C:
    for t in c["cov"]: by_tt[t].append(c)

wb = openpyxl.Workbook(); sv = wb.active; sv.title = "Итог"
head = lambda ws: [setattr(c, "font", Font(bold=True, color="FFFFFF")) or setattr(c, "fill", PatternFill("solid", fgColor="305496")) for c in ws[ws.max_row] if c.value is not None]
sv.append(["Камчатскэнерго 32616346943 — не закрыто на 27.10.2026: 147 позиций и кому запросить КП"]); sv["A1"].font = Font(bold=True, size=12)
sv.append(["Контакты — только проверенные (телефон и почта совпадают минимум в 2 источниках). Цены в файле не указаны — файл можно пересылать."])
sv.append(["Группа", "Позиций", "Шт", "Кому запросить (проверенные)", "Из них новых", "Позиций без поставщика"]); head(sv)
for g in ("Краны", "Затворы", "Клапаны", "Задвижки", "Приводы и прочее"):
    gs = [o for o in O if o["grp"] == g]
    cs = {c["company"] for o in gs for c in by_tt[o["tt"]]}; nw = {c["company"] for o in gs for c in by_tt[o["tt"]] if c["src"].startswith("новый")}
    sv.append([g, len(gs), sum(o["q"] for o in gs), len(cs), len(nw), sum(1 for o in gs if not by_tt[o["tt"]])])
sv.append(["ИТОГО", len(O), sum(o["q"] for o in O), len({c["company"] for c in C}), len({c["company"] for c in C if c["src"].startswith("новый")}), sum(1 for o in O if not by_tt[o["tt"]])])
for c in sv[sv.max_row]: c.font = Font(bold=True)
sv.append([]); sv.append(["Уже запрашивали", "Что", "Статус"]); head(sv)
for a in ASKED: sv.append(list(a))
for row in sv.iter_rows(min_row=3):
    for c in row:
        if c.value is not None: c.border = BOX
for i, w in enumerate([40, 50, 30, 26, 14, 22], 1): sv.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

ks = wb.create_sheet("Кому запросить")
ks.append(["№", "Компания", "Статус", "Тип", "Город", "Телефон", "E-mail", "Сайт", "Позиций закрывает", "Какие позиции (табл./№ ТЗ)", "Что именно", "Реестр Минпромторга", "Источники контактов"]); head(ks)
C.sort(key=lambda c: (-len(c["cov"]), c["src"]))
for i, c in enumerate(C, 1):
    ks.append([i, c["company"], c["src"], c["type"], c["city"], c["phone"], c["email"], c["site"], len(c["cov"]), ", ".join(c["cov"])[:300], c["note"], c["reg"], c["srcs"]])
    if c["src"].startswith("новый"): ks.cell(ks.max_row, 3).fill = PatternFill("solid", fgColor="C6EFCE")
    for x in ks[ks.max_row]: x.border = BOX
for i, w in enumerate([4, 36, 16, 11, 16, 20, 26, 24, 9, 40, 40, 20, 40], 1): ks.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
ks.freeze_panes = "C2"; ks.auto_filter.ref = f"A1:M{ks.max_row}"

us = wb.create_sheet("Кандидаты без проверки")
us.append(["№", "Компания", "Тип", "Город", "Телефон (не подтверждён в 2 источниках)", "E-mail", "Сайт", "Позиций", "Какие позиции", "Что именно / почему не подтверждено", "Источники"]); head(us)
for i, c in enumerate(sorted(U, key=lambda c: -len(c["cov"])), 1):
    us.append([i, c["company"], c["type"], c["city"], c["phone"], c["email"], c["site"], len(c["cov"]), ", ".join(c["cov"])[:300], c["note"], c["srcs"]])
    for x in us[us.max_row]: x.border = BOX
for i, w in enumerate([4, 36, 12, 18, 22, 26, 24, 8, 40, 50, 40], 1): us.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

F = ["Наименование", "ГОСТ/ТУ", "Ду, мм", "Проход", "Давление, МПа", "Температура, °С", "Среда", "Материал корпуса", "Климатическое исполнение", "Присоединение",
     "Приводное устройство: наименование", "Приводное устройство: технические характеристики", "Установочное положение", "Материал штока, шара, уплотнений", "Габариты", "Примечание"]
for g in ("Краны", "Затворы", "Клапаны", "Задвижки", "Приводы и прочее"):
    ws = wb.create_sheet(g); ws.append(["№", "Табл./№ ТЗ", *F, "Кол-во", "Ед.", "Кому запросить", "Почему не закрыто"]); head(ws)
    for i, o in enumerate(sorted([o for o in O if o["grp"] == g], key=lambda o: (o["tt"].split("/")[0], int(o["tt"].split("/")[1]))), 1):
        who = "; ".join(f'{c["company"].split("(")[0].strip()} {c["phone"]}' for c in by_tt[o["tt"]]) or "нет проверенного — искать дальше"
        ws.append([i, o["tt"], *[ILLEGAL_CHARACTERS_RE.sub("", str(o["tz"].get(k, ""))) for k in F], o["q"], o["tz"].get("unit", "шт"), who[:400], o["why"]])
        X = ws.max_row
        ws.cell(X, 2).fill = PatternFill("solid", fgColor="FFF2CC")
        if not by_tt[o["tt"]]: ws.cell(X, 21).fill = PatternFill("solid", fgColor="F8CBAD")
        for x in ws[X]: x.border = BOX
    for i, w in enumerate([4, 9, 36, 14, 7, 10, 9, 12, 14, 14, 10, 14, 14, 16, 10, 22, 12, 14, 7, 5, 50, 30], 1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "D2"; ws.auto_filter.ref = f"A1:V{ws.max_row}"
pl = wb.create_sheet("Письма", 2)
pl.append(["№", "Кому", "E-mail", "Тема", "Текст письма (позиции — во вложении: лист группы, отфильтровать по «Табл./№ ТЗ»)"]); head(pl)
TT = {o["tt"]: o for o in O}
for i, c in enumerate(C, 1):
    lst = "\n".join(f'{t} — {TT[t]["name"]} — {TT[t]["q"]:g} {TT[t]["tz"].get("unit", "шт")}' for t in sorted(c["cov"], key=lambda t: (t.split("/")[0], int(t.split("/")[1]))))
    txt = ("Добрый день!\n\nООО «ГИДРОПАС» (ИНН 9729231965) готовит заявку на поставку трубопроводной арматуры для АО «Камчатскэнерго» (закупка 32616346943). "
           "Просим коммерческое предложение на позиции ниже — строго по техническим требованиям (полные характеристики — во вложении):\n\n" + lst +
           "\n\nВ КП просим указать: цену за единицу с НДС, срок изготовления/отгрузки, условия оплаты, страну производства и номер реестровой записи Минпромторга (ПП 719), "
           "наличие сертификатов/деклараций ТР ТС 010/2011 и 032/2013, вес и габариты для доставки до Петропавловска-Камчатского. "
           "Если какая-то позиция выпускается в другом исполнении — укажите отличия от ТЗ.\n\nОтвет нужен до 30.10.2026.\n\nС уважением,\nООО «ГИДРОПАС»")
    pl.append([i, c["company"], c["email"], f"Запрос КП: арматура для Камчатскэнерго, {len(c['cov'])} поз.", txt])
    for x in pl[pl.max_row]: x.border = BOX
for i, w in enumerate([4, 34, 26, 30, 120], 1): pl.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
tidy(wb, row_h=48)
for r in range(2, pl.max_row + 1): pl.row_dimensions[r].height = min(409, 60 + 15 * len(C[r - 2]["cov"]))
wb.save("../KAMCHATKA_OPEN_ZAPROS_2710.xlsx")
print("позиций", len(O), "компаний", len({c["company"] for c in C}), "новых", len({c["company"] for c in C if c["src"].startswith("новый")}), "без поставщика", sum(1 for o in O if not by_tt[o["tt"]]))
