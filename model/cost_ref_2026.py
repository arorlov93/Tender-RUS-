"""Справочник себестоимости 2026: группы товаров (медиана по проверенным ссылкам) + все источники с датой и контактами."""
import json, statistics as st, openpyxl
from openpyxl.styles import Font, PatternFill, Border, Side
import ref2026 as R
bench={b["item"]:b for b in json.load(open("data/deep/bench_clothes.json"))}
TEN={"Костюм летний рабочий":"костюм летний","Костюм летний спасателя/МЧС":"костюм летний","Костюм зимний рабочий (I–III пояс)":"костюм зимний","Костюм зимний северный (IV/особый пояс)":"костюм зимний",
 "Куртка утепленная":"куртка утепленная","Полукомбинезон / брюки зимние":"полукомбинезон","Футболка":"футболка","Халат рабочий":"халат","Жилет сигнальный":"жилет сигнальный",
 "Ботинки кожаные с МП (летние)":"ботинки кожаные с МП","Сапоги ПВХ":"сапоги ПВХ","Сапоги утепленные / валенки":"сапоги утепленные","Костюм сварщика":"костюм сварщика",
 "Костюм от электродуги летний":"костюм от ЭДП","Термобельё (комплект)":"термобелье","Перчатки х/б с ПВХ (пара)":"перчатки х/б с ПВХ","Перчатки трикотажные":"перчатки трикотажные","Рукавицы (пара)":"рукавицы","Каска":"каска"}
DLV={"центр":15,"регионы":35,"Север/ДВ":120}
MONEY='#,##0" ₽"'; thin=Side(style="thin",color="BFBFBF")
wb=openpyxl.Workbook(); ws=wb.active; ws.title="Себестоимость 2026"
ws.append(["Себестоимость по рынку 2026 — одежда, обувь, СИЗ (только проверенные ссылки, 14.10.2026)"]); ws["A1"].font=Font(bold=True,size=13)
ws.append(["Медиана оптовых/прайсовых цен 2026 (без распродаж). Класс A — прайс с датой 2026, B — живой магазин © 2026. Доставка ТК [оценка]: 15/35/120 ₽/кг."])
H=["Товар","Источников","из них A (дата 2026)","Мин., ₽","Медиана рынка, ₽","Макс., ₽","Себестоимость для расчёта, ₽","Вес, кг"]+[f"С доставкой: {k}, ₽" for k in DLV]+["Медиана НМЦК 44-ФЗ, ₽","Наценка к НМЦК (центр)","Лучший источник"]
ws.append(H)
for c in ws[3]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
for g in sorted(R.G):
    r=[x for x in R.G[g] if not R.bad(x)]; ps=[x["price"] for x in r]
    b=R.base(g); w=R.weight(g); dl=[b+w*v for v in DLV.values()]
    tm=bench.get(TEN.get(g,""),{}).get("tender_median"); best=R.sources(g,1)[0]
    ws.append([g,len(r),sum(1 for x in r if x["date_evidence"].startswith("A")),min(ps),st.median(ps),max(ps),b,w]+dl+[tm,(tm/dl[0]-1) if tm else None,best["site"]])
    Rr=ws.max_row
    for j in [4,5,6,7,9,10,11,12]: ws.cell(Rr,j).number_format=MONEY
    ws.cell(Rr,13).number_format="0%"; ws.cell(Rr,14).hyperlink=best["url"]; ws.cell(Rr,14).font=Font(color="0563C1",underline="single")
    if tm: ws.cell(Rr,13).fill=PatternFill("solid",fgColor="C6EFCE" if tm/dl[0]-1>0.3 else ("FFEB9C" if tm/dl[0]-1>0 else "FFC7CE"))
    if g in R.CLIENT: ws.cell(Rr,7).font=Font(bold=True,color="C00000")
for i,w in enumerate([36,9,9,11,12,11,13,7,13,13,13,12,10,22],1): ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
ws.freeze_panes="B4"; ws.auto_filter.ref=f"A3:N{ws.max_row}"
ws.append([]); ws.append(["Красным — ваша себестоимость (летний 3 500 / зимний 5 000); рынок 2026 указан в «Медиана рынка»."])
s2=wb.create_sheet("Источники 2026")
s2.append(["Группа","Компания","Город","Товар","Цена, ₽","Цена до, ₽","Ед.","Тип цены","НДС","Условия","Подтверждение 2026","Телефон","E-mail","Ссылка"])
for c in s2[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
allrows=sorted([x for g in R.G for x in R.G[g]]+R.SALTROWS,key=lambda x:(x["group"],x["price"] or 0))
for x in allrows:
    s2.append([x["group"],x.get("company"),x.get("city"),x.get("product"),x.get("price"),x.get("price_max"),x.get("unit")+(" (KZT)" if x.get("currency")=="KZT" else ""),x.get("price_type"),x.get("vat"),x.get("conditions"),x.get("date_evidence"),x.get("phone"),x.get("email"),x.get("url")])
    Rr=s2.max_row
    for j in (5,6): s2.cell(Rr,j).number_format='#,##0.##'
    s2.cell(Rr,14).hyperlink=x["url"]; s2.cell(Rr,14).font=Font(color="0563C1",underline="single")
    ev=x.get("date_evidence") or ""
    s2.cell(Rr,11).fill=PatternFill("solid",fgColor="C6EFCE" if ev.startswith("A") else ("FFEB9C" if ev.startswith(("B","C+")) else "FCE4D6"))
    if R.bad(x): s2.cell(Rr,4).font=Font(color="808080",italic=True)
for i,w in enumerate([30,28,14,44,11,11,8,14,11,40,40,18,26,50],1): s2.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
s2.freeze_panes="B2"; s2.auto_filter.ref=f"A1:N{s2.max_row}"
s3=wb.create_sheet("Соль 2026 — по лотам")
s3.append(["Лот","Себестоимость с доставкой, ₽/т","Как получено","Ссылки"])
for c in s3[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="305496")
for n,(c,t,U) in R.SALT2026.items():
    s3.append([n,c,t,"\n".join(U)]); s3.cell(s3.max_row,2).number_format=MONEY
for i,w in enumerate([22,16,80,60],1): s3.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
from tidy import tidy; tidy(wb,row_h=30)
wb.save("../COST_REFERENCE_2026.xlsx")
print(len(allrows),"источников")
